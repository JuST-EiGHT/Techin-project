"""Шаг 10.13: «Все товары» — продано, выручка и прибыль за период (из «Прибыли по товарам»: без продаж при данных — 0,
без данных за период — нет данных, без себестоимости — нет прибыли), сортировка по новым колонкам, индикаторы проблем
(убыток, низкая маржа, нет себестоимости, недавно изменилась цена), выбор всех по фильтру, выгрузка выбранных, шаблон
себестоимости по выбранным."""
import csv
import io
import uuid
from datetime import datetime, time, timedelta, timezone
from decimal import Decimal

import openpyxl

from app.engine import products as engine
from app.models.product import Product, ProductPrice
from app.services.sync.finance import MSK, today_msk
from tests.test_pnl_api import add_api_day, user_id
from tests.test_report_pnl import OWNER, auth, set_profile

D = Decimal


def test_problem_flags():
    row = {"period_profit": D("-50"), "margin": D("4"), "stock_status": "critical", "days_of_supply": D("3.4"),
           "cost_price": D("100")}
    flags = engine.problem_flags(row, D(10), (today_msk(), D("120"), D("100")))
    assert [f["code"] for f in flags] == ["loss", "low_margin", "low_stock", "price_changed"]
    assert flags[0]["text"] == "Убыток за период: −50 ₽" and flags[2]["text"] == "Критично мало остатков: запаса на 3,4 дн."
    assert flags[3]["text"].endswith("120 ₽ → 100 ₽")
    row = {"period_profit": None, "margin": D("-3"), "stock_status": "normal", "days_of_supply": None, "cost_price": None,
           "price": D("970"), "breakeven_price": D("1000")}
    assert [f["text"] for f in engine.problem_flags(row, D(10), None)] == [
        "Цена 970 ₽ ниже минимальной 1 000 ₽: маржа −3 %", "Нет себестоимости — маржу и прибыль не посчитать"]


async def seed(client, tokens, sessions):
    await set_profile(client, tokens, OWNER)
    for offer, cost in (("A-1", "100"), ("B-2", "400")):
        await client.put(f"/api/v1/costs/{offer}", json={"cost_price": cost, "vat_multiplier": "1"}, headers=auth(tokens))
    uid = await user_id(client, tokens)
    today = today_msk()
    # A: 2 шт. по 500 — прибыль 1000 − 150 − 200 = 650; B: 1 шт. за 500, логистика 600 — убыток; C — без себестоимости
    await add_api_day(sessions, uid, today - timedelta(days=3), [
        ("Выручка", "1000", 2, "A-1"), ("Вознаграждение за продажу", "-150", None, "A-1"),
        ("Выручка", "500", 1, "B-2"), ("Логистика", "-600", None, "B-2"), ("Выручка", "300", 3, "C-3"),
    ])
    now = datetime.now(timezone.utc)
    async with sessions() as db:
        products = {p.offer_id: p for p in (await db.execute(Product.__table__.select())).mappings()}
        db.add(Product(user_id=uid, ozon_product_id=99, ozon_sku=9999, offer_id="D-4", name="Без продаж", synced_at=now))
        at = lambda days: datetime.combine(today - timedelta(days=days), time(12), MSK).astimezone(timezone.utc)  # noqa: E731
        a = products["A-1"]["id"]
        db.add_all([ProductPrice(product_id=a, snapshot_at=at(20), price=D("500")),
                    ProductPrice(product_id=a, snapshot_at=at(2), price=D("550"))])
        await db.commit()
    return uid


async def page(client, tokens, **params):
    r = await client.get("/api/v1/products", params=params, headers=auth(tokens))
    assert r.status_code == 200, r.text
    return r.json()


async def test_period_columns_sort_and_flags(client, tokens, sessions):
    await seed(client, tokens, sessions)
    body = await page(client, tokens, sort="profit", order="desc")
    assert body["period"]["data_source"] == "api" and body["period"]["days_total"] == 30
    rows = {r["offer_id"]: r for r in body["items"]}
    assert (rows["A-1"]["sold_qty"], rows["A-1"]["revenue"], rows["A-1"]["period_profit"]) == (2, "1000.00", "650.00")
    assert (rows["B-2"]["period_profit"], rows["C-3"]["period_profit"], rows["C-3"]["sold_qty"]) == ("-500.00", None, 3)
    assert (rows["D-4"]["sold_qty"], rows["D-4"]["revenue"], rows["D-4"]["period_profit"]) == (0, "0.00", "0.00")
    # по прибыли: 650, 0, −500, без себестоимости — в конце
    assert [r["offer_id"] for r in body["items"]] == ["A-1", "D-4", "B-2", "C-3"]
    by_sold = await page(client, tokens, sort="sold", order="desc")
    assert [r["offer_id"] for r in by_sold["items"]] == ["C-3", "A-1", "B-2", "D-4"]
    by_revenue = await page(client, tokens, sort="revenue", order="asc")
    assert [r["offer_id"] for r in by_revenue["items"]] == ["D-4", "C-3", "B-2", "A-1"]

    flags = {o: [f["code"] for f in r["flags"]] for o, r in rows.items()}
    assert "loss" in flags["B-2"] and flags["C-3"][-1] == "no_cost" and "no_cost" in flags["D-4"]
    price = next(f for f in rows["A-1"]["flags"] if f["code"] == "price_changed")
    assert price["text"].endswith("500 ₽ → 550 ₽")

    # период без начислений — нет данных, а не 0
    old = await page(client, tokens, date_from="2020-01-01", date_to="2020-01-31")
    assert old["period"]["data_source"] == "none"
    assert {r["sold_qty"] for r in old["items"]} == {None} and {r["period_profit"] for r in old["items"]} == {None}
    bad = await client.get("/api/v1/products", params={"date_from": "2026-02-01", "date_to": "2026-01-01"},
                           headers=auth(tokens))
    assert bad.status_code == 400


async def test_select_by_filter_export_and_cost_template(client, tokens, sessions):
    await seed(client, tokens, sessions)
    ids = (await client.get("/api/v1/products/ids", params={"filters": "no_cost"}, headers=auth(tokens))).json()
    listed = await page(client, tokens, filters="no_cost")
    assert ids["total"] == 2 and set(ids["ids"]) == {r["id"] for r in listed["items"]}
    everything = (await client.get("/api/v1/products/ids", headers=auth(tokens))).json()
    assert everything["total"] == 4

    rows = {r["offer_id"]: r["id"] for r in (await page(client, tokens))["items"]}
    r = await client.post("/api/v1/products/export/selected", json={"product_ids": [rows["B-2"], rows["A-1"]]},
                          headers=auth(tokens))
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    table = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert "Прибыль до налога за период, руб." in table[0] and [line[0] for line in table[1:]] == ["A-1", "B-2"]
    profit_col = table[0].index("Прибыль до налога за период, руб.")
    assert table[2][profit_col] == "-500,00" and "Убыток за период" in table[2][table[0].index("Проблемы")]

    r = await client.post("/api/v1/costs/template/selected", json={"product_ids": [rows["C-3"], rows["A-1"]]},
                          headers=auth(tokens))
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    offers = [ws.cell(i, 1).value for i in range(2, ws.max_row + 1) if ws.cell(i, 1).value]
    assert sorted(offers) == ["A-1", "C-3"]
    bad = await client.post("/api/v1/costs/template/selected", json={"product_ids": [str(uuid.uuid4())]},
                            headers=auth(tokens))
    assert bad.status_code == 400 and "нет в этом магазине" in bad.json()["detail"]
