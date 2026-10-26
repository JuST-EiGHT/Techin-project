"""Планирование поставок (шаг 5.2, BUSINESS_LOGIC §5): расчёт, API, Excel, параметры пользователя."""
import io
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.engine.stocks import Thresholds
from app.engine.supply_planner import CRITICAL, PLANNED, PURCHASE, SUPPLY, URGENT, Demand, Params, Point, plan, recommend
from app.models.posting import Posting, PostingItem, PostingSync
from app.models.product import Product
from app.services.sync.finance import today_msk
from tests.test_sync_stocks import FBS_WH, api_user, auth
from tests.test_sync_stocks import ozon as stocks_ozon  # noqa: F401 — фикстура остатков


def point(stock: int, sold_30: int, sold_7: int, transit: int = 0, days_30: int = 30, days_7: int = 7,
          offer: str = "A", dest: str = "москва", kind: str = SUPPLY) -> Point:
    return Point(offer, offer, kind, dest, stock, transit, Demand(sold_30, sold_7, days_30, days_7))


# ── расчёт ──
def test_formula_by_hand():
    """60 шт. за 30 дней (2 в день), за 7 дней — 21 (3 в день, тренд 1,5); на складе 30, в пути 10.
    Хватит на 30 / 3 = 10 дней — «Скоро» (меньше 14); цель 2 × 45 = 90; к приходу поставки (через 7 дней) нужно
    2 × (45 + 7) = 104 → отправить 104 − 30 − 10 = 64."""
    r = recommend(point(30, 60, 21, transit=10))
    assert (r.avg_daily, r.rate, r.trend) == (Decimal("2.000"), Decimal("3.000"), Decimal("1.50"))
    assert r.days_left == Decimal("10.0")
    assert (r.target_stock, r.need, r.to_order, r.urgency, r.arrives_late) == (90, 104, 64, URGENT, False)


def test_out_of_stock_is_critical_and_late():
    r = recommend(point(0, 15, 7))  # 0,5 в день, кончился
    assert (r.days_left, r.urgency, r.arrives_late, r.to_order) == (Decimal("0.0"), CRITICAL, True, 26)  # ⌈0,5 × 52⌉
    r = recommend(point(5, 30, 7))  # 1 в день: хватит на 5 дней — поставка (7 дней) не успеет
    assert (r.urgency, r.arrives_late, r.to_order) == (CRITICAL, True, 47)


def test_no_recommendation():
    assert recommend(point(200, 60, 14)) is None  # запаса хватает: 104 ≤ 200
    assert recommend(point(50, 60, 14, transit=60)) is None  # уже везут
    assert recommend(point(0, 0, 0)) is None  # продаж за 30 дней нет — не планируем


def test_no_sales_last_week_is_planned():
    r = recommend(point(10, 30, 0))
    assert (r.days_left, r.trend, r.urgency, r.to_order) == (None, Decimal("0.00"), PLANNED, 42)


def test_parameters_and_thresholds():
    r = recommend(point(30, 60, 21), Params(target_days=30, lead_time_days=0), Thresholds(3, 5, 60))
    assert (r.need, r.to_order, r.urgency) == (60, 30, PLANNED)  # 10 дней при порогах 3 / 5 — плановая
    r = recommend(point(30, 60, 21), Params(target_days=60, lead_time_days=14))
    assert (r.need, r.to_order, r.arrives_late) == (148, 118, True)  # 10 дней < 14 в пути


def test_partial_coverage():
    """Заказы загружены только за 10 дней — среднее по ним, а не по 30."""
    r = recommend(point(0, 20, 14, days_30=10, days_7=7))
    assert (r.avg_daily, r.need) == (Decimal("2.000"), 104)


def test_plan_order():
    recs = plan([
        point(100, 60, 14, offer="planned"),  # 50 дней
        point(0, 30, 7, offer="empty"),  # 0 дней
        point(20, 60, 14, offer="soon"),  # 10 дней
        point(5, 30, 7, offer="critical", kind=PURCHASE, dest="wh"),  # 5 дней
        point(500, 60, 14, offer="enough"),
    ])
    assert [(r.point.offer_id, r.urgency) for r in recs] == [
        ("empty", CRITICAL), ("critical", CRITICAL), ("soon", URGENT), ("planned", PLANNED),
    ]


# ── API ──
TH, ACB, RS, CM = "6001-2RS.C3.SRL.NAC.TH", "ACB305523.NAC.JP", "6000-2RS.C3.SRL.NAC.JP", "6000.CM.NAC.JP"
URL = "/api/v1/supply/recommendations"
# Остатки после синхронизации (test_sync_stocks): TH — Екатеринбург 9, СПб 10 (в пути 4), свой склад 12 (резерв 2);
# ACB — СПб 5 (в пути 1); RS — Краснодар 0 (в пути 20); CM — свой склад 1.


class Orders:
    def __init__(self, sessions, user_id: uuid.UUID) -> None:
        self.sessions, self.user_id, self.n = sessions, user_id, 0
        self.end = today_msk() - timedelta(days=1)

    async def add(self, scheme: str, days_ago: int, qty: dict[str, int], *, cluster: str | None = None,
                  warehouse: int | None = None, status: str = "delivered") -> None:
        async with self.sessions() as db:
            ids = dict((await db.execute(select(Product.offer_id, Product.id).where(Product.user_id == self.user_id))).all())
            self.n += 1
            day = self.end - timedelta(days=days_ago)
            posting = Posting(user_id=self.user_id, scheme=scheme, posting_number=f"T-{self.n}", status=status,
                              ordered_at=datetime(day.year, day.month, day.day, 9, tzinfo=timezone.utc), day=day,
                              warehouse_id=warehouse, cluster_to=cluster, synced_at=datetime.now(timezone.utc))
            posting.items = [PostingItem(offer_id=o, product_id=ids.get(o), quantity=q) for o, q in qty.items()]
            db.add(posting)
            await db.commit()

    async def loaded(self, scheme: str, days: int = 90) -> None:
        async with self.sessions() as db:
            db.add(PostingSync(user_id=self.user_id, scheme=scheme, loaded_from=today_msk() - timedelta(days=days - 1),
                               synced_at=datetime.now(timezone.utc)))
            await db.commit()


@pytest.fixture
async def seller(sessions, client, tokens, stocks_ozon):  # noqa: F811
    user_id = await api_user(sessions, client, tokens, stocks_ozon)
    orders = Orders(sessions, user_id)
    await orders.loaded("FBO")
    await orders.loaded("FBS")
    # TH в Екатеринбурге: 60 шт. за 30 дней (2 в день), из них 21 за последние 7 (3 в день)
    await orders.add("FBO", 0, {TH: 21}, cluster="Екатеринбург")
    await orders.add("FBO", 10, {TH: 39}, cluster="Екатеринбург")
    # TH в СПб: 6 шт. — на складе 10 и в пути 4, этого хватает
    await orders.add("FBO", 3, {TH: 2}, cluster="Санкт-Петербург и СЗО")
    await orders.add("FBO", 20, {TH: 4}, cluster="Санкт-Петербург и СЗО")
    # ACB покупают в Новосибирске, где товара нет: 15 шт., 7 из них за неделю; отменённый заказ не считается
    await orders.add("FBO", 2, {ACB: 7}, cluster="Новосибирск")
    await orders.add("FBO", 15, {ACB: 8}, cluster="Новосибирск")
    await orders.add("FBO", 1, {ACB: 10}, cluster="Санкт-Петербург и СЗО", status="cancelled")
    # RS — кластер в заказах написан иначе («краснодар »); на складе 0, в пути 20
    await orders.add("FBO", 4, {RS: 7}, cluster="краснодар ")
    await orders.add("FBO", 12, {RS: 23}, cluster="Краснодар")
    # старое (за окном) и чужое — не считаются
    await orders.add("FBO", 40, {CM: 50}, cluster="Екатеринбург")
    await orders.add("FBO", 1, {"UNKNOWN": 5}, cluster="Екатеринбург")
    # свой склад: TH 9 шт. за 30 дней, 3 за неделю; свободно 12 − 2 = 10
    await orders.add("FBS", 1, {TH: 3}, warehouse=FBS_WH, status="awaiting_packaging")
    await orders.add("FBS", 9, {TH: 6}, warehouse=FBS_WH)
    return orders


def rows(data) -> list[tuple]:
    return [(i["product"]["offer_id"], i["kind"], i["destination"]["name"], i["urgency"], i["days_left"], i["to_order"])
            for i in data["items"]]


async def test_requires_auth(client):
    assert (await client.get(URL)).status_code == 401
    assert (await client.get("/api/v1/supply/export")).status_code == 401


async def test_recommendations(client, tokens, seller):
    data = (await client.get(URL, headers=auth(tokens))).json()
    assert data["params"] == {"target_days": 45, "lead_time_days": 7, "safety_days": 0} == data["saved"] and data["is_saved"] is True
    assert data["postings_known"] is True and data["fbo"]["days"] == 30 and data["fbo"]["enough"] is True
    assert data["fbo"]["window_to"] == (today_msk() - timedelta(days=1)).isoformat()
    assert rows(data) == [
        (RS, "supply", "Краснодар", "critical", "0.0", 32),  # ⌈1 × 52⌉ − 0 − 20 в пути
        (ACB, "supply", "Новосибирск", "critical", "0.0", 26),  # ⌈0,5 × 52⌉, склада там нет
        (TH, "supply", "Екатеринбург", "critical", "3.0", 95),  # 9 / 3 в день; ⌈2 × 52⌉ − 9
        (TH, "purchase", "Склад Подольск", "planned", "23.3", 6),  # 10 / (3 / 7); ⌈0,3 × 52⌉ − 10
    ]
    th = data["items"][2]
    assert (th["sold_30"], th["sold_7"], th["avg_daily"], th["rate"], th["trend"]) == (60, 21, "2.000", "3.000", "1.50")
    assert (th["stock"], th["in_transit"], th["target_stock"], th["need"], th["arrives_late"]) == (9, 0, 90, 104, True)
    assert th["destination"]["warehouses"] == ["ЕКАТЕРИНБУРГ_РФЦ_НОВЫЙ"]
    rs = data["items"][0]
    assert (rs["stock"], rs["in_transit"], rs["destination"]["warehouses"]) == (0, 20, [])
    own = data["items"][3]
    assert (own["stock"], own["destination"]["warehouse_type"]) == (10, "FBS")  # свободно: 12 − 2 в резерве
    assert data["totals"] == {"critical": 3, "urgent": 0, "planned": 1, "supply_units": 153, "purchase_units": 6,
                              "products": 3}


async def test_parameters_override_and_save(client, tokens, seller):
    data = (await client.get(URL, params={"target_days": 30, "lead_time_days": 0}, headers=auth(tokens))).json()
    assert data["is_saved"] is False
    th = next(i for i in data["items"] if i["product"]["offer_id"] == TH and i["kind"] == "supply")
    assert (th["need"], th["to_order"], th["arrives_late"]) == (60, 51, False)
    r = await client.patch("/api/v1/settings", json={"supply_target_days": 30, "supply_lead_time_days": 0},
                           headers=auth(tokens))
    assert (r.json()["supply_target_days"], r.json()["supply_lead_time_days"]) == (30, 0)
    data = (await client.get(URL, headers=auth(tokens))).json()
    assert data["is_saved"] is True and data["params"] == {"target_days": 30, "lead_time_days": 0, "safety_days": 0}
    # срочность — по порогам остатков пользователя (5.1)
    await client.patch("/api/v1/settings", json={"stock_critical_days": 2, "stock_low_days": 5}, headers=auth(tokens))
    data = (await client.get(URL, headers=auth(tokens))).json()
    assert {(i["product"]["offer_id"], i["kind"]): i["urgency"] for i in data["items"]}[(TH, "supply")] == "urgent"

    r = await client.get(URL, params={"target_days": 0}, headers=auth(tokens))
    assert r.status_code == 400 and r.json()["detail"].startswith("Запас после поставки")


async def test_little_or_no_order_data(sessions, client, tokens, stocks_ozon):  # noqa: F811
    user_id = await api_user(sessions, client, tokens, stocks_ozon)
    data = (await client.get(URL, headers=auth(tokens))).json()
    assert (data["postings_known"], data["fbo"], data["items"]) == (False, None, [])
    orders = Orders(sessions, user_id)
    await orders.loaded("FBO", days=6)  # загружено 5 полных дней — мало, не считаем
    await orders.add("FBO", 0, {TH: 30}, cluster="Екатеринбург")
    data = (await client.get(URL, headers=auth(tokens))).json()
    assert (data["fbo"]["days"], data["fbo"]["enough"], data["items"]) == (5, False, [])


async def test_isolation(client, tokens, seller):
    other = {"email": "other@example.com", "password": "strongpass2"}
    await client.post("/api/v1/auth/register", json=other)
    t2 = (await client.post("/api/v1/auth/login", json=other)).json()
    data = (await client.get(URL, headers=auth(t2))).json()
    assert (data["items"], data["postings_known"], data["has_key"]) == ([], False, False)


async def test_export_xlsx(client, tokens, seller):
    r = await client.get("/api/v1/supply/export", headers=auth(tokens))
    assert r.status_code == 200 and "supply.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Поставки", "Параметры"]
    table = [list(row) for row in wb["Поставки"].iter_rows(values_only=True)]
    assert table[0][:6] == ["Тип", "Куда", "Срочность", "Артикул", "Название", "Отправить / закупить, шт."]
    assert [r[:6] for r in table[1:5]] == [  # по месту — удобно собирать поставку в один кластер
        ["Поставка в Ozon", "Екатеринбург", "Срочно", TH, None, 95],
        ["Поставка в Ozon", "Краснодар", "Срочно", RS, None, 32],
        ["Поставка в Ozon", "Новосибирск", "Срочно", ACB, None, 26],
        ["Закупка на свой склад", "Склад Подольск", "Плановая", TH, None, 6],
    ]
    assert table[-1][0] == "Итого" and table[-1][5] == 159
    notes = [c.value for c in wb["Параметры"]["A"] if c.value]
    assert "Итого: в Ozon — 153 шт., закупить на свои склады — 6 шт." in notes
