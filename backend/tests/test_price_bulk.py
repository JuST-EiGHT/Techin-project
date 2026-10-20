"""Массовое изменение цен (шаг 6.3): стратегии, группы предпросмотра, применение по 100 с паузой (ошибка одного запроса
не останавливает остальные), подтверждения, предупреждения, файл Excel / CSV туда и обратно. Ozon — фейковый (6.2)."""
import io
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal as D
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models.audit import AuditLog
from app.models.price_change import PriceChange
from app.models.product import Product, ProductPrice
from app.models.product_cost import BEGINNING, ProductCost
from app.services import price_bulk
from tests.test_price_updater import OFFER, journal, latest_price, ozon, ready  # noqa: F401 — фикстура фейкового Ozon
from tests.test_pricing import auth

NOW = lambda: datetime.now(timezone.utc)  # noqa: E731


@pytest.fixture(autouse=True)
def pauses(monkeypatch):
    """Пауза между запросами в Ozon — без ожидания, но запоминаем."""
    slept: list[float] = []

    async def sleep(seconds):
        slept.append(seconds)

    monkeypatch.setattr(price_bulk, "asyncio", SimpleNamespace(sleep=sleep))
    return slept


async def bulk_preview(client, tokens, **body):
    return await client.post("/api/v1/prices/bulk/preview", json=body, headers=auth(tokens))


async def bulk_apply(client, tokens, batch_id, **body):
    return await client.post("/api/v1/prices/bulk/apply", json={"batch_id": str(batch_id), **body},
                             headers={**auth(tokens), "user-agent": "pytest-browser"})


def ids(env, *pids):
    return [str(env["products"][p]) for p in pids]


def by_offer(items):
    return {i["product"]["offer_id"]: i for i in items}


async def add_products(sessions, ozon, user_id, n, start=5000, price="1000") -> list[uuid.UUID]:  # noqa: F811
    """n товаров по 1000 ₽ с себестоимостью 400 ₽ × 1,22 (граница 512,40 ₽); Ozon знает их цены."""
    out = []
    async with sessions() as db:
        for i in range(n):
            pid = start + i
            p = Product(user_id=user_id, ozon_product_id=pid, offer_id=f"BULK-{pid}", name=f"Товар {pid}", synced_at=NOW())
            db.add(p)
            await db.flush()
            db.add(ProductPrice(product_id=p.id, snapshot_at=NOW(), price=D(price), commission_pct=D("38"),
                                acquiring=D(price) * D("0.015")))
            db.add(ProductCost(user_id=user_id, offer_id=p.offer_id, cost_price=D("400"), vat_multiplier=D("1.22"),
                               effective_from=BEGINNING, source="manual", created_at=NOW()))
            ozon.items[pid] = {"price": price, "old_price": "0", "min_price": "0", "marketing_seller_price": price}
            out.append(p.id)
        await db.commit()
    return out


async def test_strategies_and_groups(client, tokens, sessions, ozon):  # noqa: F811
    env = await ready(client, tokens, sessions)
    r = await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "10"}, product_ids=ids(env, 1001, 1002, 1003))
    assert r.status_code == 200, r.text
    p = r.json()
    items = by_offer(p["items"])
    assert p["counts"] == {"ok": 2, "warn": 0, "blocked": 1, "unchanged": 0, "total": 3}
    assert (items[OFFER]["group"], items[OFFER]["price_before"], items[OFFER]["price_after"], items[OFFER]["change_pct"]) == (
        "ok", "540.00", "594.00", "10.0")
    assert items["6000-2RS.C3.SRL.NAC.JP"]["price_after"] == "548.90"  # 499 × 1,1
    blocked = items["ACB305523.NAC.JP"]
    assert blocked["group"] == "blocked" and any("Нет себестоимости" in e for e in blocked["errors"])
    assert [i["group"] for i in p["items"]] == ["ok", "ok", "blocked"]  # применится — первыми
    assert p["description"] == "+10 % к текущей цене" and p["price_fresh"] and ozon.imports == []
    async with sessions() as db:
        rows = list((await db.execute(select(PriceChange).where(PriceChange.batch_id == uuid.UUID(p["batch_id"])))).scalars())
    assert len(rows) == 3 and {(r.status, r.is_dry_run, r.change_type) for r in rows} == {("preview", True, "manual_bulk")}
    assert all(r.reason == "+10 % к текущей цене" for r in rows)

    # +30 %: «не больше 20 % за раз» — ближайшая допустимая цена; без него — заблокировано
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "30"}, product_ids=ids(env, 1001))).json()
    [item] = p["items"]
    assert (item["group"], item["price_after"]) == ("ok", "648.00")
    assert "По стратегии — 702,00 ₽, за один раз можно до 648,00 ₽" in item["note"]
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "30", "limit_step": False},
                            product_ids=ids(env, 1001))).json()
    assert p["items"][0]["group"] == "blocked" and any("больше 20 %" in e for e in p["items"][0]["errors"])

    # целевая маржа и минимальная + X % — калькулятором 4.2 (у 1001 логистика 30 ₽ по продажам)
    calc = (await client.post("/api/v1/price/calculate", json={"product_id": str(env["products"][1001])},
                              headers=auth(tokens))).json()["result"]
    p = (await bulk_preview(client, tokens, strategy={"kind": "target_margin", "value": "20"}, product_ids=ids(env, 1001))).json()
    assert p["items"][0]["price_after"] == calc["target_price"] and p["description"] == "Целевая маржа 20 %"
    p = (await bulk_preview(client, tokens, strategy={"kind": "min_plus", "value": "10"}, product_ids=ids(env, 1001))).json()
    wanted = (D(calc["min_price"]) * D("1.1")).quantize(D("0.01"))  # ниже −20 % — за один раз только до 432 ₽
    assert p["items"][0]["price_after"] == "432.00" and f"По стратегии — {wanted:.2f}".replace(".", ",") in p["items"][0]["note"]

    # цена не меняется — менять нечего
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "0"}, product_ids=ids(env, 1001, 1003))).json()
    assert p["batch_id"] is None and p["counts"]["unchanged"] == 2 and p["items"] == []
    # готовые цены и проверки запроса
    p = (await bulk_preview(client, tokens, items=[{"product_id": ids(env, 1001)[0], "price": "560"}], reason="ровняем")).json()
    assert p["items"][0]["price_after"] == "560.00" and p["description"] == "ровняем"
    assert (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "5"})).status_code == 400
    r = await bulk_preview(client, tokens, strategy={"kind": "target_margin", "value": "95"}, product_ids=ids(env, 1001))
    assert r.status_code == 400 and "от 0 до 90" in r.json()["detail"]
    assert (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "5"},
                               product_ids=[str(uuid.uuid4())])).status_code == 404


async def test_apply_bulk(client, tokens, sessions, ozon):  # noqa: F811
    env = await ready(client, tokens, sessions)
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "10"},
                            product_ids=ids(env, 1001, 1002, 1003))).json()
    r = await bulk_apply(client, tokens, p["batch_id"])
    assert r.status_code == 200, r.text
    s = r.json()
    assert s["finished"] and s["counts"]["applied"] == 2 and s["counts"]["not_applied"] == 1  # заблокированный не идёт
    assert "Цены отправлены в Ozon: 2" in s["message"] and "1–3 минуты" in s["message"]
    assert s["batches"] == [{"no": 1, "size": 2, "state": "done", "applied": 2, "rejected": 0, "failed": 0, "unknown": 0}]
    assert ozon.imports == [{"prices": [  # по артикулу
        {"product_id": 1003, "price": "548.90", "currency_code": "RUB"},
        {"product_id": 1001, "price": "594.00", "currency_code": "RUB"}]}]
    async with sessions() as db:
        attempts = list((await db.execute(select(PriceChange).where(
            PriceChange.batch_id == uuid.UUID(p["batch_id"]), PriceChange.is_dry_run.is_(False)))).scalars())
    assert {(a.offer_id, a.status, a.batch_no, a.preview_id is not None) for a in attempts} == {
        (OFFER, "applied", 1, True), ("6000-2RS.C3.SRL.NAC.JP", "applied", 1, True)}
    entries = await journal(sessions, env["user_id"])
    assert len(entries) == 2 and {e.status for e in entries} == {"applied"}
    assert all(len(e.ozon_request["body"]["prices"]) == 1 and e.ozon_request["batch_no"] == 1 for e in entries)
    [http] = [a for a in await journal(sessions, env["user_id"], "http") if a.path == "/api/v1/prices/bulk/apply"]
    assert {e.request_id for e in entries} == {http.request_id}
    assert not [a for a in await journal(sessions, env["user_id"], "http") if "/bulk/preview" in a.path]
    assert (await latest_price(sessions, env["products"][1001])).price == D("594")
    # повторно — нельзя; через час — пауза; GET — тот же итог
    assert (await bulk_apply(client, tokens, p["batch_id"])).status_code == 409
    again = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "1"}, product_ids=ids(env, 1001))).json()
    assert any("мин. назад" in e for e in again["items"][0]["errors"])
    got = (await client.get(f"/api/v1/prices/bulk/{p['batch_id']}", headers=auth(tokens))).json()
    assert got["finished"] and by_offer(got["items"])["ACB305523.NAC.JP"]["status"] == "not_applied"


async def test_batches_limits_and_confirmation(client, tokens, sessions, ozon, pauses):  # noqa: F811
    env = await ready(client, tokens, sessions)
    many = [str(i) for i in await add_products(sessions, ozon, env["user_id"], 501)]
    r = await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "5"}, product_ids=many)
    assert r.status_code == 400 and "не больше 500 товаров, выбрано 501" in r.json()["detail"]

    # больше 10 — подтверждение
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "5"}, product_ids=many[:11])).json()
    assert p["counts"]["ok"] == 11
    r = await bulk_apply(client, tokens, p["batch_id"])
    assert r.status_code == 400 and "подтвердите" in r.json()["detail"] and ozon.imports == []
    assert (await bulk_apply(client, tokens, p["batch_id"], confirmed=True)).json()["counts"]["applied"] == 11

    # больше 50 — слово; 120 товаров — два запроса (100 и 20) с паузой 2 с; ошибка первого не останавливает второй
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "5"}, product_ids=many[100:220])).json()
    assert p["counts"]["ok"] == 120
    r = await bulk_apply(client, tokens, p["batch_id"], confirmed=True)
    assert r.status_code == 400 and "«ПОДТВЕРЖДАЮ»" in r.json()["detail"]
    ozon.fail_next_batch = True  # первый запрос (с повторами клиента) — ошибка Ozon
    r = await bulk_apply(client, tokens, p["batch_id"], confirm_word="подтверждаю")
    assert r.status_code == 200, r.text
    s = r.json()
    assert pauses == [2]
    assert [(b["no"], b["size"], b["state"], b["failed"], b["applied"]) for b in s["batches"]] == [
        (1, 100, "done", 100, 0), (2, 20, "done", 0, 20)]
    assert s["counts"]["applied"] == 20 and s["counts"]["failed"] == 100
    assert "ничего не изменилось" in s["items"][0]["error"]


async def test_warnings_and_changed_price(client, tokens, sessions, ozon):  # noqa: F811
    env = await ready(client, tokens, sessions)
    [wid] = await add_products(sessions, ozon, env["user_id"], 1, start=7000, price="1000")
    async with sessions() as db:  # месяц цена была 600 — сейчас 1000 дальше 50 % от медианы: предупреждение
        db.add(ProductPrice(product_id=wid, snapshot_at=NOW() - timedelta(days=40), price=D("600"),
                            commission_pct=D("38"), acquiring=D("9")))
        await db.commit()
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "5"},
                            product_ids=[str(wid)])).json()
    assert p["items"][0]["group"] == "warn" and "дальше медианы" in p["items"][0]["warnings"][0]
    r = await bulk_apply(client, tokens, p["batch_id"])
    assert r.status_code == 400 and "Применять нечего" in r.json()["detail"]
    assert (await bulk_apply(client, tokens, p["batch_id"], include_warnings=True)).json()["counts"]["applied"] == 1

    # цена в Ozon изменилась после предпросмотра — этот товар не меняем, остальные — да
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "10"},
                            product_ids=ids(env, 1001, 1003))).json()
    ozon.items[1003]["price"] = "520"
    s = (await bulk_apply(client, tokens, p["batch_id"])).json()
    items = by_offer(s["items"])
    assert items[OFFER]["status"] == "applied" and items["6000-2RS.C3.SRL.NAC.JP"]["status"] == "blocked"
    assert "изменилась после предпросмотра: была 499,00 ₽, сейчас 520,00 ₽" in items["6000-2RS.C3.SRL.NAC.JP"]["error"]
    assert (await journal(sessions, env["user_id"]))[-1].status in ("applied", "rejected")


async def test_rejections_expired_and_guards(client, tokens, sessions, ozon):  # noqa: F811
    env = await ready(client, tokens, sessions)
    ozon.import_answer = lambda body: {"result": [
        {"product_id": p["product_id"], "updated": p["product_id"] != 1003,
         "errors": [] if p["product_id"] != 1003 else [{"code": "price_out_of_range", "message": "x"}]}
        for p in body["prices"]]}
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "10"},
                            product_ids=ids(env, 1001, 1003))).json()
    s = (await bulk_apply(client, tokens, p["batch_id"])).json()
    items = by_offer(s["items"])
    assert items[OFFER]["status"] == "applied" and items["6000-2RS.C3.SRL.NAC.JP"]["status"] == "rejected"
    assert "порог категории" in items["6000-2RS.C3.SRL.NAC.JP"]["error"]

    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "-5"}, product_ids=ids(env, 1003))).json()
    async with sessions() as db:
        for row in (await db.execute(select(PriceChange).where(PriceChange.batch_id == uuid.UUID(p["batch_id"])))).scalars():
            row.created_at = row.created_at - timedelta(minutes=16)
        await db.commit()
    r = await bulk_apply(client, tokens, p["batch_id"])
    assert r.status_code == 410 and "устарел" in r.json()["detail"]
    # часть массового изменения одним товаром не применить
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "-5"}, product_ids=ids(env, 1003))).json()
    r = await client.post("/api/v1/prices/apply", json={"preview_id": p["items"][0]["preview_id"]}, headers=auth(tokens))
    assert r.status_code == 400 and "Цены списком" in r.json()["detail"]
    # чужой — не найден
    await client.post("/api/v1/auth/register", json={"email": "other@example.com", "password": "strongpass1"})
    t2 = (await client.post("/api/v1/auth/login", json={"email": "other@example.com", "password": "strongpass1"})).json()
    assert (await bulk_apply(client, t2, p["batch_id"])).status_code == 404
    assert (await client.get(f"/api/v1/prices/bulk/{p['batch_id']}", headers=auth(t2))).status_code == 404
    assert (await bulk_preview(client, t2, strategy={"kind": "percent", "value": "5"}, product_ids=ids(env, 1001))).status_code == 404


# ── Excel / CSV ──
async def download(client, tokens, fmt="xlsx", **params) -> bytes:
    r = await client.get("/api/v1/prices/bulk/template", params={"format": fmt, **params}, headers=auth(tokens))
    assert r.status_code == 200, r.text
    return r.content


async def upload(client, tokens, content: bytes, name="prices.xlsx", path="/api/v1/prices/bulk/upload"):
    return await client.post(path, files={"file": (name, content, "application/octet-stream")}, headers=auth(tokens))


def sheet(content: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(content))
    return [list(r) for r in wb["Цены"].iter_rows(values_only=True)]


async def test_excel_round_trip(client, tokens, sessions, ozon):  # noqa: F811
    env = await ready(client, tokens, sessions)
    content = await download(client, tokens)
    rows = sheet(content)
    assert rows[0] == ["Артикул", "Название (справочно)", "Цена сейчас, руб. (справочно)",
                       "Минимальная по расчёту (справочно)", "Целевая по расчёту (справочно)", "Новая цена, руб."]
    data = {r[0]: r for r in rows[1:]}
    assert set(data) == {OFFER, "ACB305523.NAC.JP", "6000-2RS.C3.SRL.NAC.JP"}
    assert data[OFFER][2] == data[OFFER][5] == 540  # «Новая цена» = текущая
    # файл без правок — 0 изменений
    p = (await upload(client, tokens, content)).json()
    assert p["batch_id"] is None and p["counts"]["total"] == 0 and p["file"] == {"rows": 3, "unchanged": 3, "errors": []}
    csv = await download(client, tokens, "csv")
    p = (await upload(client, tokens, csv, "prices.csv")).json()
    assert p["counts"]["total"] == 0 and p["file"]["unchanged"] == 3

    # правки: цена, ошибка в цене, чужой артикул, повтор; 1002 — без себестоимости (не пройдёт проверки)
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Цены"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == OFFER:
            row[5].value = 600
        elif row[0].value == "ACB305523.NAC.JP":
            row[5].value = 2200
        elif row[0].value == "6000-2RS.C3.SRL.NAC.JP":
            row[5].value = "дорого"
    ws.append(["NOPE-1", None, None, None, None, 100])
    ws.append([OFFER, None, None, None, None, 610])
    buf = io.BytesIO()
    wb.save(buf)
    edited = buf.getvalue()
    r = await upload(client, tokens, edited)
    assert r.status_code == 200, r.text
    p = r.json()
    items = by_offer(p["items"])
    assert set(items) == {OFFER, "ACB305523.NAC.JP"}
    assert items[OFFER]["group"] == "ok" and items[OFFER]["price_after"] == "600.00" and items[OFFER]["line"] == 3
    assert items["ACB305523.NAC.JP"]["group"] == "blocked" and p["description"] == "Цены из файла"
    errors = {e["offer_id"]: e["message"] for e in p["file"]["errors"]}
    assert set(errors) == {"6000-2RS.C3.SRL.NAC.JP", "NOPE-1", OFFER}
    assert "новая цена" in errors["6000-2RS.C3.SRL.NAC.JP"].lower() and "нет среди ваших товаров" in errors["NOPE-1"]
    assert "уже есть в строке 3" in errors[OFFER]
    # файл ошибок: и ошибки разбора, и что не прошло проверки
    r = await upload(client, tokens, edited, path="/api/v1/prices/bulk/upload/errors")
    assert r.status_code == 200
    out = sheet(r.content)
    assert out[0][0] == "Ошибка" and len(out) == 5
    reasons = {row[1]: row[0] for row in out[1:]}
    assert "Нет себестоимости" in reasons["ACB305523.NAC.JP"] and "NOPE-1" in reasons
    # применяем то, что из файла
    s = (await bulk_apply(client, tokens, p["batch_id"])).json()
    assert s["counts"]["applied"] == 1 and ozon.imports[-1]["prices"] == [
        {"product_id": 1001, "price": "600.00", "currency_code": "RUB"}]
    # чужой файл
    wb = Workbook()
    wb.active.append(["что-то", "другое"])
    buf = io.BytesIO()
    wb.save(buf)
    r = await upload(client, tokens, buf.getvalue())
    assert r.status_code == 400 and "«Артикул» и «Новая цена»" in r.json()["detail"]
    fresh = await download(client, tokens)  # без правок — строк с ошибками нет
    r = await upload(client, tokens, fresh, path="/api/v1/prices/bulk/upload/errors")
    assert r.status_code == 400 and "нет строк с ошибками" in r.json()["detail"]
    assert (await client.get("/api/v1/prices/bulk/template", params={"filters": "nope"}, headers=auth(tokens))).status_code == 400
    assert not [a for a in await journal(sessions, env["user_id"], "http") if "/bulk/upload" in a.path]


async def test_audit_entries_kind(client, tokens, sessions, ozon):  # noqa: F811
    env = await ready(client, tokens, sessions)
    p = (await bulk_preview(client, tokens, strategy={"kind": "percent", "value": "10"}, product_ids=ids(env, 1001))).json()
    await bulk_apply(client, tokens, p["batch_id"])
    log = (await client.get("/api/v1/audit/prices", params={"kind": "price"}, headers=auth(tokens))).json()
    assert [(i["action_title"], i["status"]) for i in log["items"]] == [("Изменение цены", "applied")]
    async with sessions() as db:
        assert (await db.scalar(select(AuditLog.ozon_request).where(AuditLog.entity_type == "product_price")))["batch_no"] == 1
