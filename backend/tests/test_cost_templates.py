"""Шаблоны себестоимости (шаг 2.4): скачал → заполнил → загрузил, файл ошибок, записи Excel.

Главное правило (14): файл, скачанный из сервиса и загруженный обратно, принимается без ошибок.
"""

import io
import re
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from urllib.parse import unquote

import openpyxl
import pytest
from sqlalchemy import func, select

from app.models.product import Product
from app.models.product_cost import ProductCost
from app.services.cost_import import COST_SPEC, COST_TITLE
from app.services.templates import (
    FORMULA,
    INSTRUCTION_SHEET,
    CellError,
    parse_day,
    parse_key,
    parse_money,
    parse_rate,
    rate_text,
)
from tests.test_costs import ART, XLSX_TYPE, auth, get_report, import_file, upload_report, xlsx

TEMPLATES = Path(__file__).parent / "fixtures" / "templates"
TOP = "6004-2RS.CM.SRL.NAC.TH"  # без себестоимости, продано 19 шт. — первый в шаблоне


async def setup_articles(client, tokens, sessions) -> None:
    """17 артикулов: 16 из отчёта (у одного есть цена) и 1 только из каталога Ozon."""
    me = (await client.get("/api/v1/auth/me", headers=auth(tokens))).json()
    async with sessions() as db:
        db.add(
            Product(
                user_id=uuid.UUID(me["id"]),
                ozon_product_id=1,
                ozon_sku=777,
                offer_id="CATALOG-ONLY",
                name="Товар из каталога",
                synced_at=datetime.now(timezone.utc),
            )
        )
        await db.commit()
    await upload_report(client, tokens)
    r = await client.put(
        f"/api/v1/costs/{ART}",
        json={"cost_price": "164", "vat_multiplier": "1.22", "note": "партия 1"},
        headers=auth(tokens),
    )
    assert r.status_code == 200


async def template(client, tokens, rows="all", fmt="xlsx"):
    r = await client.get("/api/v1/costs/template", params={"rows": rows, "format": fmt}, headers=auth(tokens))
    assert r.status_code == 200, r.text
    return r


def book(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content))


def save(wb: openpyxl.Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


async def count_records(sessions) -> int:
    async with sessions() as db:
        return (await db.execute(select(func.count()).select_from(ProductCost))).scalar_one()


# ---------- шаблон ----------


async def test_xlsx_template_layout_and_rows(client, tokens, sessions):
    await setup_articles(client, tokens, sessions)
    r = await template(client, tokens)
    assert r.headers["content-type"] == XLSX_TYPE
    disposition = r.headers["content-disposition"]
    assert disposition.startswith('attachment; filename="costs.xlsx"')
    assert unquote(disposition.split("filename*=UTF-8''")[1]).startswith("Себестоимость — все товары ")

    wb = book(r.content)
    assert wb.sheetnames == ["Себестоимость", INSTRUCTION_SHEET]  # таблица — первой, её и откроет Excel
    ws = wb["Себестоимость"]
    assert [c.value for c in ws[1]] == [
        "Артикул",
        "Название (справочно)",
        "SKU (справочно)",
        COST_TITLE,
        "Ставка НДС поставщика",
        "Действует с",
        "Комментарий",
    ]
    assert ws.freeze_panes == "A2"
    # артикул — «Текст», в том числе у строк, которые пользователь допишет: Excel не превратит 1-2 в дату
    assert ws["A2"].number_format == "@" and ws.column_dimensions["A"].number_format == "@"
    assert ws["D2"].number_format == "#,##0.00" and ws.column_dimensions["F"].number_format == "DD.MM.YYYY"
    validations = {str(dv.sqref).split(":")[0]: dv for dv in ws.data_validations.dataValidation}
    assert validations["D2"].type == "decimal" and validations["F2"].type == "date"
    assert validations["E2"].type == "list" and "без НДС" in validations["E2"].formula1
    assert validations["E2"].showErrorMessage is False  # можно ввести и «1,22»

    rows = table(ws)
    assert len(rows) == 17
    assert rows[0]["Артикул"] == TOP and rows[0][COST_TITLE] is None  # без цены и самые продаваемые — первыми
    priced = next(x for x in rows if x["Артикул"] == ART)
    assert (priced[COST_TITLE], priced["Ставка НДС поставщика"], priced["Комментарий"]) == (164, "22 %", "партия 1")
    assert priced["Действует с"] is None  # новая цена из шаблона — с сегодня, а не задним числом
    catalog = next(x for x in rows if x["Артикул"] == "CATALOG-ONLY")
    assert (catalog["Название (справочно)"], catalog["SKU (справочно)"]) == ("Товар из каталога", "777")

    missing = table(book((await template(client, tokens, "missing")).content)["Себестоимость"])
    assert len(missing) == 16 and ART not in {x["Артикул"] for x in missing}
    assert table(book((await template(client, tokens, "empty")).content)["Себестоимость"]) == []


async def test_csv_template_opens_in_russian_excel(client, tokens, sessions):
    await setup_articles(client, tokens, sessions)
    r = await template(client, tokens, fmt="csv")
    assert r.headers["content-type"].startswith("text/csv")
    assert r.content.startswith(b"\xef\xbb\xbf")  # BOM: Excel сам поймёт UTF-8
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0].split(";")[:4] == ["Артикул", "Название (справочно)", "SKU (справочно)", COST_TITLE]
    assert f"{ART};" in next(x for x in lines if x.startswith(ART)) and ";164,00;22 %;;партия 1" in "\n".join(lines)
    assert len(lines) == 18


@pytest.mark.parametrize("fmt", ["xlsx", "csv"])
@pytest.mark.parametrize("rows", ["all", "missing", "empty"])
async def test_downloaded_template_uploads_back_without_errors(client, tokens, sessions, rows, fmt):
    await setup_articles(client, tokens, sessions)
    content = (await template(client, tokens, rows, fmt)).content
    body = (await import_file(client, tokens, f"template.{fmt}", content, dry_run=True)).json()
    expected = {
        "all": {"new": 0, "changed": 0, "unchanged": 1, "skipped": 16, "error": 0},
        "missing": {"new": 0, "changed": 0, "unchanged": 0, "skipped": 16, "error": 0},
        "empty": {"new": 0, "changed": 0, "unchanged": 0, "skipped": 0, "error": 0},
    }[rows]
    assert body["counts"] == expected, body["rows"][:3]
    # применение без правок ничего не меняет
    applied = (await import_file(client, tokens, f"template.{fmt}", content)).json()
    assert applied["applied"] is True and applied["recosted_reports"] == 0
    assert await count_records(sessions) == 1


async def test_filled_template_applies_and_exported_prices_come_back_unchanged(client, tokens, sessions):
    await setup_articles(client, tokens, sessions)
    report = (await client.get("/api/v1/reports", headers=auth(tokens))).json()["items"][0]
    wb = book((await template(client, tokens, "missing")).content)
    ws = wb["Себестоимость"]
    ws["D2"] = 254.76  # TOP, как пользователь: число
    ws["E2"] = "10 %"
    ws["D3"] = "1 250,50 ₽"  # текстом с пробелами и рублём
    ws["F3"] = date(2026, 9, 1)
    ws["G3"] = "партия 7"
    ws.append(["NEW-ART", None, None, 99, "без НДС"])  # дописанная строка
    body = (await import_file(client, tokens, "filled.xlsx", save(wb))).json()
    assert body["counts"] == {"new": 3, "changed": 0, "unchanged": 0, "skipped": 14, "error": 0}
    assert body["recosted_reports"] == 1
    rows = {r["offer_id"]: r for r in body["rows"]}
    assert (rows[TOP]["cost_price"], rows[TOP]["vat_multiplier"]) == ("254.76", "1.100")
    second = next(r for r in body["rows"] if r["line"] == 3)
    assert (second["cost_price"], second["effective_from"], second["note"]) == ("1250.50", "2026-09-01", "партия 7")
    assert rows["NEW-ART"]["vat_multiplier"] == "1.000"
    assert (await get_report(client, tokens, report["id"]))["summary"]["by_article"][TOP]["cost"] == "254.76"

    # выгрузка справочника → загрузка без правок: ничего не меняется (в том числе ставка «без НДС»)
    records = await count_records(sessions)
    for fmt in ("xlsx", "csv"):
        content = (await template(client, tokens, "all", fmt)).content
        again = (await import_file(client, tokens, f"all.{fmt}", content)).json()
        assert again["counts"]["unchanged"] == 4 and again["counts"]["new"] + again["counts"]["changed"] == 0
        assert again["counts"]["error"] == 0
    assert await count_records(sessions) == records


async def test_csv_resaved_by_excel_in_windows_1251(client, tokens, sessions):
    await setup_articles(client, tokens, sessions)
    text = (await template(client, tokens, "missing", "csv")).content.decode("utf-8-sig")
    lines = text.splitlines()
    cells = lines[1].split(";")
    cells[3] = "1 250,5"  # русский Excel: запятая, пробел в разрядах
    lines[1] = ";".join(cells)
    content = "\r\n".join(lines).encode("cp1251", errors="replace")  # как Excel: чего нет в кодировке — «?»
    body = (await import_file(client, tokens, "Себестоимость.csv", content)).json()
    assert body["counts"]["new"] == 1 and body["counts"]["error"] == 0
    assert next(r for r in body["rows"] if r["status"] == "new")["cost_price"] == "1250.50"


@pytest.mark.parametrize("name", ["costs_libreoffice.xlsx", "costs_libreoffice.csv"])
async def test_template_saved_by_libreoffice(client, tokens, name):
    # наш шаблон, заполненный и сохранённый в LibreOffice Calc: формула посчитана, разделитель разрядов
    # в CSV («1,250.50»), артикулы «1-2» и «0042» остались текстом
    body = (await import_file(client, tokens, name, (TEMPLATES / name).read_bytes(), dry_run=True)).json()
    assert body["counts"] == {"new": 5, "changed": 0, "unchanged": 0, "skipped": 0, "error": 0}
    rows = {r["offer_id"]: (r["cost_price"], r["vat_multiplier"], r["effective_from"]) for r in body["rows"]}
    assert rows == {
        "6001": ("164.00", "1.220", None),
        "A-2": ("1250.50", "1.100", "2026-09-01"),
        "A-3": ("164.00", "1.000", None),  # =100+64
        "0042": ("55.00", None, None),
        "1-2": ("99.00", None, None),
    }


# ---------- записи значений ----------


def test_vat_rate_notations():
    for value, expected in [
        ("22 %", "1.220"),
        ("22%", "1.220"),
        ("22", "1.220"),
        (22, "1.220"),
        ("0,22", "1.220"),
        (0.22, "1.220"),
        ("1,22", "1.220"),
        (1.22, "1.220"),
        ("10 %", "1.100"),
        ("1,1", "1.100"),
        ("5", "1.050"),
        ("7 %", "1.070"),
        ("20 %", "1.200"),
        ("без НДС", "1.000"),
        ("Без ндс", "1.000"),
        ("0", "1.000"),
        (0, "1.000"),
        ("0 %", "1.000"),
        ("1", "1.000"),
        ("не облагается", "1.000"),
    ]:
        assert parse_rate(value) == Decimal(expected), value
    assert parse_rate(None) is None and parse_rate("  ") is None
    for bad in ("3", "1,3", "abc", "30 %", 2):
        with pytest.raises(CellError):
            parse_rate(bad)
    assert [rate_text(Decimal(v)) for v in ("1.220", "1.100", "1.000")] == ["22 %", "10 %", "без НДС"]
    assert rate_text(None) == ""


def test_article_mangled_by_excel():
    assert parse_key(6001) == "6001" and parse_key(6001.0) == "6001" and parse_key(" 6001-2RS ") == "6001-2RS"
    assert parse_key("1-2") == "1-2" and parse_key(12.5) == "12.5"
    for bad, text in [
        (datetime(2026, 1, 2), "Excel превратил артикул в дату (02.01.2026)"),
        ("02.янв", "похоже, Excel превратил артикул в дату («02.янв»)"),
        ("2-Jan", "похоже, Excel превратил артикул в дату («2-Jan»)"),
        ("1,23457E+11", "Excel сократил длинный артикул до «1,23457E+11»"),
        (FORMULA, "формула без сохранённого значения"),
    ]:
        with pytest.raises(CellError, match=re.escape(text)):
            parse_key(bad)


def test_money_and_date_values():
    assert parse_money("1 250,50 ₽") == Decimal("1250.50") and parse_money(164) == 164 and parse_money("") is None
    with pytest.raises(CellError, match="«дорого» — не число"):
        parse_money("дорого")
    assert parse_day("01.09.2026") == date(2026, 9, 1) and parse_day(datetime(2026, 9, 1)) == date(2026, 9, 1)
    assert parse_day(46266) == date(2026, 9, 1)  # дата Excel числом — колонке сняли формат даты
    assert parse_day(None) is None
    with pytest.raises(CellError):
        parse_day("вчера")


async def test_article_errors_through_api(client, tokens):
    content = xlsx(
        [
            ["Артикул", COST_TITLE],
            [datetime(2026, 1, 2), 10],
            ["1,23457E+11", 10],
            [6001, 10],
            ["02.янв", None],  # без цены — просто пропуск, даже если артикул испорчен
        ]
    )
    rows = (await import_file(client, tokens, "costs.xlsx", content, dry_run=True)).json()["rows"]
    assert rows[0]["message"].startswith("Excel превратил артикул в дату (02.01.2026) — задайте колонке формат")
    assert rows[1]["message"].startswith("Excel сократил длинный артикул")
    assert (rows[2]["offer_id"], rows[2]["status"]) == ("6001", "new")
    assert rows[3]["status"] == "skipped"


async def test_cost_with_vat_column(client, tokens):
    content = "Артикул;Себестоимость за шт. с НДС, ₽;Ставка НДС поставщика\nA;244;22 %\nB;122;\nC;150;без НДС\n"
    body = (await import_file(client, tokens, "costs.csv", content.encode())).json()
    rows = {r["offer_id"]: r for r in body["rows"]}
    assert (rows["A"]["cost_price"], rows["A"]["cost_with_vat"]) == ("200.00", "244")
    assert rows["B"]["cost_price"] == "100.00"  # ставка не указана, профиля нет — общая ставка 22 %
    assert rows["C"]["cost_price"] == "150.00"
    # НДС поставщика по умолчанию — из налогового профиля
    profile = {"legal_form": "ip", "regime": "usn_income", "supplier_vat_pct": "10"}
    assert (await client.put("/api/v1/tax-profile", json=profile, headers=auth(tokens))).status_code == 200
    content = "Артикул;Себестоимость с НДС\nD;110\n".encode()
    assert (await import_file(client, tokens, "d.csv", content)).json()["rows"][0]["cost_price"] == "100.00"


async def test_formula_without_saved_value(client, tokens):
    wb = openpyxl.Workbook()
    wb.active.append(["Артикул", COST_TITLE, "Ставка НДС поставщика"])
    wb.active.append(["A", "=100+64", None])
    wb.active.append(["B", 10, "=A1"])
    rows = (await import_file(client, tokens, "costs.xlsx", save(wb), dry_run=True)).json()["rows"]
    assert rows[0]["message"] == (
        "Себестоимость: формула без сохранённого значения — откройте файл в Excel и сохраните его"
    )
    assert rows[1]["message"].startswith("Ставка НДС: формула без сохранённого значения")


async def test_instruction_sheet_is_not_a_table(client, tokens):
    wb = book((await template(client, tokens, "empty")).content)
    del wb["Себестоимость"]
    r = await import_file(client, tokens, "costs.xlsx", save(wb), dry_run=True)
    assert r.status_code == 400 and r.json()["detail"].startswith("Не найдены колонки «Артикул» и «Себестоимость»")
    texts = [c.value for row in wb[INSTRUCTION_SHEET].iter_rows() for c in row if c.value]
    assert any(t.startswith("1. Заполните колонку") for t in texts)
    assert {c.title for c in COST_SPEC.template_columns} <= set(texts)


# ---------- файл ошибок ----------


async def errors_file(client, tokens, name: str, content: bytes):
    kind = "text/csv" if name.endswith(".csv") else XLSX_TYPE
    return await client.post("/api/v1/costs/import/errors", headers=auth(tokens), files={"file": (name, content, kind)})


async def test_error_rows_file_fix_and_upload_again(client, tokens, sessions):
    await setup_articles(client, tokens, sessions)
    wb = book((await template(client, tokens, "missing")).content)
    ws = wb["Себестоимость"]
    ws["D2"], ws["D3"], ws["E3"], ws["D4"], ws["D5"] = 100, 0, "3 %", "дорого", 50
    content = save(wb)
    preview = (await import_file(client, tokens, "filled.xlsx", content, dry_run=True)).json()
    assert preview["counts"]["error"] == 2 and preview["counts"]["new"] == 2

    r = await errors_file(client, tokens, "filled.xlsx", content)
    assert r.status_code == 200 and r.headers["content-type"] == XLSX_TYPE
    assert 'filename="costs-errors.xlsx"' in r.headers["content-disposition"]
    fixed = book(r.content)
    ws = fixed["Себестоимость"]
    rows = table(ws)
    assert [c.value for c in ws[1]][:2] == ["Ошибка", "Артикул"]
    assert [x["Ошибка"] for x in rows] == [
        "Себестоимость должна быть больше 0; ставка НДС: нужно 22 %, 20 %, 10 %, 7 %, 5 %, 0 % или «без НДС» "
        "(можно и 1,22)",
        "Себестоимость: «дорого» — не число",
    ]
    assert rows[0]["Ставка НДС поставщика"] == "3 %" and rows[1][COST_TITLE] == "дорого"  # как ввёл пользователь
    assert rows[0]["Название (справочно)"]  # справочные колонки — из каталога и отчётов
    ws["E2"], ws["F2"], ws["E3"] = 12.5, "22 %", 30  # колонки сдвинуты: «Ошибка» — первая
    body = (await import_file(client, tokens, "fixed.xlsx", save(fixed))).json()
    assert body["counts"] == {"new": 2, "changed": 0, "unchanged": 0, "skipped": 0, "error": 0}

    # CSV → файл ошибок тоже CSV; нет ошибок — понятный отказ
    r = await errors_file(client, tokens, "c.csv", "Артикул;Себестоимость\nX;-1\n".encode())
    assert r.headers["content-type"].startswith("text/csv")
    assert r.content.decode("utf-8-sig").splitlines()[1].startswith("Себестоимость должна быть больше 0;X;")
    r = await errors_file(client, tokens, "ok.csv", "Артикул;Себестоимость\nX;1\n".encode())
    assert r.status_code == 400 and r.json()["detail"] == "В файле нет строк с ошибками."


async def test_error_file_keeps_cost_with_vat_column(client, tokens):
    content = "Артикул;Себестоимость с НДС\n;122\n".encode()
    text = (await errors_file(client, tokens, "c.csv", content)).content.decode("utf-8-sig")
    assert "Себестоимость за шт. с НДС, руб." in text.splitlines()[0]
    body = (await import_file(client, tokens, "again.csv", text.replace(";;", ";Z;", 1).encode())).json()
    assert body["counts"]["new"] == 1 and body["rows"][0]["cost_price"] == "100.00"


# ---------- доступ ----------


async def test_template_only_own_articles_and_params(client, tokens, sessions):
    await setup_articles(client, tokens, sessions)
    await client.post("/api/v1/auth/register", json={"email": "other@example.com", "password": "strongpass1"})
    other = (
        await client.post("/api/v1/auth/login", json={"email": "other@example.com", "password": "strongpass1"})
    ).json()
    assert table(book((await template(client, other)).content)["Себестоимость"]) == []
    r = await client.get("/api/v1/costs/template", params={"format": "pdf"}, headers=auth(tokens))
    assert r.status_code == 400 and "Формат файла" in r.json()["detail"]
    assert (await client.get("/api/v1/costs/template")).status_code == 401
    page = (await client.get("/api/v1/costs", params={"missing": True}, headers=auth(tokens))).json()
    assert (page["all_count"], page["missing_count"], page["total"]) == (17, 16, 16)
