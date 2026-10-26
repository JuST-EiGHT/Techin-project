"""Шаг 10.8: «Аналитика продаж» — продажи по товарам (заказы, продано, возвраты и их доля, выручка, средняя цена, к прошлому
периоду), события (цена, акция, нет остатка, снова в наличии), выгрузка в Excel."""
import io
import uuid
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal

from openpyxl import load_workbook

from app.engine import leaders as lead
from app.engine import sales_events as ev
from app.models.accrual import AccrualDay
from app.models.posting import FBO, Posting, PostingItem, PostingSync
from app.models.price_change import APPLIED, MANUAL_SINGLE, PriceChange
from app.models.product import Product, ProductPrice
from app.models.promotion import Promotion, PromotionProduct
from app.models.sales_daily import SalesDaily
from app.models.warehouse import ProductStock, Warehouse
from app.services.sync.finance import MSK, today_msk
from tests.test_products import user_id_of

D = Decimal
D1, D2, D3, D4, D5 = (date(2026, 10, d) for d in (1, 2, 3, 4, 5))


# ── чистые функции ──
def test_price_changes_skip_empty_price():
    points = [(D1, D("100")), (D2, D("0")), (D3, D("120")), (D4, None), (D5, D("110"))]
    assert ev.price_changes(D("100"), points) == [(D3, D("100"), D("120")), (D5, D("120"), D("110"))]
    assert ev.price_changes(None, [(D1, D("100")), (D2, D("90"))]) == [(D2, D("100"), D("90"))]  # до периода не знали


def test_stock_changes_total_over_warehouses():
    points = [(D1, "w1", 0), (D2, "w2", 3), (D3, "w1", 2), (D4, "w2", 0), (D5, "w1", 0)]
    assert ev.stock_changes({"w1": 5, "w2": 0}, points) == [(D1, ev.STOCKOUT), (D2, ev.RESTOCK), (D5, ev.STOCKOUT)]
    assert ev.stock_changes({}, [(D1, "w1", 0), (D2, "w1", 4)]) == [(D2, ev.RESTOCK)]  # первое значение — только запомнить
    assert ev.stock_changes({}, [(D1, "w1", 3)]) == []


def test_events_by_day_one_or_many():
    a, b = ev.ProductRef(uuid.uuid4(), "A-1", None), ev.ProductRef(uuid.uuid4(), "B-2", None)
    raw = [ev.price_raw(D2, a, D("100"), D("120"), True), ev.price_raw(D2, b, D("50"), D("45"), False),
           ev.stock_raw(D1, ev.STOCKOUT, b), ev.promo_raw(D2, ev.PROMO_IN, a, uuid.uuid4(), "Распродажа", True)]
    events = ev.by_day(raw)
    assert [(e.day, e.kind, e.count) for e in events] == [(D1, "stockout", 1), (D2, "price", 2), (D2, "promo_in", 1)]
    assert events[0].title == "B-2 закончился на складах" and events[0].link == f"/dashboard/products/{b.id}"
    assert events[1].title == "Цена изменилась у 2 товаров" and events[1].link == "/dashboard/products/prices/history"
    assert [i.text for i in events[1].items] == ["↑ 100 ₽ → 120 ₽, через сервис", "↓ 50 ₽ → 45 ₽"]
    assert events[2].title == "A-1 вошёл в акцию «Распродажа» — Ozon добавил сам"
    [one] = ev.by_day([ev.price_raw(D2, a, D("100"), D("120"), False)])
    assert one.title == "Цена A-1 ↑ 100 ₽ → 120 ₽" and one.link == f"/dashboard/products/prices/history?product={a.id}"


def test_compare_qty_with_days_and_noise():
    assert lead.compare_qty(30, 120, 30, 30) == lead.QtyChange(D("120.0"), D("-75.0"), lead.DOWN)
    assert lead.compare_qty(60, 60, 30, 30).direction is None
    assert lead.compare_qty(3, 2, 30, 30).direction is None  # шум
    assert lead.compare_qty(40, 10, 30, 15) == lead.QtyChange(D("20.0"), D("100.0"), lead.UP)
    assert lead.compare_qty(5, 5, 30, 6) is None


# ── API ──
def auth(tokens):
    return {"Authorization": f"Bearer {tokens['access_token']}"}


def at(day: date, hour: int = 12) -> datetime:
    return datetime.combine(day, time(hour), MSK).astimezone(timezone.utc)


async def seed(sessions, uid) -> dict:
    """X: 2 шт. в день и сейчас, и раньше; 10 заказов по 2 шт. и 2 отменённых; цена 100 → 120 (через сервис) → 115.
    Y: 1 шт. в день (раньше — 4), 1 возврат; акция с −8 по −3 дня; остаток 5 → 0 (−6) → 10 (−2)."""
    today, now = today_msk(), datetime.now(timezone.utc)
    back = lambda n: today - timedelta(days=n)  # noqa: E731
    async with sessions() as db:
        x = Product(user_id=uid, ozon_product_id=1, ozon_sku=11, offer_id="X-1", name="Подшипник X", synced_at=now)
        y = Product(user_id=uid, ozon_product_id=2, ozon_sku=22, offer_id="Y-2", name="Подшипник Y", synced_at=now)
        w = Warehouse(user_id=uid, ozon_warehouse_id=7, name="Хоругвино", type="FBO", cluster_name="Москва", synced_at=now)
        db.add_all([x, y, w])
        await db.flush()
        for n in range(1, 61):
            day = back(n)
            db.add(AccrualDay(user_id=uid, day=day, synced_at=now, entries=1))
            y_qty = 1 if n <= 30 else 4
            for p, qty, price in ((x, 2, 100), (y, y_qty, 300)):
                db.add(SalesDaily(user_id=uid, day=day, offer_id=p.offer_id, product_id=p.id, qty_sold=qty,
                                  qty_returned=1 if (p is y and n == 4) else 0, revenue=D(qty * price),
                                  returns_amount=D(300) if (p is y and n == 4) else D(0), source="api", calculated_at=now))
        db.add(PostingSync(user_id=uid, scheme=FBO, loaded_from=back(40), synced_at=now))
        for i in range(12):
            post = Posting(user_id=uid, scheme=FBO, posting_number=f"p-{i}", day=back(i + 1),
                           status="cancelled" if i >= 10 else "delivered", ordered_at=at(back(i + 1)), synced_at=now)
            post.items = [PostingItem(sku=11, offer_id="X-1", product_id=x.id, quantity=2)]
            db.add(post)
        db.add(ProductPrice(product_id=x.id, snapshot_at=at(back(40)), price=D(100)))
        db.add(ProductPrice(product_id=x.id, snapshot_at=at(back(10), 13), price=D(120)))
        db.add(ProductPrice(product_id=x.id, snapshot_at=at(back(5)), price=D(115)))
        db.add(PriceChange(user_id=uid, product_id=x.id, offer_id="X-1", created_at=at(back(10)), change_type=MANUAL_SINGLE,
                           price_before=D(100), price_after=D(120), status=APPLIED, is_applied=True))
        promo = Promotion(user_id=uid, ozon_action_id=5, title="Распродажа", synced_at=now)
        db.add(promo)
        await db.flush()
        db.add(PromotionProduct(promotion_id=promo.id, product_id=y.id, action_price=D(250), is_active=True,
                                added_automatically=True, first_seen_at=at(back(8)), left_at=at(back(3)), synced_at=now))
        for when, present in ((back(45), 5), (back(6), 0), (back(2), 10)):
            db.add(ProductStock(product_id=y.id, warehouse_id=w.id, snapshot_at=at(when), present=present))
        await db.commit()
        return {"x": x.id, "y": y.id, "promo": promo.id}


async def test_api_sales_table_events_and_export(client, tokens, sessions):
    uid = await user_id_of(client, tokens)
    ids = await seed(sessions, uid)
    today = today_msk()
    back = lambda n: str(today - timedelta(days=n))  # noqa: E731
    period = {"date_from": back(30), "date_to": back(1)}
    get = lambda path, **kw: client.get(f"/api/v1/analytics/{path}", params={**period, **kw}, headers=auth(tokens))  # noqa: E731

    table = (await get("sales")).json()
    assert (table["days"], table["prev_days"], table["orders_days"], table["compared"]) == (30, 30, 30, True)
    rows = {r["product"]["offer_id"]: r for r in table["items"]}
    x, y = rows["X-1"], rows["Y-2"]
    assert (x["orders"], x["ordered_qty"], x["sold_qty"], x["revenue"], x["avg_price"]) == (10, 20, 60, "6000.00", "100.00")
    assert (x["prev_sold_qty"], x["change_pct"], x["direction"]) == (60, "0.00", None)
    assert (y["orders"], y["sold_qty"], y["returned_qty"], y["return_share"], y["returns_amount"]) == (0, 30, 1, "3.30", "300.00")
    assert (y["prev_sold_qty"], y["expected"], y["change_pct"], y["direction"]) == (120, "120.00", "-75.00", "down")
    assert [r["product"]["offer_id"] for r in table["items"]] == ["Y-2", "X-1"]  # по выручке: 9000 и 6000
    assert table["totals"] == {"orders": 10, "ordered_qty": 20, "sold_qty": 90, "returned_qty": 1, "revenue": "15000.00",
                               "return_share": "1.10"}

    events = (await get("sales/events")).json()
    assert [(e["day"], e["kind"]) for e in events] == [
        (back(10), "price"), (back(8), "promo_in"), (back(6), "stockout"), (back(5), "price"), (back(3), "promo_out"),
        (back(2), "restock")]
    assert events[0]["title"] == "Цена X-1 ↑ 100 ₽ → 120 ₽, через сервис"
    assert events[3]["title"] == "Цена X-1 ↓ 120 ₽ → 115 ₽"  # изменили в кабинете Ozon
    assert events[1]["title"] == "Y-2 вошёл в акцию «Распродажа» — Ozon добавил сам"
    assert events[1]["link"] == f"/dashboard/products/promotions/{ids['promo']}"
    only_y = (await get("sales/events", product_id=str(ids["y"]))).json()
    assert {e["kind"] for e in only_y} == {"promo_in", "stockout", "promo_out", "restock"}

    one = (await get("sales", product_id=str(ids["x"]))).json()
    assert [r["product"]["offer_id"] for r in one["items"]] == ["X-1"]

    r = await get("sales/export")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(1, 1).value == "Артикул" and [ws.cell(i, 1).value for i in (2, 3)] == ["Y-2", "X-1"]
    assert ws.cell(2, 7).value == 3.3 and ws.cell(5, 1).value == "Итого" and ws.cell(5, 5).value == 90

    assert (await get("sales", product_id=str(uuid.uuid4()))).status_code == 404
    bad = {"date_from": back(1), "date_to": back(30)}
    assert (await client.get("/api/v1/analytics/sales", params=bad, headers=auth(tokens))).status_code == 400
