"""Бюджеты и сценарии (шаг 10.21): план-факт по месяцам, кварталам и году (суммы месяцев), подстатьи (план и факт
входят в статью), отклонения и «хуже плана», версии плана, заполнение «как в прошлом году / как прошлый период»,
шаблон Excel туда-обратно, сценарии и риск кассового разрыва."""
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.api.v1 import budgets as budgets_api
from app.engine import budget as eng
from app.services import budgets
from tests.test_plan_calendar import TODAY, call, fixed_today  # noqa: F401 — фикстура «сегодня»
from tests.test_report_pnl import auth

D = Decimal
URL = "/api/v1/planning"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def budgets_today(monkeypatch):
    for module in (budgets, budgets_api):
        monkeypatch.setattr(module, "today_msk", lambda: TODAY)


def test_engine_periods_cells_split_scenarios():
    q = eng.periods(2026, "quarter")
    assert [(p.key, p.start, p.end) for p in q][3] == ("2026-Q4", date(2026, 10, 1), date(2026, 12, 31))
    assert [p.label for p in eng.periods(2026, "month")][:2] == ["янв 2026", "фев 2026"]
    assert eng.periods(2026, "year")[0].months[-1] == date(2026, 12, 1)
    # хуже плана: поступлений меньше, расходов больше
    assert eng.cell(D(100000), D(90000), "income") == eng.Cell(D(100000), D(90000), D("-10000.00"), D("-10.0"), True)
    assert eng.cell(D(30000), D(32000), "expense").worse is True
    assert eng.cell(D(15000), D(11000), "expense").worse is False
    assert eng.cell(None, D(500), "expense").diff is None  # плана нет — отклонения нет
    assert eng.split(D(100000), 3) == [D("33333.33"), D("33333.33"), D("33333.34")]
    assert eng.group(None, [None, None]) is None and eng.group(D(5000), [D(10000), None]) == D(15000)
    opt, pes = eng.scenarios(D(10), D(-5), D(-10), D(10))[1:]
    assert (opt.apply(D(1000)), opt.apply(D(-1000)), pes.apply(D(1000)), pes.apply(D(-1000))) == (
        D("1100.00"), D("-950.00"), D("900.00"), D("-1100.00"))
    assert eng.remainder(D(30000), D(5000), D(20000)) == D(5000) and eng.remainder(None, None, D(0)) == 0


async def setup(client, tokens, balance="100000"):
    await call(client, tokens, "POST", "/setup", json={
        "supplier_payment": "upfront", "accounts": [{"name": "Расчётный счёт", "kind": "checking", "balance": balance}],
        "ozon_account": 0, "categories": ["rent", "salary"]})
    await call(client, tokens, "PUT", "/settings", json={"gen_ozon": False, "gen_taxes": False})
    cats = {c["code"] or c["name"]: c["id"] for c in await call(client, tokens, "GET", "/categories")}
    mk = await call(client, tokens, "POST", "/categories", 201, json={"name": "Маркетинг", "kind": "expense"})
    ads = await call(client, tokens, "POST", "/categories", 201, json={
        "name": "Реклама", "kind": "expense", "parent_id": mk["id"]})
    cats.update({"mk": mk["id"], "ads": ads["id"]})
    acc = (await call(client, tokens, "GET", "/accounts"))[0]["id"]
    return cats, acc


async def plan(client, tokens, category_id, start, amount, granularity="month", expect=200):
    return await call(client, tokens, "PUT", "/budgets/plan", expect, json={
        "category_id": category_id, "start": start, "amount": amount, "granularity": granularity})


async def op(client, tokens, acc, kind, amount, day, category=None):
    body = {"type": kind, "amount": amount, "day": day, "account_id": acc}
    if category:
        body["category_id"] = category
    return await call(client, tokens, "POST", "/operations", 201, json=body)


def row(section, category_id):
    return next(r for r in section["rows"] if r["category"]["id"] == category_id)


async def test_plan_fact_subcategories_quarter_year(client, tokens):
    cats, acc = await setup(client, tokens)
    assert (await plan(client, tokens, cats["rent"], "2026-10-01", "30000"))["changed"] == 1
    assert (await plan(client, tokens, cats["salary"], "2026-11-15", "90000", "quarter"))["changed"] == 3
    await plan(client, tokens, cats["mk"], "2026-10-01", "5000")
    await plan(client, tokens, cats["ads"], "2026-10-01", "10000")
    assert (await plan(client, tokens, cats["marketplace"], "2026-01-01", "1200000", "year"))["changed"] == 12
    await op(client, tokens, acc, "expense", "32000", "2026-10-05", cats["rent"])
    await op(client, tokens, acc, "expense", "8000", "2026-10-06", cats["ads"])
    await op(client, tokens, acc, "expense", "3000", "2026-10-07", cats["mk"])
    await op(client, tokens, acc, "income", "90000", "2026-10-02", cats["marketplace"])
    await op(client, tokens, acc, "expense", "500", "2026-10-08")  # без статьи — «вне бюджета»

    body = await call(client, tokens, "GET", "/budgets", params={"year": 2026, "granularity": "month"})
    assert [p["key"] for p in body["periods"]][9] == "2026-10" and body["periods"][9]["partial"] is True
    oct_ = 9
    rent = row(body["expense"], cats["rent"])["cells"][oct_]
    assert rent == {"plan": "30000.00", "fact": "32000.00", "diff": "2000.00", "pct": "6.7", "worse": True}
    mk = row(body["expense"], cats["mk"])
    assert mk["cells"][oct_] == {"plan": "15000.00", "fact": "11000.00", "diff": "-4000.00", "pct": "-26.7",
                                 "worse": False}  # подстатья входит в статью: план 5 000 + 10 000, факт 3 000 + 8 000
    assert mk["children"][0]["cells"][oct_]["plan"] == "10000.00" and mk["own"][oct_]["plan"] == "5000.00"
    market = row(body["income"], cats["marketplace"])["cells"][oct_]
    assert (market["fact"], market["diff"], market["pct"], market["worse"]) == ("90000.00", "-10000.00", "-10.0", True)
    assert row(body["expense"], cats["rent"])["cells"][10]["fact"] is None  # ноябрь ещё не начался
    assert row(body["expense"], cats["rent"])["cells"][8] == {"plan": None, "fact": "0.00", "diff": None, "pct": None,
                                                               "worse": None}  # сентябрь: плана нет, трат нет
    assert body["expense"]["other"]["cells"][oct_]["fact"] == "500.00"
    assert (body["expense"]["cells"][oct_]["plan"], body["expense"]["cells"][oct_]["fact"]) == ("75000.00", "43500.00")
    assert body["net"]["cells"][oct_]["plan"] == "25000.00" and body["net"]["cells"][oct_]["fact"] == "46500.00"

    # квартал и год — суммы месяцев
    body = await call(client, tokens, "GET", "/budgets", params={"year": 2026, "granularity": "quarter"})
    assert row(body["expense"], cats["salary"])["cells"][3]["plan"] == "90000.00"
    assert row(body["expense"], cats["rent"])["cells"][3]["plan"] == "30000.00"
    assert row(body["income"], cats["marketplace"])["cells"][0]["plan"] == "300000.00"
    body = await call(client, tokens, "GET", "/budgets", params={"year": 2026, "granularity": "year"})
    assert row(body["income"], cats["marketplace"])["cells"][0]["plan"] == "1200000.00"
    assert row(body["income"], cats["marketplace"])["total"]["fact"] == "90000.00"

    # правка — новая версия, прежняя в истории; пусто — план убран
    await plan(client, tokens, cats["rent"], "2026-10-20", "31000")
    await plan(client, tokens, cats["rent"], "2026-10-01", None)
    history = await call(client, tokens, "GET", "/budgets/history", params={"category_id": cats["rent"], "month": "2026-10-01"})
    assert [(h["amount"], h["superseded_at"] is None) for h in history] == [
        (None, True), ("31000.00", False), ("30000.00", False)]
    body = await call(client, tokens, "GET", "/budgets", params={"year": 2026})
    assert row(body["expense"], cats["rent"])["cells"][oct_]["plan"] is None

    # переводы и статьи вне бюджета не планируются
    transfer = next(c for c in await call(client, tokens, "GET", "/categories") if c["kind"] == "transfer")
    r = await client.put(f"{URL}/budgets/plan", headers=auth(tokens), json={
        "category_id": transfer["id"], "start": "2026-10-01", "amount": "1"})
    assert r.json()["detail"].startswith("Бюджет — по статьям поступлений и расходов")
    await call(client, tokens, "PATCH", f"/categories/{cats['salary']}", json={"in_budget": False})
    r = await client.put(f"{URL}/budgets/plan", headers=auth(tokens), json={
        "category_id": cats["salary"], "start": "2026-10-01", "amount": "1"})
    assert r.json()["detail"] == "Статья «Зарплата» не участвует в бюджетах — отметьте «Бюджеты» в «Статьях»"
    assert (await client.put(f"{URL}/budgets/plan", headers=auth(tokens), json={
        "category_id": cats["rent"], "start": "2026-10-01", "amount": "-5"})).status_code == 400


async def test_fill_from_last_year_and_previous_period(client, tokens):
    cats, acc = await setup(client, tokens)
    await op(client, tokens, acc, "expense", "29000", "2025-11-05", cats["rent"])
    await op(client, tokens, acc, "expense", "28000", "2025-11-25", cats["salary"])
    await op(client, tokens, acc, "expense", "32000", "2026-10-05", cats["rent"])
    await plan(client, tokens, cats["salary"], "2026-11-01", "30000")

    dry = await call(client, tokens, "POST", "/budgets/fill", json={
        "source": "last_year", "granularity": "month", "start": "2026-11-01"})
    assert (dry["changed"], dry["kept"], dry["from"]) == (1, 1, "2025-11-01")  # зарплата — уже есть свой план
    assert dry["examples"] == [{"category": "Аренда", "month": "2026-11-01", "before": None, "after": "29000.00"}]
    body = await call(client, tokens, "GET", "/budgets", params={"year": 2026})
    assert row(body["expense"], cats["rent"])["cells"][10]["plan"] is None  # проверка ничего не пишет
    done = await call(client, tokens, "POST", "/budgets/fill", json={
        "source": "last_year", "granularity": "month", "start": "2026-11-01", "overwrite": True, "dry_run": False})
    assert done["changed"] == 2
    body = await call(client, tokens, "GET", "/budgets", params={"year": 2026})
    assert [row(body["expense"], cats[c])["cells"][10]["plan"] for c in ("rent", "salary")] == ["29000.00", "28000.00"]
    # «как прошлый период»: декабрь ← ноябрь (факта ещё нет — не заполняем), ноябрь ← октябрь
    prev = await call(client, tokens, "POST", "/budgets/fill", json={
        "source": "prev_period", "granularity": "month", "start": "2026-12-01"})
    assert prev["changed"] == 0
    prev = await call(client, tokens, "POST", "/budgets/fill", json={
        "source": "prev_period", "granularity": "month", "start": "2026-11-01", "overwrite": True})
    assert prev["examples"] == [{"category": "Аренда", "month": "2026-11-01", "before": "29000.00", "after": "32000.00"}]
    history = await call(client, tokens, "GET", "/budgets/history", params={"category_id": cats["rent"], "month": "2026-11-01"})
    assert history[0]["source"] == "last_year"


async def test_template_round_trip(client, tokens):
    cats, acc = await setup(client, tokens)
    await plan(client, tokens, cats["rent"], "2026-01-01", "360000", "year")
    r = await client.get(f"{URL}/budgets/template", params={"year": 2026}, headers=auth(tokens))
    assert r.status_code == 200 and r.content[:2] == b"PK"
    files = {"file": ("budget.xlsx", r.content, XLSX)}
    same = (await client.post(f"{URL}/budgets/import", params={"dry_run": True}, headers=auth(tokens), files=files)).json()
    assert same["counts"]["error"] == 0 and same["counts"]["unchanged"] == 1  # скачанный файл принимается обратно
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Бюджет"]
    header = [c.value for c in ws[1]]
    rent_row = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, 1).value == "Аренда")
    ws.cell(rent_row, header.index("Декабрь") + 1).value = 35000
    ws.cell(rent_row, header.index("Ноябрь") + 1).value = None  # пусто — плана нет
    ads_row = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, 1).value == "Реклама")
    ws.cell(ads_row, header.index("Январь") + 1).value = 5000
    ws.append(["Хостинг", "Расход", None, 2026, 990])
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("budget.xlsx", buf.getvalue(), XLSX)}
    preview = (await client.post(f"{URL}/budgets/import", params={"dry_run": True}, headers=auth(tokens), files=files)).json()
    statuses = {r["category"] or r["line"]: (r["status"], r["comment"], r["message"]) for r in preview["rows"]}
    assert statuses["Аренда"] == ("changed", "изменится месяцев: 2", None)
    assert statuses["Реклама"][0] == "new"
    error = next(r for r in preview["rows"] if r["status"] == "error")
    assert error["message"].startswith("Статьи «Хостинг» нет среди статей бюджета")
    applied = (await client.post(f"{URL}/budgets/import", params={"dry_run": False}, headers=auth(tokens),
                                 files=files)).json()
    assert applied["applied"] is True
    body = await call(client, tokens, "GET", "/budgets", params={"year": 2026})
    rent = row(body["expense"], cats["rent"])
    assert [c["plan"] for c in rent["cells"][10:]] == [None, "35000.00"]
    assert row(body["expense"], cats["mk"])["cells"][0]["plan"] == "5000.00"  # подстатья — в статье
    history = await call(client, tokens, "GET", "/budgets/history", params={"category_id": cats["rent"], "month": "2026-12-01"})
    assert [h["source"] for h in history] == ["import", "manual"]
    errors = await client.post(f"{URL}/budgets/import/errors", headers=auth(tokens), files=files)
    assert errors.status_code == 200 and errors.content[:2] == b"PK"


async def test_scenarios_forecast_gap(client, tokens):
    cats, acc = await setup(client, tokens, balance="50000")
    for month in ("2026-11-01", "2026-12-01"):
        await plan(client, tokens, cats["rent"], month, "30000")
        await plan(client, tokens, cats["marketplace"], month, "15000")
    await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "expense", "title": "Аренда", "amount": "20000", "start_date": "2026-11-05", "category_id": cats["rent"]})
    body = await call(client, tokens, "GET", "/budgets/forecast", params={"months": 3})
    assert (body["start"], body["date_to"]) == ("50000.00", "2026-12-31")
    # аренду ноября частично уже видно в календаре: из бюджета добавляем только остаток
    assert [(a["category"], a["day"], a["amount"]) for a in body["budget_added"]] == [
        ("Аренда", "2026-11-01", "-10000.00"), ("Поступления от маркетплейсов", "2026-11-30", "15000.00"),
        ("Аренда", "2026-12-01", "-30000.00"), ("Поступления от маркетплейсов", "2026-12-31", "15000.00")]
    sc = {s["key"]: s for s in body["scenarios"]}
    assert (sc["base"]["end_balance"], sc["base"]["min_balance"], sc["base"]["min_day"], sc["base"]["first_below_zero"]) == (
        "20000.00", "5000.00", "2026-12-01", None)
    assert (sc["pessimistic"]["first_below_zero"], sc["pessimistic"]["min_balance"]) == ("2026-12-01", "-2500.00")
    assert sc["optimistic"]["end_balance"] == "26000.00"
    await call(client, tokens, "PUT", "/budgets/scenarios", 204, json={"scen_pes_expense": "0", "scen_pes_income": "0"})
    sc = {s["key"]: s for s in (await call(client, tokens, "GET", "/budgets/forecast"))["scenarios"]}
    assert sc["pessimistic"]["end_balance"] == sc["base"]["end_balance"]
    assert (await client.put(f"{URL}/budgets/scenarios", headers=auth(tokens),
                             json={"scen_opt_income": "500"})).status_code == 400
