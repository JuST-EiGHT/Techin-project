"""«Операции» (шаг 10.19): подстатьи (одна вложенность, проверки, факт подстатьи входит в статью — ДДС и фильтр),
«где учитывается статья», единый журнал факт + план (статусы, источники, фильтры, сортировка, «Оплачено», Excel),
массовое изменение (новые версии, всё или ничего)."""
import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.engine.cash_flow import Flow, cash_flow
from tests.test_plan_calendar import TODAY, call, fixed_today, setup  # noqa: F401 — фикстура «сегодня»
from tests.test_planning_overview import overview_today  # noqa: F401
from tests.test_report_pnl import auth

D = Decimal
URL = "/api/v1/planning"


async def op(client, tokens, kind, amount, day, account, category=None, **extra):
    body = {"type": kind, "amount": amount, "day": day, "account_id": account, **extra}
    if category:
        body["category_id"] = category
    return await call(client, tokens, "POST", "/operations", 201, json=body)


async def err(client, tokens, method, path, **kw) -> str:
    r = await client.request(method, f"{URL}{path}", headers=auth(tokens), **kw)
    assert r.status_code == 400, r.text
    return r.json()["detail"]


def test_cash_flow_subcategories_roll_into_parent():
    flows = [
        Flow(date(2026, 9, 5), D(-10000), "ads", "expense"),  # статья напрямую
        Flow(date(2026, 9, 7), D(-3000), "yandex", "expense"),  # подстатьи
        Flow(date(2026, 10, 2), D(-2000), "vk", "expense"),
        Flow(date(2026, 10, 3), D(500), "vk", "expense"),  # возврат по подстатье — уменьшает расход статьи
        Flow(date(2026, 10, 4), D(-12000), "rent", "expense"),
        Flow(date(2026, 10, 5), D(-700), "print", "expense"),  # подстатья статьи без своих операций
    ]
    parents = {"yandex": "ads", "vk": "ads", "print": "office"}
    flat = cash_flow(D(50000), flows, date(2026, 9, 1), date(2026, 10, 14))
    r = cash_flow(D(50000), flows, date(2026, 9, 1), date(2026, 10, 14), parents=parents)
    assert [(row.category, row.values, row.total, [c.category for c in row.children]) for row in r.rows] == [
        ("ads", [D("-13000.00"), D("-1500.00")], D("-14500.00"), ["yandex", "vk"]),
        ("rent", [D("0.00"), D("-12000.00")], D("-12000.00"), []),
        ("office", [D("0.00"), D("-700.00")], D("-700.00"), ["print"]),
    ]
    ads = r.rows[0]
    assert [(c.category, c.total) for c in ads.children] == [("yandex", D("-3000.00")), ("vk", D("-1500.00"))]
    # итоги разделов и остатки — те же: каждая операция считается один раз
    assert (r.expense, r.closing) == (flat.expense, flat.closing) == (D("-27200.00"), D("22800.00"))
    assert r.opening + sum(row.total for row in r.rows) == r.closing


async def test_subcategories_rules_and_where_counted(client, tokens):
    cats, accounts = await setup(client, tokens)
    mk = await call(client, tokens, "POST", "/categories", 201, json={"name": "Маркетинг", "kind": "expense"})
    assert (mk["parent_id"], mk["in_pnl"], mk["in_budget"], mk["has_children"]) == (None, True, True, False)
    ya = await call(client, tokens, "POST", "/categories", 201, json={
        "name": "Реклама Яндекс", "kind": "expense", "parent_id": mk["id"]})
    assert ya["parent_id"] == mk["id"] and ya["in_pnl"] is True  # подстатья — как статья
    listed = {c["name"]: c for c in await call(client, tokens, "GET", "/categories")}
    assert listed["Маркетинг"]["has_children"] is True
    # где учитывается по умолчанию: налоги и поставщик — не расходы P&L; переводы — не в бюджетах
    assert (listed["Налоги и взносы"]["in_pnl"], listed["Аренда"]["in_pnl"], listed["Поступления от маркетплейсов"]["in_pnl"]) == (
        False, True, False)

    errors = {
        "«Реклама Яндекс» сама подстатья — вложенность одна: выберите статью верхнего уровня":
            ("POST", "/categories", {"name": "Директ", "kind": "expense", "parent_id": ya["id"]}),
        "Подстатья — того же вида, что статья: поступление в поступлении, расход в расходе":
            ("POST", "/categories", {"name": "Возвраты", "kind": "income", "parent_id": mk["id"]}),
        "В P&L идут только расходы: доходы магазина P&L берёт из начислений Ozon":
            ("POST", "/categories", {"name": "Прочие доходы", "kind": "income", "in_pnl": True}),
        "У статьи есть подстатьи — сначала перенесите их, потом делайте её подстатьёй":
            ("PATCH", f"/categories/{mk['id']}", {"parent_id": cats["rent"]}),
        "Статья не может быть подстатьёй самой себя": ("PATCH", f"/categories/{cats['rent']}", {"parent_id": cats["rent"]}),
        "Эта статья нужна сервису — она остаётся статьёй верхнего уровня":
            ("PATCH", f"/categories/{cats['marketplace']}", {"parent_id": cats["taxes"]}),
        "У статьи есть подстатьи — сначала перенесите их в другую статью или в архив":
            ("PATCH", f"/categories/{mk['id']}", {"archived": True}),
    }
    for message, (method, path, body) in errors.items():
        assert await err(client, tokens, method, path, json=body) == message

    # перенести существующую статью в подстатьи и обратно; отметки «где учитывается» меняются
    moved = await call(client, tokens, "PATCH", f"/categories/{cats['salary']}", json={"parent_id": mk["id"]})
    assert moved["parent_id"] == mk["id"]
    top = await call(client, tokens, "PATCH", f"/categories/{cats['salary']}", json={"parent_id": None, "in_pnl": False,
                                                                                      "in_budget": False})
    assert (top["parent_id"], top["in_pnl"], top["in_budget"]) == (None, False, False)
    untouched = await call(client, tokens, "PATCH", f"/categories/{cats['salary']}", json={"name": "Зарплата"})
    assert untouched["parent_id"] is None  # parent_id не передан — не меняется

    # архив: подстатья — можно; потом статья; подстатья из архива возвращается на верхний уровень
    await call(client, tokens, "PATCH", f"/categories/{ya['id']}", json={"archived": True})
    await call(client, tokens, "PATCH", f"/categories/{mk['id']}", json={"archived": True})
    assert await err(client, tokens, "POST", "/categories", json={
        "name": "SMM", "kind": "expense", "parent_id": mk["id"]}) == "Статья «Маркетинг» в архиве — выберите действующую"
    back = await call(client, tokens, "PATCH", f"/categories/{ya['id']}", json={"archived": False})
    assert (back["archived"], back["parent_id"]) == (False, None)


async def test_subcategory_fact_counts_in_category(client, tokens):
    cats, accounts = await setup(client, tokens)
    acc = accounts["Расчётный счёт"]
    mk = await call(client, tokens, "POST", "/categories", 201, json={"name": "Маркетинг", "kind": "expense"})
    ya = await call(client, tokens, "POST", "/categories", 201, json={
        "name": "Реклама Яндекс", "kind": "expense", "parent_id": mk["id"]})
    vk = await call(client, tokens, "POST", "/categories", 201, json={
        "name": "Реклама VK", "kind": "expense", "parent_id": mk["id"]})
    await op(client, tokens, "expense", "4000", "2026-09-10", acc, mk["id"])
    await op(client, tokens, "expense", "3000", "2026-09-12", acc, ya["id"])
    await op(client, tokens, "expense", "2500", "2026-10-03", acc, vk["id"])
    await op(client, tokens, "expense", "30000", "2026-10-05", acc, cats["rent"])

    body = await call(client, tokens, "GET", "/cash-flow", params={"date_from": "2026-09-01", "date_to": "2026-10-14"})
    rows = {r["name"]: r for r in body["rows"]}
    assert (rows["Маркетинг"]["values"], rows["Маркетинг"]["total"]) == (["-7000.00", "-2500.00"], "-9500.00")
    assert [(c["name"], c["total"]) for c in rows["Маркетинг"]["children"]] == [
        ("Реклама Яндекс", "-3000.00"), ("Реклама VK", "-2500.00")]
    assert "Реклама Яндекс" not in rows and body["expense"] == "-39500.00"  # подстатьи — внутри статьи, не отдельно

    # фильтр по статье — вместе с подстатьями
    page = await call(client, tokens, "GET", "/operations", params={"date_from": "2026-09-01", "category_id": mk["id"]})
    assert sorted(o["amount"] for o in page["items"]) == ["-2500.00", "-3000.00", "-4000.00"]
    assert {o["category"]["parent_id"] for o in page["items"]} == {None, mk["id"]}
    page = await call(client, tokens, "GET", "/operations", params={"date_from": "2026-09-01", "category_id": ya["id"]})
    assert [o["amount"] for o in page["items"]] == ["-3000.00"]
    counts = {c["name"]: c["operations"] for c in await call(client, tokens, "GET", "/categories")}
    assert (counts["Маркетинг"], counts["Реклама Яндекс"], counts["Аренда"]) == (1, 1, 1)


async def test_journal_fact_and_plan_filters_paid_export(client, tokens):
    cats, accounts = await setup(client, tokens)
    acc = accounts["Расчётный счёт"]
    await call(client, tokens, "PUT", "/settings", json={"gen_ozon": False, "gen_taxes": False})
    rent = await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "expense", "title": "Аренда склада", "amount": "30000", "start_date": "2026-10-10", "period": "month",
        "category_id": cats["rent"], "account_id": acc, "counterparty": "ИП Петров"})
    await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "income", "title": "Возврат депозита", "amount": "5000", "start_date": "2026-10-20"})
    await op(client, tokens, "expense", "1200", "2026-10-12", acc, cats["salary"], counterparty="Courier LLC")
    await op(client, tokens, "income", "800", "2026-10-13", acc, comment="Возврат от поставщика")
    params = {"date_from": "2026-10-01", "date_to": "2026-10-31"}

    body = await call(client, tokens, "GET", "/journal", params=params)
    rows = [(r["status"], r["type"], r["day"], r["amount"], r["source"]) for r in body["items"]]
    assert rows == [
        ("planned", "income", "2026-10-20", "5000.00", "plan"),
        ("fact", "income", "2026-10-13", "800.00", "manual"),
        ("fact", "expense", "2026-10-12", "-1200.00", "manual"),
        ("overdue", "expense", "2026-10-10", "-30000.00", "plan"),  # срок прошёл, оплаты нет
    ]
    assert body["counts"] == {"fact": 2, "planned": 1, "expected": 0, "overdue": 1}
    assert body["totals"]["fact"] == {"income": "800.00", "expense": "-1200.00", "net": "-400.00"}
    assert body["totals"]["plan"] == {"income": "5000.00", "expense": "-30000.00", "net": "-25000.00"}
    overdue = body["items"][3]
    assert (overdue["category"]["name"], overdue["account"]["name"], overdue["counterparty"], overdue["comment"]) == (
        "Аренда", "Расчётный счёт", "ИП Петров", "Аренда склада")
    assert overdue["plan"]["key"] == f"plan:{rent['plan_id']}:2026-10-10" and overdue["operation"] is None

    async def journal(**extra):
        return [(r["status"], r["amount"]) for r in (await call(client, tokens, "GET", "/journal",
                                                                params={**params, **extra}))["items"]]

    assert await journal(status="fact") == [("fact", "800.00"), ("fact", "-1200.00")]
    assert await journal(type="expense", sort="amount", order="asc") == [("fact", "-1200.00"), ("overdue", "-30000.00")]
    assert await journal(category_id=cats["rent"]) == [("overdue", "-30000.00")]
    assert await journal(uncategorized="true") == [("planned", "5000.00"), ("fact", "800.00")]
    assert await journal(source="plan") == [("planned", "5000.00"), ("overdue", "-30000.00")]
    assert await journal(search="петров") == [("overdue", "-30000.00")]
    assert await journal(search="courier") == [("fact", "-1200.00")]

    # плановая → фактическая: «Оплачено» с датой и суммой — строка становится фактом и помнит платёж
    await call(client, tokens, "POST", "/calendar/paid", json={"key": overdue["plan"]["key"], "day": "2026-10-11",
                                                               "amount": "29500"})
    body = await call(client, tokens, "GET", "/journal", params=params)
    paid = next(r for r in body["items"] if r["plan"] and r["plan"]["title"] == "Аренда склада")
    assert (paid["status"], paid["day"], paid["amount"], paid["category"]["name"]) == (
        "fact", "2026-10-11", "-29500.00", "Аренда")
    assert paid["operation"]["source"] == "manual" and body["counts"] == {"fact": 3, "planned": 1, "expected": 0, "overdue": 0}

    r = await client.get(f"{URL}/journal/export", params=params, headers=auth(tokens))
    assert r.status_code == 200 and r.content[:2] == b"PK"
    ws = load_workbook(io.BytesIO(r.content)).active
    table = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=5)]
    assert table[0][:5] == ["Дата", "Статус", "Тип", "Сумма, ₽", "Статья"]
    assert [(row[1], row[3], row[4], row[9]) for row in table[1:]] == [
        ("План", 5000.0, "Без статьи", "Плановый платёж"), ("Факт", 800.0, "Без статьи", "Вручную"),
        ("Факт", -1200.0, "Зарплата", "Вручную"), ("Факт", -29500.0, "Аренда", "Вручную")]

    # без планирования плановых нет — и сказано почему
    await call(client, tokens, "PUT", "/settings", json={"enabled": False})
    off = await call(client, tokens, "GET", "/journal", params=params)
    assert off["counts"]["planned"] == 0 and off["plan_note"].startswith("Плановые платежи появятся")


async def test_bulk_change_versions_and_all_or_nothing(client, tokens):
    cats, accounts = await setup(client, tokens)
    acc = accounts["Расчётный счёт"]
    cash = (await call(client, tokens, "POST", "/accounts", 201, json={"name": "Наличные", "kind": "cash"}))["id"]
    a = await op(client, tokens, "expense", "100", "2026-10-01", acc)
    b = await op(client, tokens, "expense", "200", "2026-10-02", acc, cats["salary"])
    c = await op(client, tokens, "income", "300", "2026-10-03", acc)
    t = await op(client, tokens, "transfer", "50", "2026-10-04", acc, to_account_id=cash)
    ids = [a["op_id"], b["op_id"]]
    transfer_cat = next(c for c in await call(client, tokens, "GET", "/categories") if c["kind"] == "transfer")
    assert (transfer_cat["in_budget"], transfer_cat["in_pnl"]) == (False, False)  # перевод — не доход и не расход

    done = await call(client, tokens, "POST", "/operations/bulk", json={
        "op_ids": ids, "category_id": cats["rent"], "account_id": cash, "counterparty": "ИП Петров"})
    assert done == {"changed": 2, "unchanged": 0}
    again = await call(client, tokens, "POST", "/operations/bulk", json={"op_ids": ids, "counterparty": "ИП Петров"})
    assert again == {"changed": 0, "unchanged": 2}
    history = await call(client, tokens, "GET", f"/operations/{b['op_id']}/history")
    assert [(h["version"], h["category"], h["account"]) for h in history] == [
        (2, "Аренда", "Наличные"), (1, "Зарплата", "Расчётный счёт")]  # прежняя версия — в истории
    page = await call(client, tokens, "GET", "/operations", params={"date_from": "2026-10-01"})
    edited = {o["op_id"]: o for o in page["items"]}
    assert edited[a["op_id"]]["category_by"] == "user" and edited[a["op_id"]]["counterparty"] == "ИП Петров"

    # не подходит хоть одна — ничего не меняется, сказано какая
    assert await err(client, tokens, "POST", "/operations/bulk", json={
        "op_ids": [a["op_id"], c["op_id"]], "category_id": cats["salary"]}) == (
        "Среди выбранных есть поступления — для них нужна статья поступлений. Выберите только поступления или только платежи")
    assert await err(client, tokens, "POST", "/operations/bulk", json={
        "op_ids": [a["op_id"], t["op_id"]], "account_id": acc}) == (
        "Операция 04.10.2026 на −50 ₽: у перевода меняются сумма, дата и комментарий; другие счета — удалите перевод "
        "и добавьте заново")
    page = await call(client, tokens, "GET", "/operations", params={"date_from": "2026-10-01"})
    assert next(o for o in page["items"] if o["op_id"] == a["op_id"])["account"]["name"] == "Наличные"  # без изменений
    assert await err(client, tokens, "POST", "/operations/bulk", json={"op_ids": ids}) == "Выберите, что изменить"

    # убрать статью у выбранных
    cleared = await call(client, tokens, "POST", "/operations/bulk", json={"op_ids": ids, "category_id": None})
    assert cleared["changed"] == 2
    page = await call(client, tokens, "GET", "/operations", params={"date_from": "2026-10-01", "uncategorized": "true"})
    assert {o["op_id"] for o in page["items"]} >= set(ids)
