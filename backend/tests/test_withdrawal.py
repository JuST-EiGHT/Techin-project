"""Кандидаты на вывоз (шаг 5.1, BUSINESS_LOGIC §6): расчёт, API, параметры пользователя, Excel.
10.17: причина и приоритет, замороженные деньги, история продаж, фильтры, решения «вывозим» / «оставить», вывозы Ozon."""
import io
import uuid
from datetime import datetime, timedelta
from decimal import Decimal

import httpx
import pytest
import respx
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.config import settings
from app.engine import withdrawal as engine
from app.engine.stocks import Thresholds
from app.engine.withdrawal import IDLE, SELLING, Item, Params, Place, plan, surplus
from app.models.withdrawal import WithdrawalDecision
from app.services import withdrawal as withdrawal_svc
from app.services.sales import sold_by_product
from app.services.sync.finance import today_msk
from tests.test_pnl_api import add_api_day
from tests.test_sync_stocks import api_user, auth, ozon  # noqa: F401 — фикстура ozon
from tests.test_warehouses import ACB, CM, RS, TH, fake_sales, product_ids

SPB, MSK, EKB = "1", "2", "3"


def place(wid: str, cluster: str | None, present: int) -> Place:
    names = {SPB: "Санкт-Петербург и СЗО", MSK: "Москва, МО и Дальние регионы", EKB: "Урал"}
    return Place(wid, f"Склад {wid}", cluster, names.get(cluster), present)


def item(offer: str, places: list[Place], avg: str | None, dws: int | None, extra: int = 0) -> Item:
    """extra — остаток на своих складах (FBS / rFBS): входит в «доступно», но вывозить его неоткуда."""
    available = sum(p.present for p in places) + extra
    return Item(offer, offer, available, None if avg is None else Decimal(avg), dws, places)


def lines(c) -> list[tuple[str, int, int]]:
    return [(line.place.warehouse_id, line.qty, line.cluster_left) for line in c.lines]


# ── расчёт ──
def test_idle_product_keeps_two_per_cluster():
    """Не продаётся: в каждом кластере оставляем 2 шт., остальное — на вывоз; сначала склад с большим остатком."""
    it = item("A", [place("w1", MSK, 10), place("w2", MSK, 3), place("w3", EKB, 2), place("w4", SPB, 1)], "0", 45)
    p = plan([it], Params())
    (c,) = p.candidates
    assert (c.kind, c.keep, c.surplus, c.possible, c.qty) == (IDLE, 2, None, 11, 11)
    assert lines(c) == [("w1", 10, 2), ("w2", 1, 2)]  # Урал (2) и СПб (1) — не больше, чем оставляем
    assert (p.possible, p.planned, p.left_out, p.no_sales_data) == (11, 11, 0, False)


def test_selling_product_withdraws_only_surplus():
    """Продаётся, но запаса больше порога «Избыток»: в кластере оставляем 10+, всего — не больше излишка (остаток сверх
    запаса на 60 дней), чтобы товар не кончился."""
    it = item("B", [place("w1", MSK, 70), place("w2", EKB, 30)], "1", 2)  # 100 шт. при 1 шт./день — 100 дней
    assert surplus(it) == 40
    (c,) = plan([it], Params()).candidates
    assert (c.kind, c.keep, c.surplus, c.possible) == (SELLING, 10, 40, 40)
    assert lines(c) == [("w1", 40, 30)]  # излишек забрали из кластера с большим остатком, Урал не трогаем
    assert c.days_of_supply == Decimal("100.0")


def test_own_warehouses_count_but_are_not_withdrawn():
    """Остаток на своих складах (FBS) входит в запас товара, но вывозится только со складов Ozon."""
    it = item("C", [place("w1", MSK, 50)], "1", 0, extra=40)  # 90 шт. — 90 дней, излишек 30
    (c,) = plan([it], Params()).candidates
    assert (c.surplus, lines(c)) == (30, [("w1", 30, 20)])


def test_not_candidates():
    normal = item("N", [place("w1", MSK, 100)], "2", 0)  # 50 дней — норма
    unknown = item("U", [place("w1", MSK, 100)], None, None)  # продажи неизвестны
    little = item("L", [place("w1", MSK, 2), place("w2", EKB, 1)], "0", 90)  # в кластерах не больше, чем оставляем
    p = plan([normal, unknown, little], Params())
    assert (p.candidates, p.possible, p.no_sales_data) == ([], 0, False)
    assert plan([unknown], Params()).no_sales_data is True


def test_thresholds_and_keep_are_parameters():
    it = item("B", [place("w1", MSK, 70), place("w2", EKB, 30)], "1", 2)
    assert plan([it], Params(), Thresholds(7, 14, 120)).candidates == []  # 100 дней при пороге 120 — не избыток
    (c,) = plan([it], Params(keep_selling=60), Thresholds(7, 14, 30)).candidates
    assert (c.surplus, lines(c)) == (70, [("w1", 10, 60)])  # излишек 70, но в кластере оставляем 60


def test_excluded_cluster_is_not_touched():
    it = item("A", [place("w1", SPB, 20), place("w2", MSK, 5)], "0", 30)
    (c,) = plan([it], Params(excluded=frozenset({SPB}))).candidates
    assert (lines(c), c.excluded_units) == ([("w2", 3, 2)], 20)
    only_spb = item("S", [place("w1", SPB, 20)], "0", 30)
    assert plan([only_spb], Params(excluded=frozenset({SPB}))).candidates == []


def test_worst_first():
    """Дольше не продававшиеся — первыми; продаж в данных нет — самые первые; при равных — больше дней запаса."""
    items = [
        item("recent", [place("w1", MSK, 10)], "0", 31),
        item("never", [place("w1", MSK, 10)], "0", None),
        item("selling", [place("w1", MSK, 300)], "1", 40),
        item("old", [place("w1", MSK, 10)], "0", 200),
        item("selling-more", [place("w1", MSK, 600)], "1", 40),
    ]
    order = [c.item.offer_id for c in plan(items, Params()).candidates]
    assert order == ["never", "old", "selling-more", "selling", "recent"]


def test_total_limit():
    """«Вывезти всего N шт.»: по порядку «хуже — первыми», последняя строка — частично."""
    a = item("A", [place("w1", MSK, 10), place("w2", MSK, 3)], "0", 100)  # можно 11: 10 + 1
    b = item("B", [place("w1", MSK, 100)], "1", 2)  # 100 дней — можно 40
    p = plan([a, b], Params(total=15))
    assert (p.possible, p.planned, p.left_out) == (51, 15, 0)
    assert [lines(c) for c in p.candidates] == [[("w1", 10, 2), ("w2", 1, 2)], [("w1", 4, 96)]]
    p = plan([a, b], Params(total=5))
    assert (p.planned, p.left_out, [lines(c) for c in p.candidates]) == (5, 1, [[("w1", 5, 8)]])


# ── API ──
URL = "/api/v1/warehouses/withdrawal"
SPB_CLUSTER, EKB_CLUSTER = "2", "154"


@pytest.fixture
async def seller(sessions, client, tokens, ozon, monkeypatch):  # noqa: F811
    """TH: 31 шт. (Екатеринбург 9, СПб 10, свой склад 12) по 0,2 в день — 155 дней, излишек 31 − 12 = 19;
    ACB: СПб 5 шт., не продаётся 45 дней; RS — кончился; CM — 1 шт. только на своём складе, не продаётся."""
    user_id = await api_user(sessions, client, tokens, ozon)
    ids = await product_ids(sessions, user_id)
    fake_sales(monkeypatch, ids, {TH: "0.2", ACB: "0", RS: "1", CM: "0"}, {TH: 3, ACB: 45, RS: 2})
    return ids


def summary(data) -> list[tuple]:
    return [
        (it["product"]["offer_id"], it["kind"], it["qty"],
         [(line["cluster"], line["warehouse"], line["qty"], line["cluster_left"]) for line in it["lines"]])
        for it in data["items"]
    ]


async def test_report_defaults(client, tokens, seller):
    data = (await client.get(URL, headers=auth(tokens))).json()
    assert data["params"] == {"total": None, "keep_idle": 2, "keep_selling": 10, "excluded": []} == data["saved"]
    assert data["is_saved"] is True and data["sales_known"] is True
    # TH продаётся: в кластерах по 9 и 10 шт. при «оставлять 10» вывозить нечего
    assert summary(data) == [(ACB, "idle", 3, [("Санкт-Петербург и СЗО", "СПБ_БУГРЫ_РФЦ", 3, 2)])]
    (acb,) = data["items"]
    assert (acb["reason"], acb["keep"], acb["days_without_sales"], acb["possible"]) == ("Нет продаж 45 дн.", 2, 45, 3)
    assert (data["possible"], data["planned"], data["left_out"]) == (3, 3, 0)
    assert [(c["key"], c["name"], c["present"], c["products"]) for c in data["clusters"]] == [
        (SPB_CLUSTER, "Санкт-Петербург и СЗО", 15, 2), (EKB_CLUSTER, "Екатеринбург", 9, 1),
    ]  # Краснодар пуст — не показываем


async def test_report_overrides_and_saved_params(client, tokens, seller):
    get = lambda **q: client.get(URL, params=q, headers=auth(tokens))  # noqa: E731
    data = (await get(keep_selling=5)).json()
    assert data["is_saved"] is False and data["params"]["keep_selling"] == 5
    assert summary(data) == [
        (ACB, "idle", 3, [("Санкт-Петербург и СЗО", "СПБ_БУГРЫ_РФЦ", 3, 2)]),
        (TH, "selling", 9, [("Санкт-Петербург и СЗО", "СПБ_БУГРЫ_РФЦ", 5, 5),
                            ("Екатеринбург", "ЕКАТЕРИНБУРГ_РФЦ_НОВЫЙ", 4, 5)]),
    ]
    th = data["items"][1]
    assert (th["surplus"], th["days_of_supply"], th["reason"]) == (19, "155.0", "Запас на 155 дн. при продажах 0,2 шт./день")

    # сохранить: СПб не трогаем (решение владельца — его настройка), оставлять 5 у продающихся
    body = {"total": 0, "keep_idle": 2, "keep_selling": 5, "excluded": [SPB_CLUSTER]}
    r = await client.put(f"{URL}/params", json=body, headers=auth(tokens))
    assert r.json() == {"total": None, "keep_idle": 2, "keep_selling": 5, "excluded": [SPB_CLUSTER]}
    assert (await client.get(f"{URL}/params", headers=auth(tokens))).json() == r.json()
    data = (await get()).json()
    assert data["is_saved"] is True
    assert summary(data) == [(TH, "selling", 4, [("Екатеринбург", "ЕКАТЕРИНБУРГ_РФЦ_НОВЫЙ", 4, 5)])]
    assert data["items"][0]["excluded_units"] == 10

    # в запросе: ни один кластер не исключён и вывезти всего 5 шт. — по порядку «хуже — первыми»
    data = (await get(excluded="", total=5)).json()
    assert (data["possible"], data["planned"]) == (12, 5)
    assert summary(data) == [
        (ACB, "idle", 3, [("Санкт-Петербург и СЗО", "СПБ_БУГРЫ_РФЦ", 3, 2)]),
        (TH, "selling", 2, [("Санкт-Петербург и СЗО", "СПБ_БУГРЫ_РФЦ", 2, 8)]),
    ]
    data = (await get(excluded="", total=3)).json()
    assert (data["planned"], data["left_out"], len(data["items"])) == (3, 1, 1)


async def test_report_validation_and_isolation(client, tokens, seller):
    r = await client.get(URL, params={"keep_idle": -1}, headers=auth(tokens))
    assert r.status_code == 400 and r.json()["detail"].startswith("Оставлять у товара без продаж")
    r = await client.put(f"{URL}/params", json={"keep_selling": 20000}, headers=auth(tokens))
    assert r.status_code == 400 and r.json()["detail"].startswith("Оставлять у продающегося")
    other = {"email": "other@example.com", "password": "strongpass2"}
    await client.post("/api/v1/auth/register", json=other)
    t2 = (await client.post("/api/v1/auth/login", json=other)).json()
    data = (await client.get(URL, headers=auth(t2))).json()
    assert (data["items"], data["stocks_known"], data["clusters"]) == ([], False, [])


async def test_report_without_sales(sessions, client, tokens, ozon):  # noqa: F811
    await api_user(sessions, client, tokens, ozon)
    data = (await client.get(URL, headers=auth(tokens))).json()
    assert (data["sales_known"], data["items"], data["possible"]) == (False, [], 0)


async def test_export_xlsx(client, tokens, seller):
    await client.put(f"{URL}/params", json={"keep_selling": 5, "excluded": [SPB_CLUSTER], "total": 100},
                     headers=auth(tokens))
    r = await client.get(f"{URL}/export", params={"excluded": ""}, headers=auth(tokens))
    assert r.status_code == 200 and r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "withdrawal.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Вывоз", "Параметры"]
    ws = wb["Вывоз"]
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    assert rows[0][:5] == ["Артикул", "Название", "Кластер", "Склад", "Сколько вывезти, шт."]
    assert [r[:7] for r in rows[1:4]] == [
        [ACB, None, "Санкт-Петербург и СЗО", "СПБ_БУГРЫ_РФЦ", 3, 5, 2],
        [TH, None, "Санкт-Петербург и СЗО", "СПБ_БУГРЫ_РФЦ", 5, 10, 5],
        [TH, None, "Екатеринбург", "ЕКАТЕРИНБУРГ_РФЦ_НОВЫЙ", 4, 9, 5],
    ]
    assert rows[1][7] == "Нет продаж 45 дн." and isinstance(rows[1][8], datetime)
    assert rows[-1][0] == "Итого" and rows[-1][4] == 12
    notes = [c.value for c in wb["Параметры"]["A"] if c.value]
    assert "Вывезти всего: 100 шт." in notes
    assert "Не трогать кластеры: нет" in notes  # в запросе исключение снято — файл как на экране


# ── 10.17: причина, приоритет, замороженные деньги ──
def test_reason_priority_and_frozen_money():
    t = Thresholds(7, 14, 60)
    idle = plan([item("A", [place("w1", SPB, 5)], "0", 45)], Params(), t).candidates[0]
    assert (engine.reason_kind(idle, t), engine.PRIORITY[engine.reason_kind(idle, t)]) == (engine.NO_SALES, engine.HIGH)
    # 130 шт. по 1 в день — 130 дней ≥ двух порогов «Избыток» (120): низкая оборачиваемость
    slow = plan([item("B", [place("w1", SPB, 130)], "1", 1)], Params(keep_selling=0), t).candidates[0]
    assert engine.reason_kind(slow, t) == engine.LOW_TURNOVER and engine.PRIORITY[engine.LOW_TURNOVER] == engine.MEDIUM
    # 80 дней — больше порога, но меньше двух: избыточный запас
    extra = plan([item("C", [place("w1", SPB, 80)], "1", 1)], Params(keep_selling=0), t).candidates[0]
    assert engine.reason_kind(extra, t) == engine.EXCESS_STOCK and engine.PRIORITY[engine.EXCESS_STOCK] == engine.NORMAL
    # замороженные деньги — сколько вывезти × себестоимость с НДС; нет себестоимости — нет данных, а не 0
    assert engine.frozen_money(extra.qty, Decimal("120")) == Decimal("2400.00")  # вывезти 20 шт. сверх 60 дней
    assert engine.frozen_money(3, Decimal("33.335")) == Decimal("100.01")
    assert engine.frozen_money(5, None) is None


async def test_sold_by_product(sessions, client, tokens):
    from tests.test_pnl_api import user_id

    uid = await user_id(client, tokens)
    today = today_msk()
    for i, qty in ((1, 2), (10, 3), (60, 5)):
        await add_api_day(sessions, uid, today - timedelta(days=i), [("Выручка", "500", qty, "W-1")])
    await add_api_day(sessions, uid, today - timedelta(days=2), [("Выручка", "0", None, None)])
    async with sessions() as db:
        assert set((await sold_by_product(db, uid, 30, today)).values()) == {5}
        assert set((await sold_by_product(db, uid, 90, today)).values()) == {10}
        assert await sold_by_product(db, uid, 30, today + timedelta(days=400)) == {}  # дней с данными нет — не 0


@pytest.fixture
async def priced(client, tokens, seller, monkeypatch):
    """TH — себестоимость 100 × 1,2; у ACB себестоимости нет. Продажи: TH — 6 за 30 дней и 18 за 90; ACB — 0."""
    await client.put(f"/api/v1/costs/{TH}", json={"cost_price": "100", "vat_multiplier": "1.2"}, headers=auth(tokens))

    async def sold(db, uid, days, today=None):
        return {seller[TH]: 6 if days == 30 else 18, seller[ACB]: 0}

    monkeypatch.setattr(withdrawal_svc, "sold_by_product", sold)
    return seller


async def test_report_reasons_frozen_and_filters(client, tokens, priced):
    get = lambda **q: client.get(URL, params={"keep_selling": 5, **q}, headers=auth(tokens))  # noqa: E731
    data = (await get()).json()
    acb, th = data["items"]
    assert (acb["reason_kind"], acb["reason_title"], acb["priority"], acb["frozen"], acb["sold_30"]) == (
        "no_sales", "Нет продаж 45 дн.", "high", None, 0)
    # TH: 155 дней запаса ≥ 120 — низкая оборачиваемость; вывезти 9 × 120 ₽ = 1080 ₽ заморожено
    assert (th["reason_kind"], th["reason_title"], th["priority"], th["qty"], th["unit_cost"], th["frozen"]) == (
        "low_turnover", "Низкая оборачиваемость", "medium", 9, "120.00", "1080.00")
    assert (th["sold_30"], th["sold_90"], th["days_of_supply"], th["available"]) == (6, 18, "155.0", 31)
    assert (data["frozen"], data["without_cost"]) == ("1080.00", 1)
    assert data["counts"] == {"priority": {"high": 1, "medium": 1, "normal": 0},
                              "decision": {"undecided": 2, "remove": 0, "keep": 0}}
    assert [(w["name"], w["qty"], w["products"]) for w in data["warehouses"]] == [
        ("СПБ_БУГРЫ_РФЦ", 8, 2), ("ЕКАТЕРИНБУРГ_РФЦ_НОВЫЙ", 4, 1)]
    assert data["cabinet_url"].startswith("https://seller.ozon.ru") and "Вывоз и утилизация" in data["cabinet_path"]

    # фильтр склада — только его строки, количество и деньги по ним; план (лимит) — по всем товарам
    ekb = data["warehouses"][1]["id"]
    data = (await get(warehouse_id=ekb)).json()
    assert [(it["product"]["offer_id"], it["qty"], it["qty_all"], it["frozen"]) for it in data["items"]] == [
        (TH, 4, 9, "480.00")]
    assert (data["shown"], data["planned"]) == ({"items": 1, "qty": 4, "frozen": "480.00", "without_cost": 0}, 12)
    data = (await get(priority="high")).json()
    assert [it["product"]["offer_id"] for it in data["items"]] == [ACB]
    assert (await get(priority="urgent")).status_code == 400


async def test_decisions(sessions, client, tokens, priced):
    post = lambda body: client.post(f"{URL}/decisions", json=body, headers=auth(tokens))  # noqa: E731
    get = lambda **q: client.get(URL, params={"keep_selling": 5, **q}, headers=auth(tokens))  # noqa: E731
    acb = str(priced[ACB])
    r = await post({"product_id": acb, "decision": "remove", "qty": 3})
    assert r.status_code == 200 and (r.json()["decision"], r.json()["qty"]) == ("remove", 3)
    data = (await get()).json()
    assert data["items"][0]["decision"]["decision"] == "remove" and data["items"][1]["decision"] is None
    assert data["counts"]["decision"] == {"undecided": 1, "remove": 1, "keep": 0}
    assert [it["product"]["offer_id"] for it in (await get(decision="remove")).json()["items"]] == [ACB]
    assert [it["product"]["offer_id"] for it in (await get(decision="undecided")).json()["items"]] == [TH]

    # передумал — действует последнее; «отменить» — решения нет, записи не удаляются
    await post({"product_id": acb, "decision": "keep"})
    assert (await get()).json()["items"][0]["decision"]["decision"] == "keep"
    r = await post({"product_id": acb, "decision": "undecided"})
    assert r.status_code == 200 and r.json() is None
    assert (await get()).json()["items"][0]["decision"] is None
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(WithdrawalDecision)) == 3

    # в Excel — приоритет, замороженные деньги, решение
    await post({"product_id": str(priced[TH]), "decision": "remove", "qty": 9})
    r = await client.get(f"{URL}/export", params={"keep_selling": 5}, headers=auth(tokens))
    rows = [list(row) for row in load_workbook(io.BytesIO(r.content))["Вывоз"].iter_rows(values_only=True)]
    assert rows[0][9:12] == ["Приоритет", "Замороженные деньги, ₽", "Решение"]
    assert rows[1][9:11] == ["высокий", "нет себестоимости"]
    assert rows[2][9:12] == ["средний", 600, f"вывозим, {today_msk():%d.%m.%Y}"]  # СПб: 5 × 120
    assert rows[-1][0] == "Итого" and rows[-1][10] == 1080

    # чужой товар, неверное решение
    assert (await post({"product_id": str(uuid.uuid4()), "decision": "keep"})).status_code == 404
    assert (await post({"product_id": acb, "decision": "maybe"})).status_code == 400


async def test_removals_from_ozon(client, tokens, seller):
    calls = []

    def handler(request: httpx.Request) -> httpx.Response:
        import json

        body = json.loads(request.content)
        calls.append(body)
        if not body["last_id"]:
            rows = [{"offer_id": ACB, "name": "Подшипник", "sku": "900001002", "quantity_for_return": 2,
                     "return_created_at": "2026-10-20T06:49:37.591Z", "return_state": "В пути", "box_state": "Собран",
                     "clearing_warehouse_name": "СПБ_БУГРЫ_РФЦ", "destination_warehouse_name": "РАЗДОРЫ_4",
                     "is_auto_return": False, "stock_type": "Излишек", "delivery_type": "Вывоз с ПВЗ/СЦ"}] * 500
            return httpx.Response(200, json={"returns_summary_report_rows": rows, "last_id": "next"})
        return httpx.Response(200, json={"returns_summary_report_rows": [
            {"offer_id": "NOT-IN-CATALOG", "quantity_for_return": 1, "return_created_at": "2026-10-22T10:00:00Z"}],
            "last_id": ""})

    with respx.mock(base_url=settings.OZON_API_URL, assert_all_called=False) as mock:
        mock.post("/v1/removal/from-stock/list").mock(side_effect=handler)
        r = await client.get(f"{URL}/removals", params={"days": 30}, headers=auth(tokens))
    assert r.status_code == 200, r.text
    data = r.json()
    assert calls[0]["date_to"] == today_msk().isoformat() and calls[0]["limit"] == 500 and calls[1]["last_id"] == "next"
    assert (len(data["rows"]), data["qty"]) == (501, 1001)
    first = data["rows"][0]  # новые — первыми
    assert (first["offer_id"], first["product_id"], first["qty"]) == ("NOT-IN-CATALOG", None, 1)
    second = data["rows"][1]
    assert (second["product_id"], second["state"], second["warehouse"], second["destination"]) == (
        str(seller[ACB]), "В пути", "СПБ_БУГРЫ_РФЦ", "РАЗДОРЫ_4")

    with respx.mock(base_url=settings.OZON_API_URL, assert_all_called=False) as mock:
        mock.post("/v1/removal/from-stock/list").mock(return_value=httpx.Response(500, json={"message": "boom"}))
        r = await client.get(f"{URL}/removals", headers=auth(tokens))
    assert r.status_code == 502 and r.json()["detail"].startswith("Ozon не отдал список вывозов")
    assert (await client.get(f"{URL}/removals", params={"days": 7}, headers=auth(tokens))).status_code == 400


async def test_removals_without_key(client, tokens):
    r = await client.get(f"{URL}/removals", headers=auth(tokens))
    assert r.status_code == 400 and "ключ" in r.json()["detail"].lower()
