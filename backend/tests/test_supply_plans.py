"""Шаг 10.18: страховой запас, планы поставок с ручной корректировкой (версии строк), «уже в планах и заявках»,
календарь, заявки Ozon с историей и фактом против плана, продажи после приёмки, создание заявки в Ozon только после
проверки и подтверждения (правило 6)."""
import io
import json
import uuid
from datetime import date, timedelta

import httpx
import pytest
import respx
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.config import settings
from app.engine.supply_planner import Params, recommend
from app.models.audit import AuditLog
from app.models.supply import SupplyOrder, SupplyPlanItem
from app.services import supply_ozon
from app.services.sync.finance import today_msk
from tests.test_pnl_api import add_api_day
from tests.test_supply import point, seller  # noqa: F401 — фикстура продаж и остатков
from tests.test_supply_orders import FakeSupply, load, run_sync
from tests.test_sync_stocks import auth
from tests.test_sync_stocks import ozon as stocks_ozon  # noqa: F401
from tests.test_warehouses import ACB, CM, TH, product_ids

URL = "/api/v1/supply"
EKB_WH = 1020000000000101


def test_safety_days():
    """2 в день: цель 45 + страховой 10 дней = 110 шт.; к приходу поставки (7 дней) — 2 × 62 = 124 → везти 124 − 30 − 10."""
    r = recommend(point(30, 60, 21, transit=10), Params(safety_days=10))
    assert (r.target_stock, r.need, r.to_order) == (110, 124, 84)
    assert recommend(point(30, 60, 21, transit=10)).to_order == 64  # без страхового — как в 5.2


async def test_safety_days_setting(client, tokens, seller):  # noqa: F811
    r = await client.patch("/api/v1/settings", json={"supply_safety_days": 7}, headers=auth(tokens))
    assert r.json()["supply_safety_days"] == 7
    data = (await client.get(f"{URL}/recommendations", headers=auth(tokens))).json()
    assert data["params"]["safety_days"] == 7
    th = next(i for i in data["items"] if i["product"]["offer_id"] == TH and i["kind"] == "supply")
    assert (th["need"], th["to_order"]) == (118, 109)  # ⌈2 × (45 + 7 + 7)⌉ − 9


def rec(data, offer, kind="supply") -> dict:
    return next(i for i in data["items"] if i["product"]["offer_id"] == offer and i["kind"] == kind)


async def test_plan_corrections_and_calendar(sessions, client, tokens, seller):  # noqa: F811
    ids = await product_ids(sessions, seller.user_id)
    post = lambda body: client.post(f"{URL}/plans", json=body, headers=auth(tokens))  # noqa: E731
    day = today_msk() + timedelta(days=5)
    body = {"kind": "supply", "destination_key": "екатеринбург", "destination_name": "Екатеринбург",
            "planned_date": day.isoformat(), "items": [
                {"product_id": str(ids[TH]), "qty": 95, "recommended": 95},
                {"product_id": str(ids[ACB]), "qty": 0}]}
    r = await post(body)
    assert r.status_code == 201, r.text
    plan = r.json()
    assert (plan["status"], plan["units"], plan["products"], plan["corrected"], plan["editable"]) == ("draft", 95, 1, 0, True)
    pid = plan["id"]

    # проверки: пустой план, чужой товар, повтор товара
    assert (await post({**body, "items": [{"product_id": str(ids[TH]), "qty": 0}]})).status_code == 400
    assert (await post({**body, "items": [{"product_id": str(uuid.uuid4()), "qty": 1}]})).status_code == 400
    assert (await post({**body, "items": [{"product_id": str(ids[TH]), "qty": 1}] * 2})).status_code == 400

    # ручная корректировка — новая версия, прежняя остаётся
    r = await client.put(f"{URL}/plans/{pid}/items/{ids[TH]}", json={"qty": 80}, headers=auth(tokens))
    detail = r.json()
    th = next(i for i in detail["items"] if i["product"]["offer_id"] == TH)
    assert (th["qty"], th["recommended"], [h["qty"] for h in th["history"]], detail["corrected"]) == (80, 95, [95, 80], 1)
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(SupplyPlanItem)) == 3

    # рекомендация знает про план: везти 95, в плане 80 — нужно ещё 15
    data = (await client.get(f"{URL}/recommendations", headers=auth(tokens))).json()
    assert (rec(data, TH)["to_order"], rec(data, TH)["in_plans"], rec(data, TH)["still_needed"]) == (95, 80, 15)

    # Excel плана и календарь
    r = await client.get(f"{URL}/plans/{pid}/export", headers=auth(tokens))
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows[1][0] == TH and rows[1][3:5] == (95, 80)
    cal = (await client.get(f"{URL}/calendar", headers=auth(tokens))).json()
    assert [(e["date"], e["kind"], e["units"]) for e in cal["events"]] == [(day.isoformat(), "plan", 80)]

    # отмена: план не считается и не меняется; вернуть в работу — снова можно
    await client.patch(f"{URL}/plans/{pid}", json={"status": "cancelled"}, headers=auth(tokens))
    assert rec((await client.get(f"{URL}/recommendations", headers=auth(tokens))).json(), TH)["in_plans"] == 0
    r = await client.put(f"{URL}/plans/{pid}/items/{ids[TH]}", json={"qty": 70}, headers=auth(tokens))
    assert r.status_code == 400 and "отменён" in r.json()["detail"]
    r = await client.patch(f"{URL}/plans/{pid}", json={"status": "draft", "note": "Машина в четверг"}, headers=auth(tokens))
    assert (r.json()["status"], r.json()["note"]) == ("draft", "Машина в четверг")

    # привязать заявку, оформленную в кабинете: план фиксируется, заявка ещё готовится — «уже в заявках»
    async with sessions() as db:
        from datetime import datetime, timezone

        from app.models.supply import BUNDLE, SupplyOrderItem

        order = SupplyOrder(user_id=seller.user_id, ozon_order_id=5002, order_number="2000000005002",
                            state="READY_TO_SUPPLY", cluster_name="Екатеринбург", synced_at=datetime.now(timezone.utc))
        db.add(order)
        await db.flush()
        db.add(SupplyOrderItem(order_id=order.id, source=BUNDLE, snapshot_at=datetime.now(timezone.utc), sku=900001001,
                               offer_id=TH, product_id=ids[TH], quantity=12))
        await db.commit()
    assert (await client.patch(f"{URL}/plans/{pid}", json={"ozon_order_id": 777}, headers=auth(tokens))).status_code == 400
    r = await client.patch(f"{URL}/plans/{pid}", json={"ozon_order_id": 5002}, headers=auth(tokens))
    assert (r.json()["status"], r.json()["order"]["number"], r.json()["editable"]) == ("ordered", "2000000005002", False)
    assert (await client.put(f"{URL}/plans/{pid}/items/{ids[TH]}", json={"qty": 70}, headers=auth(tokens))).status_code == 400
    data = (await client.get(f"{URL}/recommendations", headers=auth(tokens))).json()
    assert (rec(data, TH)["in_plans"], rec(data, TH)["in_orders"], rec(data, TH)["still_needed"]) == (0, 12, 83)
    # отвязать — снова черновик
    r = await client.patch(f"{URL}/plans/{pid}", json={"ozon_order_id": None}, headers=auth(tokens))
    assert r.json()["status"] == "draft" and r.json()["order"] is None

    # чужой магазин
    other = {"email": "other@example.com", "password": "strongpass2"}
    await client.post("/api/v1/auth/register", json=other)
    t2 = (await client.post("/api/v1/auth/login", json=other)).json()
    assert (await client.get(f"{URL}/plans/{pid}", headers=auth(t2))).status_code == 404
    assert (await client.get(f"{URL}/plans", headers=auth(t2))).json() == []


async def test_orders_results_and_sales_after(sessions, client, tokens, stocks_ozon):  # noqa: F811
    from tests.test_sync_stocks import api_user

    user_id = await api_user(sessions, client, tokens, stocks_ozon)
    with respx.mock(base_url=settings.OZON_API_URL, assert_all_called=False) as mock:
        FakeSupply(mock)
        await run_sync(sessions, user_id)
    # продажи TH: до приёмки (27.09) — 1 шт. за день с данными, после — 2 и 3 шт.
    for d, qty in ((date(2026, 9, 20), 1), (date(2026, 9, 28), 2), (date(2026, 10, 1), 3)):
        await add_api_day(sessions, user_id, d, [("Выручка", str(500 * qty), qty, TH)])

    data = (await client.get(f"{URL}/orders", headers=auth(tokens))).json()
    assert data["total"] == 3 and data["counts"] == {"preparing": 1, "on_the_way": 0, "accepting": 0, "done": 1,
                                                     "failed": 1}
    done = next(o for o in data["items"] if o["ozon_order_id"] == 5001)
    assert (done["state_text"], done["cluster"], done["declared"], done["accepted"], done["accepted_on"]) == (
        "Завершена", "Санкт-Петербург и СЗО", 25, 23, "2026-09-27")
    ready = next(o for o in data["items"] if o["ozon_order_id"] == 5002)
    assert (ready["accepted"], ready["timeslot_day"], ready["timeslot_time"]) == (None, "2026-10-28", "12:00–13:00")
    assert [o["ozon_order_id"] for o in (await client.get(f"{URL}/orders", params={"search": "acb"},
                                                          headers=auth(tokens))).json()["items"]] == [5001]
    assert [o["ozon_order_id"] for o in (await client.get(f"{URL}/orders", params={"group": "failed"},
                                                          headers=auth(tokens))).json()["items"]] == [5003]

    r = await client.get(f"{URL}/orders/{done['id']}", headers=auth(tokens))
    detail = r.json()
    assert [e["state"] for e in detail["events"]] == ["COMPLETED"] and detail["has_acts"] is True
    th = next(i for i in detail["items"] if i["product"]["offer_id"] == TH)
    # заявлено 20, принято 18 (брак) — недовоз 2; после приёмки 5 шт. за 2 дня с данными (2,5 в день) против 2 до
    assert (th["declared"], th["accepted"], th["diff"], th["defect"]) == (20, 18, -2, True)
    after = th["sales_after"]
    assert (after["after_qty"], after["after_days_with_data"], after["per_day_after"], after["before_qty"]) == (
        5, 2, "2.50", 2)
    assert (after["change_pct"], th["lasts_days"]) == ("150.0", 7)  # 18 / 2,5
    assert detail["totals"] == {"planned": None, "declared": 25, "accepted": 23, "diff": -2, "short": 2}
    # нет актов — «нет данных», а не 0
    detail = (await client.get(f"{URL}/orders/{ready['id']}", headers=auth(tokens))).json()
    assert detail["has_acts"] is False and detail["totals"]["accepted"] is None
    assert [(i["product"]["offer_id"], i["product"]["id"]) for i in detail["items"]][-1] == ("UNKNOWN-1", None)

    # календарь: таймслоты по поясу склада, приёмка
    cal = (await client.get(f"{URL}/calendar", params={"date_from": "2026-09-20", "date_to": "2026-10-31"},
                            headers=auth(tokens))).json()
    assert [(e["date"], e["kind"]) for e in cal["events"]] == [
        ("2026-09-25", "timeslot"), ("2026-09-27", "accepted"), ("2026-10-28", "timeslot")]
    assert (await client.get(f"{URL}/calendar", params={"date_from": "2026-01-01", "date_to": "2026-12-31"},
                             headers=auth(tokens))).status_code == 400

    r = await client.get(f"{URL}/orders/export", headers=auth(tokens))
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows[0][0] == "Номер заявки" and {row[0] for row in rows[1:4]} == {"2000000005001", "2000000005002",
                                                                               "2000000005003"}
    assert (await client.get(f"{URL}/orders/{uuid.uuid4()}", headers=auth(tokens))).status_code == 404


class FakeDrafts:
    def __init__(self, mock: respx.MockRouter) -> None:
        self.calls: dict[str, list] = {"draft": [], "info": [], "slots": [], "create": [], "status": []}
        self.create_response = load("draft_supply_create")
        mock.post("/v1/cluster/list").mock(return_value=httpx.Response(200, json=load("cluster_list")))
        for path, key, name in ((supply_ozon.DRAFT_CREATE, "draft", "draft_create"),
                                (supply_ozon.DRAFT_INFO, "info", "draft_create_info"),
                                (supply_ozon.TIMESLOTS, "slots", "draft_timeslot_info"),
                                (supply_ozon.SUPPLY_STATUS, "status", "draft_supply_create_status")):
            mock.post(path).mock(side_effect=self._answer(key, name))
        mock.post(supply_ozon.SUPPLY_CREATE).mock(side_effect=self._create)

    def _answer(self, key: str, name: str):
        def handler(request):
            self.calls[key].append(json.loads(request.content))
            return httpx.Response(200, json=load(name))
        return handler

    def _create(self, request):
        self.calls["create"].append(json.loads(request.content))
        return httpx.Response(200, json=self.create_response)


@pytest.fixture
def drafts():
    with respx.mock(base_url=settings.OZON_API_URL, assert_all_called=False) as mock:
        yield FakeDrafts(mock)


async def test_create_order_only_after_check_and_confirm(sessions, client, tokens, seller, drafts):  # noqa: F811
    ids = await product_ids(sessions, seller.user_id)
    plan = (await client.post(f"{URL}/plans", json={
        "kind": "supply", "destination_key": "екатеринбург", "destination_name": "Екатеринбург",
        "items": [{"product_id": str(ids[TH]), "qty": 95, "recommended": 95}, {"product_id": str(ids[CM]), "qty": 4}],
    }, headers=auth(tokens))).json()
    pid = plan["id"]

    # проверка: черновик в Ozon — кластер по названию, SKU и количество; заявку не создаёт
    r = await client.post(f"{URL}/plans/{pid}/ozon-check", headers=auth(tokens))
    assert r.status_code == 200, r.text
    check = r.json()
    assert drafts.calls["draft"] == [{"cluster_info": {"items": [{"sku": 900001001, "quantity": 95},
                                                                 {"sku": 900001004, "quantity": 4}],
                                                       "macrolocal_cluster_id": 154}, "deletion_sku_mode": "PARTIAL"}]
    assert (check["status"], check["cluster_name"]) == ("ready", "Екатеринбург")
    assert [(w["name"], w["available"], w["reason"]) for w in check["warehouses"]] == [
        ("ЕКАТЕРИНБУРГ_РФЦ_НОВЫЙ", True, None), ("ЕКАТЕРИНБУРГ_2", False, "WAREHOUSE_OVERLOADED")]
    assert [(x["offer_id"], x["reasons"]) for x in check["rejected"]] == [(CM, ["SKU_IS_NOT_ALLOWED_FOR_SUPPLY"])]
    assert drafts.calls["create"] == []

    slots = (await client.get(f"{URL}/drafts/{check['id']}/timeslots", params={"warehouse_id": EKB_WH},
                              headers=auth(tokens))).json()
    assert [len(d["slots"]) for d in slots["days"]] == [2, 1] and slots["days"][0]["slots"][0]["from"] == "2026-10-29T09:00:00"
    sent = drafts.calls["slots"][0]
    assert (sent["supply_type"], sent["selected_cluster_warehouses"]) == (
        "DIRECT", [{"macrolocal_cluster_id": 154, "storage_warehouse_id": EKB_WH}])
    assert date.fromisoformat(sent["date_to"]) <= today_msk() + timedelta(days=27)

    create = lambda draft, **kw: client.post(f"{URL}/drafts/{draft}/create", json={  # noqa: E731
        "warehouse_id": EKB_WH, "timeslot_from": "2026-10-29T09:00:00", "timeslot_to": "2026-10-29T10:00:00", **kw},
        headers=auth(tokens))
    # без подтверждения, на недоступный склад — ничего не уходит в Ozon
    r = await create(check["id"])
    assert r.status_code == 400 and "Подтвердите" in r.json()["detail"]
    r = await create(check["id"], confirm=True, warehouse_id=1020000000000111)
    assert r.status_code == 400 and "не принимает" in r.json()["detail"]
    # план изменили после проверки — проверить заново
    await client.put(f"{URL}/plans/{pid}/items/{ids[CM]}", json={"qty": 0}, headers=auth(tokens))
    r = await create(check["id"], confirm=True)
    assert r.status_code == 400 and "изменили после проверки" in r.json()["detail"]
    assert drafts.calls["create"] == []

    check = (await client.post(f"{URL}/plans/{pid}/ozon-check", headers=auth(tokens))).json()
    r = await create(check["id"], confirm=True)
    assert r.status_code == 200, r.text
    assert r.json()["ozon_order_id"] == 5009 and len(drafts.calls["create"]) == 1
    assert drafts.calls["create"][0]["timeslot"] == {"from_in_timezone": "2026-10-29T09:00:00",
                                                     "to_in_timezone": "2026-10-29T10:00:00"}
    detail = (await client.get(f"{URL}/plans/{pid}", headers=auth(tokens))).json()
    assert (detail["status"], detail["order"]["ozon_order_id"], detail["order"]["state_text"]) == (
        "ordered", 5009, "Загружается")
    async with sessions() as db:
        entry = (await db.execute(select(AuditLog).where(AuditLog.entity_type == "supply_order"))).scalar_one()
        assert (entry.action, entry.status, entry.after_json["ozon_order_id"]) == ("create", "applied", 5009)
        assert entry.ozon_request["draft_id"] == 880001
    # одна попытка на проверку; план с заявкой больше не отправляется
    assert (await create(check["id"], confirm=True)).status_code == 400
    assert (await client.post(f"{URL}/plans/{pid}/ozon-check", headers=auth(tokens))).status_code == 400


async def test_create_order_rejected_by_ozon_and_purchase(sessions, client, tokens, seller, drafts):  # noqa: F811
    ids = await product_ids(sessions, seller.user_id)
    plan = (await client.post(f"{URL}/plans", json={
        "kind": "supply", "destination_key": "екатеринбург", "destination_name": "Екатеринбург",
        "items": [{"product_id": str(ids[TH]), "qty": 10}]}, headers=auth(tokens))).json()
    check = (await client.post(f"{URL}/plans/{plan['id']}/ozon-check", headers=auth(tokens))).json()
    drafts.create_response = {"draft_id": 880001, "error_reasons": ["TIMESLOT_NOT_AVAILABLE"]}
    r = await client.post(f"{URL}/drafts/{check['id']}/create", json={
        "warehouse_id": EKB_WH, "timeslot_from": "2026-10-29T09:00:00", "timeslot_to": "2026-10-29T10:00:00",
        "confirm": True}, headers=auth(tokens))
    assert r.status_code == 400 and "таймслот уже занят" in r.json()["detail"]
    async with sessions() as db:
        entry = (await db.execute(select(AuditLog).where(AuditLog.entity_type == "supply_order"))).scalar_one()
        assert (entry.status, entry.after_json) == ("rejected", None)
    assert (await client.get(f"{URL}/plans/{plan['id']}", headers=auth(tokens))).json()["status"] == "draft"

    # закупка на свой склад — без Ozon; кластера нет в Ozon — понятная ошибка
    purchase = (await client.post(f"{URL}/plans", json={
        "kind": "purchase", "destination_key": "20605650762000", "destination_name": "Склад Подольск",
        "items": [{"product_id": str(ids[TH]), "qty": 6}]}, headers=auth(tokens))).json()
    r = await client.post(f"{URL}/plans/{purchase['id']}/ozon-check", headers=auth(tokens))
    assert r.status_code == 400 and "закупка" in r.json()["detail"]
    nowhere = (await client.post(f"{URL}/plans", json={
        "kind": "supply", "destination_key": "новосибирск", "destination_name": "Новосибирск",
        "items": [{"product_id": str(ids[ACB]), "qty": 26}]}, headers=auth(tokens))).json()
    r = await client.post(f"{URL}/plans/{nowhere['id']}/ozon-check", headers=auth(tokens))
    assert r.status_code == 400 and "не найден в Ozon" in r.json()["detail"]
