"""Плановые платежи и платёжный календарь (шаг 3.9): повторения и переопределения, план-факт, перенос, разрыв и порог,
«что если», генераторы (Ozon, поставщик по профилю, налоги), сроки налогов по режимам и перенос с выходных, смена
режима посреди года — с пересчётом прошлых платежей и без, оплаченные не меняются, «Ближайшие сроки»."""
import io
import uuid
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import select

from app.api.v1 import plan_calendar as plan_api
from app.api.v1 import planning as planning_api
from app.api.v1 import tax_profile as tax_api
from app.engine import payment_calendar as cal
from app.engine.payouts import monday
from app.engine.tax_deadlines import Segment, obligations
from app.models.accrual import SOURCE_API, Accrual, AccrualDay, AccrualType
from app.models.product import Product
from app.services import ozon_payouts, plan_calendar, planning, pnl_service, tax_profile
from tests.test_payouts import WEEKS, add_days, history
from tests.test_pnl_api import user_id
from tests.test_report_pnl import auth, import_reference_costs

TODAY = date(2026, 10, 14)  # среда
URL = "/api/v1/planning"
D = Decimal
OFFER = "ACB305523.NAC.JP"  # себестоимость 851 ₽, НДС поставщика 22 % — 1 038,22 ₽ за штуку


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    for module in (planning, planning_api, plan_calendar, plan_api, tax_profile, tax_api, ozon_payouts, pnl_service):
        monkeypatch.setattr(module, "today_msk", lambda: TODAY)


async def call(client, tokens, method, path, expect=200, base=URL, **kw):
    r = await client.request(method, f"{base}{path}", headers=auth(tokens), **kw)
    assert r.status_code == expect, r.text
    return r.json() if r.content else None


async def setup(client, tokens, supplier="upfront", balance="100000", **extra):
    body = {"supplier_payment": supplier, "accounts": [{"name": "Расчётный счёт", "kind": "checking",
                                                        "balance": balance}],
            "ozon_account": 0, "categories": ["rent", "salary", "taxes", "supplier", "purchase"], **extra}
    await call(client, tokens, "POST", "/setup", json=body)
    cats = {c["code"] or c["name"]: c["id"] for c in await call(client, tokens, "GET", "/categories")}
    accounts = {a["name"]: a["id"] for a in await call(client, tokens, "GET", "/accounts")}
    return cats, accounts


async def calendar(client, tokens, date_from="2026-10-14", date_to="2026-12-12", **params):
    return await call(client, tokens, "GET", "/calendar", params={"date_from": date_from, "date_to": date_to, **params})


def lines(body, source=None):
    return [ln for ln in body["lines"] if source is None or ln["source"] == source]


async def add_sales(sessions, uid, days: list[date], revenue: str, qty: int | None = None, offer: str | None = None):
    """Начисления из API: «Выручка» каждый день (с товаром — себестоимость считается), дни отмечены загруженными."""
    async with sessions() as db:
        rev = (await db.execute(select(AccrualType).where(AccrualType.name_ru == "Выручка"))).scalar_one_or_none()
        if rev is None:
            rev = AccrualType(id=901, name_ru="Выручка")
            db.add(rev)
        product_id = None
        if offer:
            product = (await db.execute(select(Product).where(Product.offer_id == offer))).scalar_one_or_none()
            if product is None:
                product = Product(id=uuid.uuid4(), user_id=uid, ozon_product_id=1, ozon_sku=1001, offer_id=offer,
                                  name="Подшипник", synced_at=datetime.now(timezone.utc))
                db.add(product)
            product_id = product.id
        await db.flush()
        for d in days:
            db.add(Accrual(user_id=uid, product_id=product_id, accrual_type_id=rev.id, accrual_date=d, amount=D(revenue),
                           quantity=qty, source=SOURCE_API, sync_scope=f"day:{d}", line_no=0))
            db.add(AccrualDay(user_id=uid, day=d, synced_at=datetime.now(timezone.utc), entries=1))
        await db.commit()


def days(a: date, b: date) -> list[date]:
    return [a + timedelta(days=i) for i in range((b - a).days + 1)]


# ---------- движок: повторения, переопределения, прогноз ----------


def test_occurrences():
    occ = cal.occurrences
    monthly = cal.Rule(date(2026, 1, 31), "month")
    assert occ(monthly, date(2026, 1, 1), date(2026, 4, 30)) == [
        date(2026, 1, 31), date(2026, 2, 28), date(2026, 3, 31), date(2026, 4, 30)]  # 31-е — последний день месяца
    assert occ(cal.Rule(date(2026, 1, 10), "month", day=5), date(2026, 1, 1), date(2026, 3, 31)) == [
        date(2026, 2, 5), date(2026, 3, 5)]  # раньше начала — не платим
    assert occ(cal.Rule(date(2026, 1, 15), "quarter"), date(2026, 3, 1), date(2027, 1, 31)) == [
        date(2026, 4, 15), date(2026, 7, 15), date(2026, 10, 15), date(2027, 1, 15)]
    assert occ(cal.Rule(date(2024, 2, 29), "year"), date(2025, 1, 1), date(2028, 12, 31)) == [
        date(2025, 2, 28), date(2026, 2, 28), date(2027, 2, 28), date(2028, 2, 29)]
    weekly = cal.Rule(date(2026, 10, 1), "week", day=5, end=date(2026, 10, 25))  # по пятницам до 25 октября
    assert occ(weekly, date(2026, 10, 1), date(2026, 12, 31)) == [
        date(2026, 10, 2), date(2026, 10, 9), date(2026, 10, 16), date(2026, 10, 23)]
    assert occ(cal.Rule(date(2026, 10, 20)), date(2026, 10, 1), date(2026, 10, 31)) == [date(2026, 10, 20)]
    assert occ(cal.Rule(date(2026, 10, 20)), date(2026, 11, 1), date(2026, 11, 30)) == []


def item(key, day, amount, kind="expense", **kw):
    return cal.Item(key=key, source=kw.pop("source", "manual"), title=key, planned_day=day, amount=D(amount), kind=kind, **kw)


def test_overrides_forecast_gap_threshold_and_what_if():
    t = TODAY
    items = [
        item("rent", t + timedelta(days=2), "-60000"),
        item("ozon", t + timedelta(days=7), "50000", kind="income", source="ozon"),
        item("salary", t + timedelta(days=5), "-50000"),
        item("ads", t + timedelta(days=3), "-10000", enabled=False),  # выключен — «что если»
        item("old", t - timedelta(days=3), "-5000"),  # прошлый неоплаченный — в прогноз не входит
    ]
    fc = cal.forecast(D(100000), items, t, t + timedelta(days=10), threshold=D(30000))
    by_day = {d.day: d.balance for d in fc.days}
    assert by_day[t] == D(100000) and by_day[t + timedelta(days=2)] == D(40000)
    assert by_day[t + timedelta(days=5)] == D(-10000) and by_day[t + timedelta(days=7)] == D(40000)
    assert fc.first_below_zero == t + timedelta(days=5)  # денег хватит до 18.10
    assert fc.first_below_threshold == t + timedelta(days=5)
    assert [(g.date_from, g.date_to, g.min_balance, g.below_zero) for g in fc.gaps] == [
        (t + timedelta(days=5), t + timedelta(days=6), D(-10000), True)]
    assert items[4].status(t) == "overdue" and items[3].status(t) == "disabled"

    # перенос зарплаты после выплаты Ozon — разрыва ниже нуля нет, но ниже порога — да
    cal.apply_overrides(items, {"salary": cal.Override(day=t + timedelta(days=8)),
                                "rent": cal.Override(amount=D(55000))})
    assert items[2].moved and items[0].own_amount and items[0].amount == D(-55000)
    fc = cal.forecast(D(100000), items, t, t + timedelta(days=10), threshold=D(50000))
    assert fc.first_below_zero is None and fc.min_balance == D(45000)
    assert [(g.date_from, g.date_to, g.below_zero) for g in fc.gaps] == [
        (t + timedelta(days=2), t + timedelta(days=6), False), (t + timedelta(days=8), t + timedelta(days=10), False)]
    # включили рекламу — прогноз сразу другой
    items[3].enabled = True
    assert cal.forecast(D(100000), items, t, t + timedelta(days=10)).min_balance == D(35000)
    # пропуск — не платим
    cal.apply_overrides(items, {"rent": cal.Override(skipped=True)})
    assert items[0].status(t) == "skipped"
    # остаток неизвестен — прогноз движения без остатка, разрывов не ищем
    unknown = cal.forecast(None, items, t, t + timedelta(days=3))
    assert unknown.days[0].balance is None and unknown.gaps == [] and unknown.first_below_zero is None


def test_match_plan_fact():
    t = TODAY
    rent = item("rent", t - timedelta(days=1), "-30000", counterparty="ИП Петров")
    salary = item("salary", t + timedelta(days=2), "-50000", category_id="cat-salary")
    tax = item("tax", t - timedelta(days=10), "-6000")
    ozon = item("ozon", t, "40000", kind="income", source="ozon")
    marked = item("marked", t - timedelta(days=20), "-1000")
    ops = [
        cal.Op("op1", t - timedelta(days=2), D("-31000"), counterparty="Петров ИП"),  # +3 %, на день раньше
        cal.Op("op2", t, D("-50000"), category_id="other"),  # другая статья — не зарплата
        cal.Op("op3", t, D("-49000"), category_id="cat-salary"),  # заплатили раньше срока
        cal.Op("op4", t - timedelta(days=5), D("-6000")),  # 5 дней — дальше допуска
        cal.Op("op5", t, D("39500"), source="ozon_payout"),
        cal.Op("op6", t - timedelta(days=1), D("-900")),
    ]
    items = [rent, salary, tax, ozon, marked]
    cal.match(items, ops, {"marked": cal.Override(op_id="op6")})
    assert rent.fact.op_id == "op1" and rent.diff == D("-1000")  # план разошёлся с фактом — видно
    assert salary.fact.op_id == "op3" and salary.status(t) == "paid"
    assert tax.fact is None and tax.status(t) == "overdue"
    assert ozon.fact.op_id == "op5" and ozon.status(t) == "paid"  # прогноз Ozon закрыла фактическая выплата
    assert marked.fact.op_id == "op6" and marked.fact_marked
    assert cal.matched_ops(items) == {"op1", "op3", "op5", "op6"}
    # оплачено без операции — отметкой с датой и суммой
    other = item("x", t, "-700")
    cal.match([other], [], {"x": cal.Override(paid_day=t, paid_amount=D(700))})
    assert other.fact.amount == D(-700) and other.status(t) == "paid"


# ---------- сроки налогов ----------


def due(obs, title, start=None):
    return next(o for o in obs if o.title == title and (start is None or o.period_start == start))


def test_tax_deadlines_by_regime_and_weekend_shift():
    seg = Segment(date(2000, 1, 1), None, "usn_income", "ip", "none")
    obs = obligations([seg], date(2026, 1, 1), date(2027, 12, 31))
    q1 = due(obs, "Аванс УСН за I квартал", date(2026, 1, 1))
    assert (q1.due, q1.notice, q1.doc) == (date(2026, 4, 28), date(2026, 4, 27), "Уведомление")  # 25.04 — суббота
    assert due(obs, "Аванс УСН за 9 месяцев", date(2026, 7, 1)).notice == date(2026, 10, 26)  # 25.10 — воскресенье
    year = due(obs, "УСН за год", date(2026, 10, 1))
    assert (year.due, year.notice, year.doc) == (date(2027, 4, 28), date(2027, 4, 26), "Декларация")  # ИП — 28 апреля
    fixed = due(obs, "Страховые взносы ИП за год", date(2026, 1, 1))
    assert fixed.due == date(2026, 12, 28)
    assert due(obs, "Взносы ИП: 1 % с дохода свыше 300 000 ₽", date(2026, 1, 1)).due == date(2027, 7, 1)

    ooo = obligations([Segment(date(2000, 1, 1), None, "usn_income", "ooo")], date(2027, 3, 1), date(2027, 3, 31))
    assert [(o.title, o.due) for o in ooo] == [("УСН за год", date(2027, 3, 29))]  # 28.03.2027 — воскресенье

    osn = obligations([Segment(date(2000, 1, 1), None, "osn", "ooo", "general")], date(2026, 10, 1), date(2026, 12, 31))
    vat = [(o.title, o.due) for o in osn if o.what == "vat"]
    assert vat == [("НДС за квартал — 1/3", date(2026, 10, 28)), ("НДС за квартал — 2/3", date(2026, 11, 30)),
                   ("НДС за квартал — 3/3", date(2026, 12, 28))]  # 28.11 — суббота
    assert all(o.share == D(1) / 3 for o in osn if o.what == "vat")
    assert due(osn, "Аванс по налогу на прибыль за 9 месяцев").due == date(2026, 10, 28)
    assert not any(o.what == "contributions" for o in osn)  # взносы не указаны — у ООО нет и сроков

    ndfl = obligations([Segment(date(2000, 1, 1), None, "osn", "ip", "general")], date(2026, 7, 1), date(2027, 1, 31))
    assert [(o.title, o.due) for o in ndfl if o.what == "tax"] == [
        ("Аванс НДФЛ за январь–июнь", date(2026, 7, 15)), ("Аванс НДФЛ за июль–сентябрь", date(2026, 10, 15)),
        ("Аванс НДФЛ за октябрь–декабрь", date(2027, 1, 15))]

    ausn = obligations([Segment(date(2000, 1, 1), None, "ausn_income", "ooo")], date(2026, 10, 1), date(2026, 11, 30))
    assert [(o.title, o.period_start, o.due) for o in ausn] == [
        ("АУСН за месяц", date(2026, 9, 1), date(2026, 10, 26)), ("АУСН за месяц", date(2026, 10, 1), date(2026, 11, 25))]
    npd = obligations([Segment(date(2000, 1, 1), None, "npd", "self_employed")], date(2026, 11, 1), date(2026, 11, 30))
    assert [(o.title, o.due) for o in npd] == [("НПД за месяц", date(2026, 11, 30))]  # 28.11 — суббота

    patent = obligations([Segment(date(2026, 1, 1), None, "psn", "ip")], date(2026, 1, 1), date(2026, 12, 31))
    assert [(o.title, o.due, o.share) for o in patent if o.what == "patent"] == [
        ("Патент — 1/3 стоимости", date(2026, 4, 1), D(1) / 3), ("Патент — 2/3 стоимости", date(2026, 12, 28), D(2) / 3)]
    custom = obligations([Segment(date(2000, 1, 1), None, "custom")], date(2026, 1, 1), date(2026, 12, 31))
    assert custom == []  # «свой процент» — не режим: сроков нет


def test_regime_change_mid_period_splits_payment():
    segs = [Segment(date(2000, 1, 1), date(2026, 8, 14), "usn_income", "ooo"),
            Segment(date(2026, 8, 15), None, "ausn_income", "ooo")]
    obs = obligations(segs, date(2026, 9, 1), date(2026, 10, 31))
    assert [(o.title, o.period_start, o.period_end, o.due) for o in obs] == [
        ("АУСН за месяц", date(2026, 8, 15), date(2026, 8, 31), date(2026, 9, 25)),
        ("АУСН за месяц", date(2026, 9, 1), date(2026, 9, 30), date(2026, 10, 26)),
        ("Аванс УСН за 9 месяцев", date(2026, 7, 1), date(2026, 8, 14), date(2026, 10, 28)),
    ]


# ---------- API: планы, переопределения, план-факт ----------


async def test_plans_versions_and_validation(client, tokens):
    cats, accounts = await setup(client, tokens)
    body = {"kind": "expense", "amount": "30000", "start_date": "2026-10-05", "period": "month", "day": 5,
            "category_id": cats["rent"], "account_id": accounts["Расчётный счёт"], "counterparty": "ИП Петров"}
    saved = await call(client, tokens, "POST", "/plans", 201, json=body)
    plans = await call(client, tokens, "GET", "/plans")
    assert len(plans) == 1 and plans[0]["title"] == "Аренда"  # без названия — по статье
    assert plans[0]["schedule"] == "раз в месяц, 5-го числа, с 05.10.2026" and plans[0]["next_day"] == "2026-11-05"
    changed = await call(client, tokens, "PATCH", f"/plans/{saved['plan_id']}", json={"amount": "32000"})
    assert (changed["version"], changed["changed"]) == (2, True)
    same = await call(client, tokens, "PATCH", f"/plans/{saved['plan_id']}", json={"amount": "32000"})
    assert same["changed"] is False
    off = await call(client, tokens, "PATCH", f"/plans/{saved['plan_id']}", json={"enabled": False})
    history = await call(client, tokens, "GET", f"/plans/{saved['plan_id']}/history")
    assert [(h["version"], h["amount"], h["enabled"]) for h in history] == [
        (3, "32000.00", False), (2, "32000.00", True), (1, "30000.00", True)]  # прежние версии — в истории
    assert off["version"] == 3

    errors = {
        "Статья «Аренда» — для расходов: выберите другую или поменяйте «Приход / Расход»": {**body, "kind": "income"},
        "Дата окончания раньше даты начала": {**body, "end_date": "2026-01-01"},
        "День недели — от 1 (понедельник) до 7 (воскресенье)": {**body, "period": "week", "day": 9},
    }
    for message, data in errors.items():
        r = await client.post(f"{URL}/plans", json=data, headers=auth(tokens))
        assert (r.status_code, r.json()["detail"]) == (400, message)
    named = await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "expense", "amount": "990", "start_date": "2026-10-20", "category_name": "Хостинг сайта"})
    assert named["category_note"] == "created"  # вписанная статья сохранилась в список
    await call(client, tokens, "DELETE", f"/plans/{saved['plan_id']}", 204)
    assert [p["title"] for p in await call(client, tokens, "GET", "/plans")] == ["Хостинг сайта"]


async def test_calendar_overrides_paid_and_fact(client, tokens):
    cats, accounts = await setup(client, tokens)
    await call(client, tokens, "PUT", "/settings", json={"gen_ozon": False, "gen_taxes": False, "threshold": "25000"})
    rent = await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "expense", "title": "Аренда склада", "amount": "30000", "start_date": "2026-10-16", "period": "month",
        "category_id": cats["rent"], "counterparty": "ИП Петров"})
    await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "expense", "title": "Помощник", "amount": "10000", "start_date": "2026-10-02", "period": "week",
        "day": 5, "category_id": cats["salary"]})
    body = await calendar(client, tokens, date_to="2026-11-30")
    rows = [(ln["day"], ln["title"], ln["amount"], ln["status"]) for ln in lines(body)]
    assert rows[:3] == [("2026-10-16", "Аренда склада", "-30000.00", "planned"),
                        ("2026-10-16", "Помощник", "-10000.00", "planned"),
                        ("2026-10-23", "Помощник", "-10000.00", "planned")]
    assert body["start_balance"] == "100000.00"
    assert {ln["source"] for ln in lines(body)} == {"manual"}
    by_day = {d["day"]: d["balance"] for d in body["days"]}
    assert by_day["2026-10-16"] == "60000.00" and by_day["2026-11-16"] == "-10000.00"
    assert body["first_below_zero"] == "2026-11-16" and body["first_below_threshold"] == "2026-11-13"
    assert body["gaps"][0]["date_from"] == "2026-11-13" and body["gaps"][0]["below_zero"] is True  # порог 25 000

    # перенос аренды за ноябрь на 20.11 и своя сумма; пропуск помощника 30.10
    nov = f"plan:{rent['plan_id']}:2026-11-16"
    await call(client, tokens, "POST", "/calendar/override", 204, json={"key": nov, "day": "2026-11-20", "amount": "25000"})
    helper = next(ln for ln in lines(body) if ln["title"] == "Помощник" and ln["day"] == "2026-10-30")
    await call(client, tokens, "POST", "/calendar/override", 204, json={"key": helper["key"], "skipped": True})
    body = await calendar(client, tokens, date_to="2026-11-30")
    moved = next(ln for ln in lines(body) if ln["key"] == nov)
    assert (moved["day"], moved["planned_day"], moved["amount"], moved["moved"], moved["own_amount"]) == (
        "2026-11-20", "2026-11-16", "-25000.00", True, True)
    assert next(ln for ln in lines(body) if ln["key"] == helper["key"])["status"] == "skipped"
    assert body["first_below_zero"] == "2026-11-20"  # перенос сразу пересчитал прогноз: денег хватит до 19.11
    r = await client.post(f"{URL}/calendar/override", json={"key": nov, "day": "2027-06-01"}, headers=auth(tokens))
    assert r.json()["detail"] == "Перенести платёж можно не дальше чем на 90 дней от даты по плану"
    await call(client, tokens, "POST", "/calendar/override", 204, json={"key": nov, "reset": True})
    assert next(ln for ln in lines(await calendar(client, tokens)) if ln["key"] == nov)["day"] == "2026-11-16"

    # «оплачено» — операция-факт; план потом поменяли — факт тот же, расхождение видно
    oct16 = f"plan:{rent['plan_id']}:2026-10-16"
    paid = await call(client, tokens, "POST", "/calendar/paid", json={"key": oct16, "day": "2026-10-14"})
    assert paid["created"] is True
    ops = await call(client, tokens, "GET", "/operations")
    op = next(o for o in ops["items"] if o["op_id"] == paid["op_id"])
    assert (op["amount"], op["category"]["name"], op["counterparty"]) == ("-30000.00", "Аренда", "ИП Петров")
    await call(client, tokens, "PATCH", f"/plans/{rent['plan_id']}", json={"amount": "33000"})
    body = await calendar(client, tokens, date_to="2026-11-30")
    line = next(ln for ln in lines(body) if ln["key"] == oct16)
    assert (line["status"], line["fact"]["amount"], line["fact"]["marked"], line["diff"]) == (
        "paid", "-30000.00", True, "3000.00")
    assert body["start_balance"] == "70000.00"  # оплата уже в остатке — в прогнозе второй раз не вычитается
    assert {d["day"]: d["balance"] for d in body["days"]}["2026-10-16"] == "60000.00"
    await call(client, tokens, "POST", "/calendar/unpaid", 204, json={"key": oct16})
    line = next(ln for ln in lines(await calendar(client, tokens)) if ln["key"] == oct16)
    assert line["status"] == "paid" and line["fact"]["marked"] is False  # операция осталась — нашлась по сумме и дате

    # операция без плана — «вне плана»; переводы не показываем
    await call(client, tokens, "POST", "/operations", 201, json={
        "type": "expense", "amount": "1500", "day": "2026-10-14", "account_id": accounts["Расчётный счёт"]})
    facts = (await calendar(client, tokens))["facts"]
    assert [(f["amount"], f["title"]) for f in facts] == [("-1500.00", "Операция")]


async def test_generators_ozon_supplier_and_hints(client, tokens, sessions):
    cats, _ = await setup(client, tokens, supplier="after_sale")
    uid = await user_id(client, tokens)
    weeks = {w: v * 7 for w, v in WEEKS.items()}
    facts = {p.day: p.amount for p in history(weeks)}
    await add_days(sessions, uid, facts, {**weeks, monday(TODAY): D(7000)}, TODAY)
    await import_reference_costs(client, tokens)
    # продажи: по 2 шт. в день с 28.09 по 13.10 — себестоимость с НДС 2 076,44 ₽ в день
    await add_sales(sessions, uid, days(date(2026, 9, 28), date(2026, 10, 13)), "5000", 2, OFFER)
    await call(client, tokens, "PUT", "/settings", json={"gen_taxes": False})

    body = await calendar(client, tokens, date_to="2026-11-08")
    ozon = lines(body, "ozon")
    assert ozon and ozon[0]["day"] == "2026-10-21" and ozon[0]["kind"] == "income" and not ozon[0]["editable"]
    assert ozon[0]["title"] == "Выплата Ozon за 28.09–04.10" and ozon[0]["amount"] == "25000.00"  # не больше баланса
    supplier = lines(body, "supplier")
    first = supplier[0]  # за прошлую неделю — в понедельник
    assert (first["day"], first["title"], first["amount"], first["estimate"]) == (
        "2026-10-19", "Оплата поставщику за продажи 12.10–18.10", "-14535.08", True)  # 12–13.10 факт, дальше — средние
    assert first["note"].startswith("Себестоимость проданного с НДС поставщика")
    r = await client.post(f"{URL}/calendar/override", json={"key": ozon[0]["key"], "day": "2026-10-22"},
                          headers=auth(tokens))
    assert r.status_code == 400 and r.json()["detail"].startswith("Выплаты Ozon приходят по графику Ozon")

    # раз в месяц, 10-го: за сентябрь — только 28–30.09 с продажами (дни без данных ≠ 0 — по средним)
    await call(client, tokens, "PUT", "/settings", json={"supplier_period": "month", "supplier_day": 10})
    body = await calendar(client, tokens, date_to="2026-11-30")
    assert [ln["title"] for ln in lines(body, "supplier")] == ["Оплата поставщику за продажи октябрь 2026"]
    # выключили генераторы — их строк нет («что если»)
    await call(client, tokens, "PUT", "/settings", json={"gen_ozon": False, "gen_supplier": False})
    body = await calendar(client, tokens, date_to="2026-11-30")
    assert lines(body, "ozon") == [] and lines(body, "supplier") == []
    assert body["generators"]["supplier"] == {"enabled": False, "available": True, "period": "month", "day": 10}


async def test_deferred_purchase_and_own_production_hints(client, tokens):
    cats, _ = await setup(client, tokens, supplier="deferred", supplier_deferral_days=30)
    saved = await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "expense", "title": "Партия подшипников", "amount": "200000", "start_date": "2026-10-20",
        "purchase_day": "2026-10-20", "category_id": cats["supplier"]})
    body = await calendar(client, tokens, date_to="2026-12-31")
    line = lines(body, "supplier")[0]
    assert (line["day"], line["schedule"]) == ("2026-11-19", "закупка 20.10.2026, оплата 19.11.2026 (через 30 дн.)")
    assert body["deferral_days"] == 30 and body["generators"]["supplier"]["available"] is False
    # отсрочку поменяли — платёж сдвинулся сам
    await call(client, tokens, "PATCH", "", base="/api/v1/settings",
               json={"supplier_payment": "deferred", "supplier_deferral_days": 45})
    assert lines(await calendar(client, tokens, date_to="2026-12-31"), "supplier")[0]["day"] == "2026-12-04"
    assert (await call(client, tokens, "GET", "/plans"))[0]["plan_id"] == saved["plan_id"]

    await call(client, tokens, "PATCH", "", base="/api/v1/settings", json={"supplier_payment": "own_production"})
    hints = (await calendar(client, tokens))["hints"]
    assert [h["code"] for h in hints] == ["materials", "production_salary", "workshop_rent", "equipment"]
    assert all(h["why"] for h in hints)


# ---------- налоги в календаре, «Ближайшие сроки», смена режима ----------

USN = {"regime": "usn_income", "legal_form": "ooo", "vat_mode": "none"}


async def test_tax_payments_and_upcoming_block(client, tokens, sessions):
    empty = await call(client, tokens, "GET", "/upcoming", base="/api/v1/tax-profile")
    assert empty == {"has_profile": False, "planning": False, "items": []}  # профиль не задан — блока нет
    uid = await user_id(client, tokens)
    await add_sales(sessions, uid, days(date(2026, 7, 1), date(2026, 9, 30)), "1000")  # III квартал — 92 000 ₽
    await add_sales(sessions, uid, days(date(2026, 10, 1), date(2026, 10, 13)), "2000")
    await call(client, tokens, "PUT", "", base="/api/v1/tax-profile", json=USN)

    block = await call(client, tokens, "GET", "/upcoming", base="/api/v1/tax-profile")
    assert (block["has_profile"], block["planning"]) == (True, False)  # без планирования — сроки всё равно нужны
    first, second = block["items"][:2]
    assert (first["title"], first["day"], first["amount"], first["estimate"]) == (
        "Аванс УСН за 9 месяцев", "2026-10-28", "-5520.00", False)  # 92 000 × 6 %
    assert first["period_label"] == "III кв. 2026" and first["calc"] == {"date_from": "2026-07-01",
                                                                        "date_to": "2026-09-30"}
    assert first["notice"] == {"day": "2026-10-26", "doc": "Уведомление"}
    # IV квартал: 13 дней по 2 000 из 92 — оценка в том же темпе
    assert (second["title"], second["day"], second["amount"], second["estimate"]) == (
        "УСН за год", "2027-03-29", "-11040.00", True)
    assert second["note"] == "Оценка: данные за 13 дн. из 92, остальные дни — в том же темпе"

    # «оплачено» без планирования — отметка с датой и суммой
    await call(client, tokens, "POST", "/calendar/paid", json={"key": first["key"], "day": "2026-10-14"})
    block = await call(client, tokens, "GET", "/upcoming", base="/api/v1/tax-profile")
    assert block["items"][0]["status"] == "paid" and block["items"][0]["fact"]["amount"] == "-5520.00"

    # с планированием — платежи в календаре со ссылкой на расчёт
    await setup(client, tokens)
    block = await call(client, tokens, "GET", "/upcoming", base="/api/v1/tax-profile")
    assert block["planning"] is True
    body = await calendar(client, tokens, date_to="2026-11-30")
    taxes = lines(body, "tax")
    assert [(ln["title"], ln["day"]) for ln in taxes] == [("Аванс УСН за 9 месяцев", "2026-10-28")]
    assert body["generators"]["taxes"] == {"enabled": True, "available": True, "period": None, "day": None}


async def test_regime_change_keep_or_recalc_past_and_paid_unchanged(client, tokens, sessions):
    uid = await user_id(client, tokens)
    await add_sales(sessions, uid, days(date(2026, 4, 1), date(2026, 10, 13)), "1000")
    await call(client, tokens, "PUT", "", base="/api/v1/tax-profile", json=USN)
    await setup(client, tokens)
    window = {"date_from": "2026-07-01", "date_to": "2026-11-30"}

    def taxes(body):
        return {(ln["title"], ln["day"]): ln for ln in lines(body, "tax")}

    before = taxes(await calendar(client, tokens, **window))
    h1 = before[("Аванс УСН за полугодие", "2026-07-28")]
    assert (h1["amount"], h1["status"]) == ("-5460.00", "overdue")  # 91 000 × 6 %
    assert before[("Аванс УСН за 9 месяцев", "2026-10-28")]["amount"] == "-5520.00"

    # задним числом — АУСН «Доходы» 8 % с 1 июля: что изменится в прошлых платежах
    ausn = {"regime": "ausn_income", "legal_form": "ooo", "vat_mode": "none", "effective_from": "2026-07-01"}
    changes = await call(client, tokens, "POST", "/past-payments", base="/api/v1/tax-profile", json=ausn)
    assert [(c["title"], c["day"], c["old"], c["new"]) for c in changes] == [
        ("АУСН за месяц", "2026-08-25", None, "2480.00"),  # июль 31 000 × 8 %
        ("АУСН за месяц", "2026-09-25", None, "2480.00"),
    ]
    assert not await call(client, tokens, "POST", "/past-payments", base="/api/v1/tax-profile",
                          json={**ausn, "effective_from": "2026-10-14"})  # с сегодня прошлое не меняется

    # «нет» — прошлые остаются как были, будущие — по новому режиму
    await call(client, tokens, "PUT", "", base="/api/v1/tax-profile", json={**ausn, "past_payments": "keep"})
    after = taxes(await calendar(client, tokens, **window))
    assert ("АУСН за месяц", "2026-08-25") not in after and ("Аванс УСН за полугодие", "2026-07-28") in after
    assert ("Аванс УСН за 9 месяцев", "2026-10-28") not in after  # срок впереди — пересоздан по АУСН
    assert after[("АУСН за месяц", "2026-10-26")]["amount"] == "-2400.00"  # сентябрь 30 000 × 8 %
    assert after[("АУСН за месяц", "2026-11-25")]["estimate"] is True


async def test_regime_change_recalc_past_but_paid_stays(client, tokens, sessions):
    uid = await user_id(client, tokens)
    await add_sales(sessions, uid, days(date(2026, 4, 1), date(2026, 10, 13)), "1000")
    await call(client, tokens, "PUT", "", base="/api/v1/tax-profile", json=USN)
    await setup(client, tokens)
    window = {"date_from": "2026-07-01", "date_to": "2026-11-30"}
    h1 = next(ln for ln in lines(await calendar(client, tokens, **window), "tax") if ln["day"] == "2026-07-28")
    await call(client, tokens, "POST", "/calendar/paid", json={"key": h1["key"]})  # заплатили 5 460 ₽ в срок

    # ставку подняли задним числом с 1 апреля: «да» — пересчитать прошлые
    body = {**USN, "rate_pct": "7", "effective_from": "2026-04-01"}
    changes = await call(client, tokens, "POST", "/past-payments", base="/api/v1/tax-profile", json=body)
    assert changes == []  # единственный прошлый платёж оплачен — оплаченные не меняются
    await call(client, tokens, "PUT", "", base="/api/v1/tax-profile", json={**body, "past_payments": "recalc"})
    lines_after = {ln["day"]: ln for ln in lines(await calendar(client, tokens, **window), "tax")}
    paid = lines_after["2026-07-28"]
    assert (paid["status"], paid["fact"]["amount"], paid["amount"], paid["diff"]) == (
        "paid", "-5460.00", "-6370.00", "910.00")  # факт тот же, план пересчитан — расхождение видно
    assert lines_after["2026-10-28"]["amount"] == "-6440.00"  # 92 000 × 7 %


async def test_plans_template_round_trip(client, tokens):
    cats, accounts = await setup(client, tokens)
    await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "expense", "title": "Аренда", "amount": "30000", "start_date": "2026-10-05", "period": "month",
        "day": 5, "category_id": cats["rent"], "account_id": accounts["Расчётный счёт"]})
    await call(client, tokens, "POST", "/plans", 201, json={
        "kind": "income", "title": "Возврат займа", "amount": "5000", "start_date": "2026-11-01", "enabled": False})
    r = await client.get(f"{URL}/plans/template", headers=auth(tokens))
    assert r.status_code == 200
    content = r.content

    async def upload(data, dry_run=True):
        r = await client.post(f"{URL}/plans/import", params={"dry_run": dry_run}, headers=auth(tokens),
                              files={"file": ("plans.xlsx", data, "application/octet-stream")})
        assert r.status_code == 200, r.text
        return r.json()

    same = await upload(content)
    assert same["counts"] == {"new": 0, "changed": 0, "unchanged": 2, "skipped": 0, "error": 0}  # туда-обратно

    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Плановые платежи"]
    header = [c.value for c in ws[1]]
    col = {name: header.index(name) + 1 for name in header}
    ws.cell(row=2, column=col["Расход, руб."], value=35000)
    ws.append(["", "Зарплата", None, 50000, "раз в месяц", "пт", "10.11.2026", None, None, "Нет такой", None, None, "да"])
    ws.append(["", "Подписка", None, 990, "раз в неделю", "ср", "14.10.2026", None, None, None, None, None, "нет"])
    buf = io.BytesIO()
    wb.save(buf)
    preview = await upload(buf.getvalue())
    assert preview["counts"] == {"new": 1, "changed": 1, "unchanged": 1, "skipped": 0, "error": 1}
    error = next(r for r in preview["rows"] if r["status"] == "error")
    assert error["message"].startswith("Статьи «Нет такой» нет")
    applied = await upload(buf.getvalue(), dry_run=False)
    assert applied["applied"] is True
    plans = {p["title"]: p for p in await call(client, tokens, "GET", "/plans")}
    assert plans["Аренда"]["amount"] == "35000.00" and plans["Аренда"]["version"] == 2
    assert (plans["Подписка"]["period"], plans["Подписка"]["day"], plans["Подписка"]["enabled"]) == ("week", 3, False)
