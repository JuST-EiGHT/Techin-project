"""Шаг 10.17: «Остатки» — главная таблица по товарам, складам и кластерам: доступно, резерв, в пути, дни запаса, дата
окончания, стоимость запаса по себестоимости (нет себестоимости — «нет данных», не 0), фильтры, сортировка, Excel."""
import io
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from app.engine.stocks import end_date, stock_value
from app.services.sync.finance import today_msk
from tests.test_sync_stocks import SPB, auth, ozon  # noqa: F401 — фикстура ozon
from tests.test_warehouses import ACB, CM, RS, TH, seller  # noqa: F401

URL = "/api/v1/stocks/table"
D = Decimal
T = date(2026, 10, 24)


# ── эталоны (BUSINESS_LOGIC §10, «Запасы») ──
def test_end_date():
    assert end_date(31, D("6.2"), T) == T + timedelta(days=5)  # 5,0 дня запаса
    assert end_date(10, D("3"), T) == T + timedelta(days=3)  # 3,33 — закончится на четвёртый день, последний полный — 3-й
    assert end_date(0, D("1"), T) == T  # уже нет
    assert end_date(5, D("0"), T) is None  # не продаётся — даты нет
    assert end_date(5, None, T) is None  # продаж в данных нет — не знаем


def test_stock_value():
    assert stock_value(31, D("120")) == D("3720.00")
    assert stock_value(3, D("33.335")) == D("100.01")
    assert stock_value(0, D("120")) == D("0.00")
    assert stock_value(5, None) is None  # нет себестоимости — не 0


# ── через API ──
async def costs(client, tokens):
    await client.put(f"/api/v1/costs/{TH}", json={"cost_price": "100", "vat_multiplier": "1.2"}, headers=auth(tokens))
    await client.put(f"/api/v1/costs/{ACB}", json={"cost_price": "200", "vat_multiplier": "1"}, headers=auth(tokens))


def brief(data) -> list[tuple]:
    return [(r["product"]["offer_id"], (r["warehouse"] or {}).get("name") or r["cluster"], r["available"], r["stock_value"])
            for r in data["rows"]]


async def test_requires_auth(client):
    for url in (URL, f"{URL}/export"):
        assert (await client.get(url)).status_code == 401


async def test_by_products(client, tokens, seller):  # noqa: F811
    await costs(client, tokens)
    data = (await client.get(URL, headers=auth(tokens))).json()
    today = today_msk()
    assert (data["group"], data["sales_known"], data["stocks_known"]) == ("product", True, True)
    assert brief(data) == [(RS, None, 0, None), (CM, None, 1, None), (TH, None, 31, "3720.00"), (ACB, None, 5, "1000.00")]
    rs, cm, th, acb = data["rows"]
    # TH: 31 шт. по 6,2 в день — 5 дней, закончится через 5 дней; три склада, резерв 2, в пути 4
    assert (th["places"], th["reserved"], th["in_transit"], th["days_of_supply"], th["unit_cost"]) == (3, 2, 4, "5.0", "120.00")
    assert th["end_date"] == (today + timedelta(days=5)).isoformat()
    assert (rs["end_date"], rs["in_transit"], rs["status"]) == (today.isoformat(), 20, "critical")  # кончился, а продаётся
    assert (acb["end_date"], acb["sells"], acb["days_without_sales"]) == (None, False, 45)  # не продаётся — даты нет
    assert (cm["stock_value"], cm["unit_cost"]) == (None, None)  # нет себестоимости — «нет данных»
    assert data["totals"] == {"rows": 4, "products": 4, "available": 37, "in_transit": 25, "stock_value": "4720.00",
                              "without_cost": 1}


async def test_by_warehouses_and_clusters(client, tokens, seller):  # noqa: F811
    await costs(client, tokens)
    data = (await client.get(URL, params={"group": "warehouse"}, headers=auth(tokens))).json()
    assert brief(data) == [
        (RS, "КРАСНОДАР_2_РФЦ", 0, None), (CM, "Склад Подольск", 1, None),
        (TH, "ЕКАТЕРИНБУРГ_РФЦ_НОВЫЙ", 9, "1080.00"), (TH, "СПБ_БУГРЫ_РФЦ", 10, "1200.00"),
        (TH, "Склад Подольск", 12, "1440.00"), (ACB, "СПБ_БУГРЫ_РФЦ", 5, "1000.00"),
    ]
    # дни запаса и дата окончания — по товару в целом: в начислениях Ozon склада нет
    th_ekb = data["rows"][2]
    assert (th_ekb["product_available"], th_ekb["days_of_supply"], th_ekb["warehouse"]["cluster_name"]) == (31, "5.0", "Екатеринбург")
    assert (data["totals"]["rows"], data["totals"]["products"], data["totals"]["stock_value"]) == (6, 4, "4720.00")

    data = (await client.get(URL, params={"group": "cluster"}, headers=auth(tokens))).json()
    assert [(r["product"]["offer_id"], r["cluster"], r["cluster_type"], r["available"]) for r in data["rows"]] == [
        (RS, "Краснодар", "FBO", 0), (CM, "Склад Подольск", "FBS", 1),
        (TH, "Екатеринбург", "FBO", 9), (TH, "Санкт-Петербург и СЗО", "FBO", 10), (TH, "Склад Подольск", "FBS", 12),
        (ACB, "Санкт-Петербург и СЗО", "FBO", 5),
    ]


async def test_filters_and_sort(client, tokens, seller):  # noqa: F811
    await costs(client, tokens)
    get = lambda **q: client.get(URL, params=q, headers=auth(tokens))  # noqa: E731
    offers = lambda data: [r["product"]["offer_id"] for r in data["rows"]]  # noqa: E731
    assert offers((await get(status="critical")).json()) == [RS, TH]
    assert offers((await get(status="excess")).json()) == [ACB]
    assert offers((await get(status="no_sales")).json()) == [ACB]  # продаж за 30 дней нет, а лежит
    assert offers((await get(search="cm.nac")).json()) == [CM]
    whs = {w["ozon_warehouse_id"]: w["id"] for w in (await client.get("/api/v1/warehouses", headers=auth(tokens))).json()}
    data = (await get(warehouse_id=whs[SPB])).json()
    assert brief(data) == [(TH, None, 10, "1200.00"), (ACB, None, 5, "1000.00")]  # по товарам — только этот склад

    # пустые значения — в конце при любом порядке
    assert offers((await get(sort="stock_value", order="desc")).json()) == [TH, ACB, RS, CM]
    assert offers((await get(sort="stock_value", order="asc")).json()) == [ACB, TH, RS, CM]
    assert offers((await get(sort="end_date")).json()) == [RS, TH, CM, ACB]
    assert offers((await get(sort="offer_id", order="desc")).json()) == [ACB, TH, CM, RS]
    # нулевые остатки (товар ушёл со склада) — по галочке
    whole = (await get(group="warehouse")).json()["totals"]["rows"]
    assert (await get(group="warehouse", include_empty=True)).json()["totals"]["rows"] >= whole

    r = await get(group="shelf")
    assert r.status_code == 400
    assert (await get(sort="price")).status_code == 400


async def test_export_xlsx(client, tokens, seller):  # noqa: F811
    await costs(client, tokens)
    r = await client.get(f"{URL}/export", params={"group": "warehouse"}, headers=auth(tokens))
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:4] == ("Артикул", "Название", "Склад или кластер", "Доступно, шт.")
    by = {(row[0], row[2]): row for row in rows[1:7]}
    assert by[(TH, "СПБ_БУГРЫ_РФЦ")][3] == 10 and by[(TH, "СПБ_БУГРЫ_РФЦ")][10] == 1200
    assert by[(CM, "Склад Подольск")][10] == "нет себестоимости"
    assert by[(ACB, "СПБ_БУГРЫ_РФЦ")][8] == "не продаётся"
    total = next(row for row in rows if row[0] == "Итого")
    assert (total[3], total[10]) == (37, 4720)


async def test_summary_without_cost_and_isolation(client, tokens, seller):  # noqa: F811
    data = (await client.get("/api/v1/warehouses/summary", headers=auth(tokens))).json()
    assert ([r["product"]["offer_id"] for r in data["no_cost"]], data["no_cost_total"]) == ([TH, ACB, CM], 3)
    await costs(client, tokens)
    data = (await client.get("/api/v1/warehouses/summary", headers=auth(tokens))).json()
    assert ([r["product"]["offer_id"] for r in data["no_cost"]], data["no_cost_total"]) == ([CM], 1)

    other = {"email": "other@example.com", "password": "strongpass2"}
    await client.post("/api/v1/auth/register", json=other)
    t2 = (await client.post("/api/v1/auth/login", json=other)).json()
    data = (await client.get(URL, headers=auth(t2))).json()
    assert (data["rows"], data["stocks_known"], data["totals"]["stock_value"]) == ([], False, None)
