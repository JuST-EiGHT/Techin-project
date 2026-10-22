"""Планирование (шаг 3.8): мастер и профиль, счета и остатки, операции с версиями, переводы, выплаты Ozon, изоляция."""
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from sqlalchemy import func, select

from app.api.v1 import planning as planning_api
from app.models.ozon_balance import OzonBalanceDay
from app.models.planning import MoneyAccountBalance
from app.services import planning as svc
from tests.test_report_pnl import auth

TODAY = date(2026, 10, 14)
URL = "/api/v1/planning"


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    monkeypatch.setattr(svc, "today_msk", lambda: TODAY)
    monkeypatch.setattr(planning_api, "today_msk", lambda: TODAY)


async def call(client, tokens, method, path, expect=200, **kw):
    r = await client.request(method, f"{URL}{path}", headers=auth(tokens), **kw)
    assert r.status_code == expect, r.text
    return r.json() if r.content else None


async def setup(client, tokens, **body):
    data = {
        "supplier_payment": "own_production",
        "accounts": [{"name": "Расчётный счёт", "kind": "checking", "balance": "100000"},
                     {"name": "Наличные", "kind": "cash"}],
        "ozon_account": 0,
        "categories": ["salary", "materials"],
        **body,
    }
    return await call(client, tokens, "POST", "/setup", json=data)


async def ids(client, tokens):
    accounts = {a["name"]: a for a in await call(client, tokens, "GET", "/accounts")}
    cats = {c["code"] or c["name"]: c for c in await call(client, tokens, "GET", "/categories")}
    return accounts, cats


async def test_setup_wizard_profile_and_categories(client, tokens):
    state = await call(client, tokens, "GET", "/settings")
    assert (state["enabled"], state["setup_done"]) == (False, False)
    r = await client.put(f"{URL}/settings", json={"enabled": True}, headers=auth(tokens))
    assert (r.status_code, r.json()["detail"]) == (
        400, "Раздел ещё не настроен — ответьте на несколько вопросов на странице «Планирование»"
    )
    options = await call(client, tokens, "GET", "/setup")
    marketplace = next(c for c in options["categories"] if c["code"] == "marketplace")
    assert marketplace["required"] and "выплаты Ozon" in marketplace["why"]  # у статьи — зачем она
    assert options["groups_by_supplier"]["own_production"] == ["common", "production"]

    state = await setup(client, tokens)
    assert (state["enabled"], state["setup_done"], state["supplier_payment"]) == (True, True, "own_production")
    assert state["cost_label"] == "Себестоимость производства"  # подпись себестоимости в P&L
    settings = (await client.get("/api/v1/settings", headers=auth(tokens))).json()
    assert settings["cost_label"] == "Себестоимость производства"
    accounts, cats = await ids(client, tokens)
    assert set(cats) == {"marketplace", "transfer", "salary", "materials"}  # обязательные + отмеченные
    assert accounts["Расчётный счёт"]["receives_ozon"] and accounts["Расчётный счёт"]["balance"] == "100000.00"
    assert accounts["Наличные"]["balance"] is None and accounts["Наличные"]["opening"] is None  # не указан — не 0
    r = await client.post(f"{URL}/setup", json={}, headers=auth(tokens))
    assert r.status_code == 400

    # выключить — данные остаются; включить снова — без мастера
    state = await call(client, tokens, "PUT", "/settings", json={"enabled": False})
    assert (state["enabled"], state["setup_done"]) == (False, True)
    state = await call(client, tokens, "PUT", "/settings", json={"enabled": True})
    assert state["enabled"] is True and len(await call(client, tokens, "GET", "/accounts")) == 2


async def test_skipped_questions_change_nothing(client, tokens):
    state = await call(client, tokens, "POST", "/setup", json={})
    assert state["supplier_payment"] == "upfront" and state["ozon_account_id"] is None
    accounts, cats = await ids(client, tokens)
    assert list(accounts) == ["Расчётный счёт"] and accounts["Расчётный счёт"]["balance"] is None
    assert set(cats) == {"marketplace", "transfer"}  # ничего не предлагали — только нужные сервису
    r = await client.post(f"{URL}/setup", json={"supplier_payment": "deferred"}, headers=auth(tokens))
    assert r.status_code == 400  # уже настроено


async def test_deferred_needs_days_and_bad_template_code(client, tokens):
    r = await client.post(f"{URL}/setup", json={"supplier_payment": "deferred"}, headers=auth(tokens))
    assert (r.status_code, r.json()["detail"]) == (
        400, "Укажите отсрочку оплаты поставщику — через сколько дней после закупки платите"
    )
    r = await client.post(f"{URL}/setup", json={"categories": ["nope"]}, headers=auth(tokens))
    assert (r.status_code, r.json()["detail"]) == (400, "Статья не из шаблона: nope")
    state = await call(client, tokens, "POST", "/setup", json={"supplier_payment": "deferred", "supplier_deferral_days": 30})
    assert (state["supplier_payment"], state["supplier_deferral_days"]) == ("deferred", 30)


async def test_balances_versions_and_hiding(client, tokens, sessions):
    await setup(client, tokens)
    accounts, cats = await ids(client, tokens)
    main, salary = accounts["Расчётный счёт"]["id"], cats["salary"]["id"]
    op = await call(client, tokens, "POST", "/operations", 201, json={
        "type": "expense", "amount": "2500", "day": "2026-10-14", "account_id": main, "category_id": salary,
        "counterparty": "Иванов", "comment": "аванс",
    })
    await call(client, tokens, "POST", "/operations", 201, json={
        "type": "income", "amount": "1000", "day": "2026-10-10", "account_id": main,
    })
    future = await client.post(f"{URL}/operations", json={"type": "income", "amount": "5", "day": "2026-10-15",
                                                          "account_id": main}, headers=auth(tokens))
    assert future.status_code == 400 and "в будущем" in future.json()["detail"]
    page = await call(client, tokens, "GET", "/operations")
    assert [i["amount"] for i in page["items"]] == ["-2500.00", "1000.00"]
    assert page["items"][1]["before_opening"] is True  # до начального остатка — в остаток не входит
    assert page["totals"] == {"income": "1000.00", "expense": "-2500.00", "net": "-1500.00"}
    accounts, _ = await ids(client, tokens)
    assert accounts["Расчётный счёт"]["balance"] == "97500.00"

    # исправление — новая версия; ничего не изменилось — без версии
    saved = await call(client, tokens, "PATCH", f"/operations/{op['op_id']}", json={"amount": "3000"})
    assert (saved["version"], saved["changed"]) == (2, True)
    same = await call(client, tokens, "PATCH", f"/operations/{op['op_id']}", json={"amount": "3000", "comment": "аванс"})
    assert same["changed"] is False
    accounts, _ = await ids(client, tokens)
    assert accounts["Расчётный счёт"]["balance"] == "97000.00"
    await call(client, tokens, "DELETE", f"/operations/{op['op_id']}", 204)
    history = await call(client, tokens, "GET", f"/operations/{op['op_id']}/history")
    assert [(h["version"], h["amount"], h["deleted_at"] is not None) for h in history] == [
        (2, "-3000.00", True), (1, "-2500.00", False)
    ]  # удалённая — скрыта, но в истории
    accounts, _ = await ids(client, tokens)
    assert accounts["Расчётный счёт"]["balance"] == "100000.00"

    # новый начальный остаток — новая запись; операции с его даты добавляются
    acc = await call(client, tokens, "PATCH", f"/accounts/{main}", json={"balance": "50000", "as_of": "2026-10-01"})
    assert acc["balance"] == "51000.00" and acc["opening"] == {"amount": "50000.00", "as_of": "2026-10-01"}
    async with sessions() as db:
        assert await db.scalar(select(func.count()).select_from(MoneyAccountBalance)) == 2


async def test_transfers_are_not_income_or_expense(client, tokens):
    await setup(client, tokens)
    accounts, _ = await ids(client, tokens)
    main, cash = accounts["Расчётный счёт"]["id"], accounts["Наличные"]["id"]
    await call(client, tokens, "PATCH", f"/accounts/{cash}", json={"balance": "0"})
    tr = await call(client, tokens, "POST", "/operations", 201, json={
        "type": "transfer", "amount": "10000", "day": "2026-10-14", "account_id": main, "to_account_id": cash,
    })
    page = await call(client, tokens, "GET", "/operations")
    assert page["total"] == 1 and page["totals"]["net"] == "0.00"  # перевод — одной строкой, не доход и не расход
    item = page["items"][0]
    assert (item["amount"], item["transfer"]["direction"], item["transfer"]["account"]["name"]) == (
        "-10000.00", "out", "Наличные"
    )
    cash_page = await call(client, tokens, "GET", "/operations", params={"account_id": cash})
    assert cash_page["items"][0]["amount"] == "10000.00" and cash_page["items"][0]["transfer"]["direction"] == "in"
    accounts, _ = await ids(client, tokens)
    assert (accounts["Расчётный счёт"]["balance"], accounts["Наличные"]["balance"]) == ("90000.00", "10000.00")
    await call(client, tokens, "PATCH", f"/operations/{tr['op_id']}", json={"amount": "12000"})
    accounts, _ = await ids(client, tokens)
    assert (accounts["Расчётный счёт"]["balance"], accounts["Наличные"]["balance"]) == ("88000.00", "12000.00")
    bad = await client.patch(f"{URL}/operations/{tr['op_id']}", json={"account_id": cash}, headers=auth(tokens))
    assert bad.status_code == 400
    await call(client, tokens, "DELETE", f"/operations/{tr['op_id']}", 204)  # скрывает обе части
    accounts, _ = await ids(client, tokens)
    assert (accounts["Расчётный счёт"]["balance"], accounts["Наличные"]["balance"]) == ("100000.00", "0.00")
    same = await client.post(f"{URL}/operations", json={"type": "transfer", "amount": "1", "day": "2026-10-14",
                                                        "account_id": main, "to_account_id": main}, headers=auth(tokens))
    assert (same.status_code, same.json()["detail"]) == (400, "Перевод — между разными счетами")


async def add_payout(sessions, uid, day: date, amount: str):
    async with sessions() as db:
        db.add(OzonBalanceDay(user_id=uid, day=day, payments=-Decimal(amount), accrued=Decimal(100),
                              synced_at=datetime.now(timezone.utc)))
        await db.commit()


async def test_ozon_payouts_become_operations_without_duplicates(client, tokens, sessions):
    uid = uuid.UUID((await client.get("/api/v1/auth/me", headers=auth(tokens))).json()["id"])
    await add_payout(sessions, uid, date(2026, 10, 7), "8000")
    await add_payout(sessions, uid, date(2026, 10, 14), "9000")
    await add_payout(sessions, uid, date(2026, 10, 13), "0")  # день без выплаты
    await setup(client, tokens)
    page = await call(client, tokens, "GET", "/operations")
    ozon = [i for i in page["items"] if i["source"] == "ozon_payout"]
    assert [(i["day"], i["amount"], i["category"]["name"], i["account"]["name"]) for i in ozon] == [
        ("2026-10-14", "9000.00", "Поступления от маркетплейсов", "Расчётный счёт"),
        ("2026-10-07", "8000.00", "Поступления от маркетплейсов", "Расчётный счёт"),
    ]
    assert ozon[1]["before_opening"] is True  # выплата до начального остатка — уже в нём
    accounts, _ = await ids(client, tokens)
    assert accounts["Расчётный счёт"]["balance"] == "109000.00"  # 100 000 на начало дня + выплата сегодня
    page = await call(client, tokens, "GET", "/operations")
    assert sum(i["source"] == "ozon_payout" for i in page["items"]) == 2  # повторная синхронизация — без дублей

    # Ozon уточнил сумму — новая версия; свой комментарий пользователя сохраняется
    op_id = ozon[0]["op_id"]
    await call(client, tokens, "PATCH", f"/operations/{op_id}", json={"comment": "пришло утром"})
    bad = await client.patch(f"{URL}/operations/{op_id}", json={"amount": "1"}, headers=auth(tokens))
    assert bad.status_code == 400 and "обновляется сама" in bad.json()["detail"]
    assert (await client.delete(f"{URL}/operations/{op_id}", headers=auth(tokens))).status_code == 400
    await add_payout(sessions, uid, date(2026, 10, 14), "9100")
    page = await call(client, tokens, "GET", "/operations")
    item = next(i for i in page["items"] if i["op_id"] == op_id)
    assert (item["amount"], item["comment"], item["version"]) == ("9100.00", "пришло утром", 3)

    # выплаты приходят на другой счёт — операции переезжают туда
    accounts, _ = await ids(client, tokens)
    cash = accounts["Наличные"]["id"]
    bad = await client.patch(f"{URL}/accounts/{accounts['Расчётный счёт']['id']}", json={"archived": True},
                             headers=auth(tokens))
    assert bad.status_code == 400 and "выплаты Ozon" in bad.json()["detail"]
    state = await call(client, tokens, "PUT", "/settings", json={"ozon_account_id": cash})
    assert state["ozon_account_id"] == cash
    page = await call(client, tokens, "GET", "/operations", params={"account_id": cash})
    assert {i["amount"] for i in page["items"]} == {"9100.00", "8000.00"}


async def test_categories_and_isolation(client, tokens):
    await setup(client, tokens)
    accounts, cats = await ids(client, tokens)
    r = await client.patch(f"{URL}/categories/{cats['marketplace']['id']}", json={"archived": True}, headers=auth(tokens))
    assert (r.status_code, r.json()["detail"]) == (400, "Эта статья нужна сервису: в неё попадают выплаты Ozon")
    own = await call(client, tokens, "POST", "/categories", 201, json={"name": "Упаковка", "kind": "expense"})
    dup = await client.post(f"{URL}/categories", json={"name": "упаковка", "kind": "expense"}, headers=auth(tokens))
    assert (dup.status_code, dup.json()["detail"]) == (400, "Статья «Упаковка» уже есть — выберите её в списке")
    added = await call(client, tokens, "POST", "/categories/from-template", json={"codes": ["rent", "salary"]})
    assert sum(c["code"] == "rent" for c in added) == 1 and sum(c["code"] == "salary" for c in added) == 1
    archived = await call(client, tokens, "PATCH", f"/categories/{own['id']}", json={"archived": True})
    assert archived["archived"] is True
    r = await client.post(f"{URL}/operations", json={"type": "expense", "amount": "5", "day": "2026-10-14",
                                                     "account_id": accounts["Наличные"]["id"],
                                                     "category_id": own["id"]}, headers=auth(tokens))
    assert (r.status_code, r.json()["detail"]) == (400, "Статья «Упаковка» в архиве — выберите другую")
    op = await call(client, tokens, "POST", "/operations", 201, json={
        "type": "income", "amount": "5", "day": "2026-10-14", "account_id": accounts["Наличные"]["id"],
    })

    # другой пользователь не видит и не меняет чужое
    login = {"email": "other@example.com", "password": "strongpass1"}
    await client.post("/api/v1/auth/register", json=login)
    other = (await client.post("/api/v1/auth/login", json=login)).json()
    assert (await client.get(f"{URL}/accounts", headers=auth(other))).json() == []
    assert (await client.patch(f"{URL}/operations/{op['op_id']}", json={"amount": "1"}, headers=auth(other))).status_code == 404
    assert (await client.patch(f"{URL}/accounts/{accounts['Наличные']['id']}", json={"name": "x"},
                               headers=auth(other))).status_code == 404
    r = await client.post(f"{URL}/operations", json={"type": "income", "amount": "5", "day": "2026-10-14",
                                                     "account_id": accounts["Наличные"]["id"]}, headers=auth(other))
    assert (r.status_code, r.json()["detail"]) == (404, "Счёт не найден")
    page = await call(client, other, "GET", "/operations", params={"date_from": "2026-01-01"})
    assert page["total"] == 0
    assert (await client.get(f"{URL}/operations", params={"date_from": "2026-10-14", "date_to": "2026-10-01"},
                             headers=auth(tokens))).status_code == 400


# ---------- операции списком: шаблон туда-обратно ----------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def upload(client, tokens, content: bytes, name="ops.xlsx", apply=False):
    r = await client.post(f"{URL}/operations/import", params={"dry_run": str(not apply).lower()},
                          files={"file": (name, content, XLSX if name.endswith("xlsx") else "text/csv")},
                          headers=auth(tokens))
    assert r.status_code == 200, r.text
    return r.json()


async def test_template_round_trip_and_changes(client, tokens, sessions):
    import io

    from openpyxl import load_workbook

    uid = uuid.UUID((await client.get("/api/v1/auth/me", headers=auth(tokens))).json()["id"])
    await add_payout(sessions, uid, date(2026, 10, 14), "9000")  # выплата Ozon — в шаблон не входит
    await setup(client, tokens)
    accounts, cats = await ids(client, tokens)
    main, cash = accounts["Расчётный счёт"]["id"], accounts["Наличные"]["id"]
    await call(client, tokens, "POST", "/operations", 201, json={
        "type": "expense", "amount": "2500", "day": "2026-10-12", "account_id": main, "category_id": cats["salary"]["id"],
        "counterparty": "Иванов", "comment": "аванс",
    })
    await call(client, tokens, "POST", "/operations", 201, json={
        "type": "income", "amount": "700", "day": "2026-10-13", "account_id": cash, "comment": "возврат",
    })
    await call(client, tokens, "POST", "/operations", 201, json={
        "type": "transfer", "amount": "1000", "day": "2026-10-13", "account_id": main, "to_account_id": cash,
    })  # перевод — тоже не в шаблоне
    r = await client.get(f"{URL}/operations/template", params={"date_from": "2026-10-01"}, headers=auth(tokens))
    assert r.status_code == 200
    original = r.content
    wb = load_workbook(io.BytesIO(original))
    ws = wb["Операции"]
    assert [c.value for c in ws[1]] == ["ID (не меняйте)", "Дата", "Приход, руб.", "Расход, руб.", "Счёт", "Статья",
                                        "Контрагент", "Описание"]
    assert ws.max_row == 3 and ws["D2"].value == 2500 and ws["C3"].value == 700  # расход — без минуса
    lists = [dv.formula1 for dv in ws.data_validations.dataValidation if dv.type == "list"]
    assert '"Расчётный счёт,Наличные"' in lists  # счета пользователя — списком

    # скачанный файл загружается обратно без ошибок и без изменений
    body = await upload(client, tokens, original)
    assert body["counts"] == {"new": 0, "changed": 0, "unchanged": 2, "skipped": 0, "error": 0}

    ws["D2"] = 3000  # исправили сумму
    ws.append([None, date(2026, 10, 14), None, 1200, "Расчётный счёт", "Зарплата", None, "премия"])  # новая
    ws.append([None, date(2026, 10, 14), 50, None, "Сейф", None, None, None])  # нет такого счёта
    ws.append([None, date(2026, 10, 14), 50, 60, "Наличные", None, None, None])  # и приход, и расход
    ws.append([None, date(2026, 10, 20), 50, None, "Наличные", None, None, None])  # в будущем
    ws.append([None, date(2026, 10, 14), None, None, "Наличные", None, None, "без суммы"])
    buf = io.BytesIO()
    wb.save(buf)
    changed = buf.getvalue()
    body = await upload(client, tokens, changed)
    assert body["counts"] == {"new": 1, "changed": 1, "unchanged": 1, "skipped": 1, "error": 3}
    errors = {r["line"]: r["message"] for r in body["rows"] if r["status"] == "error"}
    assert errors == {
        5: "Счёта «Сейф» нет — добавьте его в «Операциях» → «Счета и выписки» или выберите из списка",
        6: "Заполните что-то одно: «Приход» или «Расход»",
        7: "Дата в будущем — здесь только то, что уже было",
    }
    body = await upload(client, tokens, changed, apply=True)
    assert body["applied"] is True
    page = await call(client, tokens, "GET", "/operations", params={"date_from": "2026-10-01"})
    amounts = sorted(i["amount"] for i in page["items"] if i["source"] in ("manual", "import") and not i["transfer"])
    assert amounts == ["-1200.00", "-3000.00", "700.00"]
    # тот же файл ещё раз — без дублей: новая строка уже есть, исправленная — без изменений
    body = await upload(client, tokens, changed)
    assert (body["counts"]["new"], body["counts"]["changed"], body["counts"]["skipped"]) == (0, 0, 2)

    r = await client.post(f"{URL}/operations/import/errors", files={"file": ("ops.xlsx", changed, XLSX)},
                          headers=auth(tokens))
    err = load_workbook(io.BytesIO(r.content))["Операции"]
    assert err["A1"].value == "Ошибка" and err.max_row == 4 and err["A2"].value.startswith("Счёта «Сейф» нет")


async def test_csv_with_signed_sum_and_single_account(client, tokens):
    await setup(client, tokens, accounts=[{"name": "Карта", "kind": "card", "balance": "1000"}])
    csv_text = "Дата;Сумма;Статья;Комментарий\n10.10.2026;-1500,50;Зарплата;тест\n11.10.2026;+300;;\n"
    body = await upload(client, tokens, csv_text.encode("utf-8"), name="ops.csv", apply=True)
    assert body["counts"]["new"] == 2 and body["counts"]["error"] == 0
    assert [(r["amount"], r["account"]) for r in body["rows"]] == [("-1500.50", "Карта"), ("300.00", "Карта")]
    page = await call(client, tokens, "GET", "/operations")
    assert {i["source"] for i in page["items"]} == {"import"}
    r = await client.post(f"{URL}/operations/import", files={"file": ("x.csv", b"a;b\n1;2\n", "text/csv")},
                          headers=auth(tokens))
    assert r.status_code == 400 and "Не найдены колонки «Дата» и «Приход» / «Расход»" in r.json()["detail"]


# ---------- доработки 3.8: быстрый старт, статья по названию, описание ----------


async def test_quick_setup_simple_categories(client, tokens):
    state = await call(client, tokens, "POST", "/setup", json={"quick": True})
    assert state["enabled"] and state["supplier_payment"] == "upfront"  # «Откуда товар» не трогаем
    accounts, cats = await ids(client, tokens)
    assert list(accounts) == ["Расчётный счёт"] and accounts["Расчётный счёт"]["receives_ozon"]
    assert {c["name"] for c in cats.values()} == {
        "Приход", "Расход", "Поступления от маркетплейсов", "Перевод между своими счетами"
    }


async def test_category_by_name_reuses_or_creates(client, tokens):
    await setup(client, tokens)
    accounts, cats = await ids(client, tokens)
    main = accounts["Расчётный счёт"]["id"]
    op = lambda **kw: call(client, tokens, "POST", "/operations", 201, json={  # noqa: E731
        "type": "expense", "amount": "100", "day": "2026-10-14", "account_id": main, **kw})
    saved = await op(category_name="  зарплата ")  # такая уже есть — без учёта регистра и пробелов
    assert (saved["category_note"], saved["category"]["id"]) == ("existing", cats["salary"]["id"])
    saved = await op(category_name="Упаковка")  # новой нет — сохраняется в список
    assert (saved["category_note"], saved["category"]["name"]) == ("created", "Упаковка")
    again = await op(category_name="упаковка")
    assert (again["category_note"], again["category"]["id"]) == ("existing", saved["category"]["id"])
    _, cats = await ids(client, tokens)
    assert sum(c["name"] == "Упаковка" for c in cats.values()) == 1  # без дублей
    dup = await client.post(f"{URL}/categories", json={"name": "УПАКОВКА", "kind": "expense"}, headers=auth(tokens))
    assert (dup.status_code, dup.json()["detail"]) == (400, "Статья «Упаковка» уже есть — выберите её в списке")
    # из архива — возвращается, а не создаётся вторая
    await call(client, tokens, "PATCH", f"/categories/{saved['category']['id']}", json={"archived": True})
    restored = await op(category_name="Упаковка")
    assert (restored["category_note"], restored["category"]["id"]) == ("restored", saved["category"]["id"])
    # приход с тем же названием — своя статья поступлений
    income = await call(client, tokens, "POST", "/operations", 201, json={
        "type": "income", "amount": "5", "day": "2026-10-14", "account_id": main, "category_name": "Упаковка"})
    assert income["category_note"] == "created" and income["category"]["id"] != saved["category"]["id"]
    # исправление операции — тоже можно вписать статью
    fixed = await call(client, tokens, "PATCH", f"/operations/{again['op_id']}", json={"category_name": "Аренда склада"})
    assert (fixed["category_note"], fixed["version"]) == ("created", 2)


async def test_long_description(client, tokens):
    await setup(client, tokens)
    accounts, _ = await ids(client, tokens)
    text = "Оплата упаковки для партии подшипников. " * 45
    body = {"type": "expense", "amount": "100", "day": "2026-10-14", "account_id": accounts["Наличные"]["id"]}
    saved = await call(client, tokens, "POST", "/operations", 201, json={**body, "comment": text[:2000]})
    page = await call(client, tokens, "GET", "/operations")
    assert page["items"][0]["comment"] == text[:2000].strip() and saved["changed"]
    r = await client.post(f"{URL}/operations", json={**body, "comment": "я" * 2001}, headers=auth(tokens))
    assert (r.status_code, r.json()["detail"]) == (400, "Описание: максимум 2000 символов")

