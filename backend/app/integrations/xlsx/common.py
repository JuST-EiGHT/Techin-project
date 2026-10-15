"""Общее для разбора XLSX-отчётов Ozon: чтение книги, текст ячеек, числа, период.

Таблицы ищутся по заголовкам (правило 4): номера листов, строк и колонок не важны.
"""
import io
import math
import re
import warnings as py_warnings
from collections.abc import Sequence
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

Row = Sequence[Any]

MAX_ROWS = 200_000  # на лист; отчёты Ozon за месяц — тысячи строк
_XLS_SIGNATURE = b"\xd0\xcf\x11\xe0"  # старый формат Excel 97–2003 (.xls)


class ReportParseError(Exception):
    """Файл не удалось разобрать. Текст — для пользователя, на русском."""


def norm(value: Any) -> str:
    """Текст ячейки для сравнения: переносы и повторные пробелы → один пробел, нижний регистр, ё → е."""
    return " ".join(cell_text(value).split()).lower().replace("ё", "е")


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # SKU и артикулы, сохранённые числом: 3145302232.0 → «3145302232»
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value).replace(" ", " ").strip()


def cell(row: Row, idx: int | None) -> Any:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def is_empty_row(row: Row) -> bool:
    return all(cell_text(v) == "" for v in row)


# ---------- числа ----------

_SPACES = re.compile(r"[\s    ]")
_CURRENCY = re.compile(r"(₽|руб\.?|р\.|rub)$", re.IGNORECASE)
_NUMBER = re.compile(r"[+-]?(\d+(\.\d*)?|\.\d+)")


def parse_number(value: Any, warnings: list[str] | None = None, context: str = "", where: str = "") -> Decimal | None:
    """Число из ячейки: «1 234,56», «1234.56», «1 234 ₽», «−304,01», «(15,00)».

    Пустая ячейка → None без предупреждения. Нечитаемое значение (#N/A, текст) → None
    и предупреждение в warnings: тихо подставлять 0 нельзя.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return _unreadable(value, warnings, context, where)
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return _unreadable(value, warnings, context, where)
        return Decimal(repr(value))  # кратчайшая запись: 482.66, а не 482.66000000000002
    if isinstance(value, Decimal):
        return value if value.is_finite() else _unreadable(value, warnings, context, where)
    if isinstance(value, (datetime, date, time)):
        return _unreadable(value, warnings, context, where)

    s = _SPACES.sub("", str(value))
    if not s:
        return None
    s = _CURRENCY.sub("", s).replace("₽", "")
    for minus in "−‒–—":
        s = s.replace(minus, "-")
    negative = False
    if s.startswith("(") and s.endswith(")"):  # бухгалтерская запись отрицательного числа
        negative, s = True, s[1:-1]
    if "," in s and "." in s:
        # правый разделитель — десятичный, другой — разряды: 1,234.56 / 1.234,56
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif s.count(",") == 1:
        s = s.replace(",", ".")
    elif s.count(",") > 1:
        s = s.replace(",", "")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    if not _NUMBER.fullmatch(s):
        return _unreadable(value, warnings, context, where)
    try:
        result = Decimal(s)
    except InvalidOperation:
        return _unreadable(value, warnings, context, where)
    return -result if negative else result


def _unreadable(value: Any, warnings: list[str] | None, context: str, where: str) -> None:
    if warnings is not None:
        shown = str(value)
        shown = shown if len(shown) <= 40 else shown[:40] + "…"
        field = f" в поле «{context}»" if context else ""
        warnings.append(f"Не удалось разобрать число «{shown}»{field}{where} — значение не учтено.")
    return None


def parse_qty(value: Any, warnings: list[str], where: str = "") -> int:
    """Количество штук. Пусто или 0 → 1, как в эталонном калькуляторе; нечитаемое → 1 с предупреждением."""
    number = parse_number(value, warnings, "Количество", where)
    if number is None or number == 0:
        return 1
    return int(number)


_ROW_NUMBER = re.compile(r"\d+(?:[.,]0+)?")


def parse_row_number(value: Any) -> int | None:
    """Порядковый номер строки таблицы (целое > 0) или None."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float):
        return int(value) if value.is_integer() and value > 0 else None
    s = cell_text(value)
    if not _ROW_NUMBER.fullmatch(s):
        return None
    n = int(re.split(r"[.,]", s)[0])
    return n if n > 0 else None


def is_numbering_row(row: Row) -> bool:
    """Строка с номерами столбцов (1, 2, 3 …) под заголовком таблицы — не данные."""
    nums = [parse_row_number(v) for v in row if cell_text(v)]
    return len(nums) >= 3 and all(n is not None for n in nums) and nums == list(range(nums[0], nums[0] + len(nums)))


_PARSE_WARNING_PREFIX = "Не удалось разобрать число"
MAX_NUMBER_WARNINGS = 20


def limit_warnings(warnings: list[str]) -> list[str]:
    """Предупреждения о нечитаемых числах — не больше 20, остальные одной строкой."""
    out: list[str] = []
    numbers = 0
    for w in warnings:
        if w.startswith(_PARSE_WARNING_PREFIX):
            numbers += 1
            if numbers > MAX_NUMBER_WARNINGS:
                continue
        out.append(w)
    if numbers > MAX_NUMBER_WARNINGS:
        out.append(f"…и ещё нечитаемых чисел: {numbers - MAX_NUMBER_WARNINGS}.")
    return out


# ---------- период ----------

_DATE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")
_PERIOD_PREFIX = re.compile(r"^.*?период\s*:?\s*", re.IGNORECASE)


def _to_date(match: re.Match) -> date | None:
    d, m, y = (int(g) for g in match.groups())
    try:
        return date(y, m, d)
    except ValueError:
        return None


_ISO_DATE = re.compile(r"(\d{4})-(\d{2})-(\d{2})")


def parse_date(value: Any) -> date | None:
    """Дата из ячейки: дата Excel, «07.09.2026» или «2026-09-07»; иначе None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = cell_text(value)
    if m := _DATE.search(text):
        return _to_date(m)
    if m := _ISO_DATE.search(text):
        y, mth, d = (int(g) for g in m.groups())
        try:
            return date(y, mth, d)
        except ValueError:
            return None
    return None


def extract_period(rows: Sequence[Row]) -> tuple[str, date | None, date | None]:
    """Период из шапки отчёта (строки над таблицей).

    Сначала ячейка со словом «период» («Период: 07.09.2026-13.09.2026», «Реализация товаров
    за период с 01.08.2026 по 31.08.2026»), иначе — первая ячейка с датой.
    """
    texts = [cell_text(v) for row in rows for v in row if cell_text(v)]
    candidates = (
        [t for t in texts if "период" in t.lower() and _DATE.search(t)]
        or [t for t in texts if "период" in t.lower()]
        or [t for t in texts if _DATE.search(t)]
    )
    if not candidates:
        return "", None, None
    text = candidates[0]
    period = _PERIOD_PREFIX.sub("", text).strip() if "период" in text.lower() else text.strip()
    dates = [_to_date(m) for m in _DATE.finditer(text)]
    start = dates[0] if dates else None
    end = dates[1] if len(dates) > 1 else start
    return period, start, end


# ---------- книга ----------


class Workbook:
    """Книга XLSX только для чтения: строки листа — кортежи значений (формулы — уже посчитанные значения)."""

    def __init__(self, content: bytes):
        if content[:4] == _XLS_SIGNATURE:
            raise ReportParseError(
                "Файл в старом формате Excel (.xls). Откройте его в Excel и сохраните как .xlsx, "
                "затем загрузите снова."
            )
        try:
            with py_warnings.catch_warnings():
                py_warnings.simplefilter("ignore")  # «Workbook contains no default style» у выгрузок Ozon
                self._wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — openpyxl бросает что угодно: BadZipFile, KeyError, XML-ошибки
            raise ReportParseError(
                "Не удалось открыть файл Excel. Нужен файл .xlsx, выгруженный из личного кабинета Ozon; "
                "если файл открывается в Excel — пересохраните его и загрузите снова."
            ) from exc
        self.sheet_names: list[str] = [ws.title for ws in self._wb.worksheets]  # без листов-диаграмм

    def rows(self, sheet_name: str, limit: int | None = None) -> list[tuple]:
        ws = self._wb[sheet_name]
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()  # размеры в файле бывают неверными — читаем до последней строки
        out: list[tuple] = []
        try:
            for row in ws.iter_rows(values_only=True):
                if limit is not None and len(out) >= limit:
                    break
                if len(out) >= MAX_ROWS:
                    raise ReportParseError(
                        f"На листе «{sheet_name}» больше {MAX_ROWS} строк — это не похоже на отчёт Ozon."
                    )
                out.append(row)
        except ReportParseError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ReportParseError(
                f"Не удалось прочитать лист «{sheet_name}»: файл повреждён. Выгрузите отчёт заново."
            ) from exc
        return out

    def close(self) -> None:
        self._wb.close()

    def __enter__(self) -> "Workbook":
        return self

    def __exit__(self, *exc) -> None:
        self.close()
