"""Чтение заполненного шаблона: XLSX или CSV, таблица ищется по заголовкам (правило 4).

Лист «Инструкция» и другие листы без нужных колонок не мешают. Формула без сохранённого
значения (файл собран программой, а не сохранён в Excel) отдаётся как FORMULA — строка получит
понятную ошибку, а не тихий пропуск.
"""

import csv
import io
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from app.integrations.xlsx.common import ReportParseError, Workbook, cell, norm
from app.services.templates.spec import TemplateSpec
from app.services.templates.values import FORMULA, is_empty

HEADER_SCAN = 20  # заголовок ищем в первых строках листа
_XLS_SIGNATURE = b"\xd0\xcf\x11\xe0"


class TemplateError(Exception):
    """Файл целиком не подходит. Текст — для пользователя."""


@dataclass
class TableRow:
    line: int  # строка файла, с 1 — как в Excel
    values: dict[str, Any]  # ключ колонки → значение ячейки (None — пусто, FORMULA — формула без значения)


@dataclass
class Table:
    format: str  # xlsx | csv
    columns: list[str]  # ключи найденных колонок
    rows: list[TableRow]


def _columns(spec: TemplateSpec, header_row: list) -> dict[str, int]:
    """Ключ колонки → индекс. Каждая ячейка заголовка — не больше чем одной колонке."""
    header = [norm(v) for v in header_row]
    found: dict[str, int] = {}
    for column in spec.read_columns:
        best: tuple[int, int] | None = None
        for j, text in enumerate(header):
            if j in found.values():
                continue
            priority = column.matches(text)
            if priority is not None and (best is None or priority < best[0]):
                best = (priority, j)
        if best is not None:
            found[column.key] = best[1]
    return found


def _has_required(spec: TemplateSpec, cols: dict[str, int]) -> bool:
    return all(any(key in cols for key in group) for group in spec.required)


def _header(spec: TemplateSpec, rows: list[list]) -> tuple[int, dict[str, int]] | None:
    for i, row in enumerate(rows[:HEADER_SCAN]):
        cols = _columns(spec, row)
        if _has_required(spec, cols):
            return i, cols
    return None


def _decode(content: bytes) -> str:
    if content[:2] in (b"\xff\xfe", b"\xfe\xff"):  # «Текст Юникод» из Excel — UTF-16 с табуляцией
        return content.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise TemplateError("Не удалось прочитать файл: нужен CSV в кодировке UTF-8 или Windows-1251, или XLSX.")


def _csv_rows(content: bytes) -> list[list]:
    text = _decode(content)
    first = text.splitlines()[0] if text else ""
    delimiter = max((";", "\t", ","), key=first.count)
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _formula_rows(content: bytes, sheet: str) -> list[tuple]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        return list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()


def _xlsx_table(spec: TemplateSpec, content: bytes) -> tuple[list[list], int, dict[str, int], list[tuple]]:
    try:
        with Workbook(content) as wb:
            if not wb.sheet_names:
                raise TemplateError("В файле нет листов.")
            for sheet in wb.sheet_names:  # первый лист, где нашлась таблица
                rows = [list(r) for r in wb.rows(sheet)]
                if found := _header(spec, rows):
                    return rows, found[0], found[1], _formula_rows(content, sheet)
    except ReportParseError as exc:
        raise TemplateError(str(exc)) from exc
    raise TemplateError(spec.missing_columns)


def read_table(spec: TemplateSpec, content: bytes, filename: str, max_rows: int) -> Table:
    name = filename.lower()
    if content[:4] == _XLS_SIGNATURE or name.endswith(".xls"):
        raise TemplateError("Файл в старом формате Excel (.xls). Сохраните его как .xlsx или .csv и загрузите снова.")
    formulas: list[tuple] = []
    if content[:2] == b"PK" or name.endswith(".xlsx"):
        fmt = "xlsx"
        rows, header_idx, cols, formulas = _xlsx_table(spec, content)
    else:
        fmt = "csv"
        rows = _csv_rows(content)
        found = _header(spec, rows)
        if found is None:
            raise TemplateError(spec.missing_columns)
        header_idx, cols = found

    out: list[TableRow] = []
    for i in range(header_idx + 1, len(rows)):
        values = {key: cell(rows[i], j) for key, j in cols.items()}
        if formulas:
            for key, j in cols.items():
                f = cell(formulas[i], j) if i < len(formulas) else None
                if values[key] is None and isinstance(f, str) and f.startswith("="):
                    values[key] = FORMULA
        if all(is_empty(v) for v in values.values()):
            continue
        if len(out) >= max_rows:
            raise TemplateError(f"В файле больше {max_rows} строк — разбейте его на части.")
        out.append(TableRow(line=i + 1, values=values))
    return Table(format=fmt, columns=list(cols), rows=out)
