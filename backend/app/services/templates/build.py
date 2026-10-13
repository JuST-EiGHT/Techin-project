"""Выгрузка шаблона в XLSX или CSV — с данными пользователя или пустого.

XLSX: лист с таблицей (закреплённый заголовок, ширины, у текстовых колонок формат «Текст»,
чтобы Excel не превращал артикулы в числа и даты, подсказки ввода и проверка значений) и
лист «Инструкция». CSV: те же колонки, UTF-8 с BOM, разделитель «;», запятая в дробных —
так файл открывается в русском Excel без мастера импорта.
"""

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.templates.spec import Column, TemplateSpec
from app.services.templates.values import csv_text

INSTRUCTION_SHEET = "Инструкция"
SPARE_ROWS = 1000  # столько пустых строк под таблицей тоже получают проверку значений

_FONT = {"name": "Calibri", "size": 11}  # шрифт по умолчанию в Excel; без имени LibreOffice берёт другой
_HEADER_FONT = Font(bold=True, **_FONT)
_HEADER_FILL = PatternFill("solid", fgColor="E3E6EE")
_MUTED = Font(color="808080", **_FONT)
_ERROR_FONT = Font(color="DC2626", **_FONT)
_FORMATS = {"text": "@", "rate": "@", "money": "#,##0.00", "date": "DD.MM.YYYY"}


def _validation(column: Column, ref: str) -> DataValidation | None:
    if column.choices:
        dv = DataValidation(type="list", formula1='"' + ",".join(column.choices) + '"', showErrorMessage=False)
    elif column.kind == "money":
        dv = DataValidation(
            type="decimal",
            operator="greaterThan",
            formula1="0",
            errorStyle="warning",
            error="Нужно число больше 0, например 1250,50",
            errorTitle="Проверьте сумму",
        )
    elif column.kind == "date":
        dv = DataValidation(
            type="date",
            operator="greaterThan",
            formula1="36526",
            errorStyle="warning",
            error="Нужна дата, например 01.09.2026",
            errorTitle="Проверьте дату",
        )
    else:
        return None
    if column.hint:
        dv.prompt, dv.promptTitle, dv.showInputMessage = column.hint[:255], column.title[:32], True
    dv.add(ref)
    return dv


def build_xlsx(spec: TemplateSpec, rows: list[dict[str, Any]], columns: list[Column] | None = None) -> bytes:
    columns = columns or list(spec.template_columns)
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet
    last_row = len(rows) + 1 + SPARE_ROWS
    for j, column in enumerate(columns, start=1):
        letter = get_column_letter(j)
        head = ws.cell(row=1, column=j, value=column.title)
        head.font, head.fill = _HEADER_FONT, _HEADER_FILL
        head.alignment = Alignment(wrap_text=True, vertical="center")
        dim = ws.column_dimensions[letter]
        dim.width = column.width
        fmt = _FORMATS[column.kind]
        dim.number_format = fmt  # новые строки, которые допишет пользователь, — в том же формате
        for i, row in enumerate(rows, start=2):
            c = ws.cell(row=i, column=j, value=row.get(column.key))
            c.number_format = fmt
            if column.key == "_error":
                c.font = _ERROR_FONT
            elif column.reference:
                c.font = _MUTED
        if not column.reference and (dv := _validation(column, f"{letter}2:{letter}{last_row}")):
            ws.add_data_validation(dv)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    _print_setup(ws)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    _instructions(wb.create_sheet(INSTRUCTION_SHEET), spec)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _print_setup(ws) -> None:
    """Печать: альбомная страница, по ширине — на один лист."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _instructions(ws, spec: TemplateSpec) -> None:
    """Лист «Инструкция». В одной строке — не больше одного названия колонки, чтобы этот лист
    никогда не приняли за таблицу."""
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    _print_setup(ws)
    ws.cell(row=1, column=1, value=f"Как заполнить лист «{spec.sheet}»").font = Font(bold=True, name="Calibri", size=14)
    r = 3
    for text in spec.instructions:
        ws.cell(row=r, column=1, value=text)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Колонки").font = _HEADER_FONT
    r += 1
    for column in spec.template_columns:
        title = ws.cell(row=r, column=1, value=column.title)
        title.font, title.alignment = (_MUTED if column.reference else _HEADER_FONT), Alignment(vertical="top")
        ws.cell(row=r, column=2, value=column.hint).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    for i, text in enumerate(spec.notes):
        note = ws.cell(row=r, column=1, value=text)
        if i == 0:
            note.font = _HEADER_FONT  # первая строка — заголовок блока
        r += 1


def build_csv(spec: TemplateSpec, rows: list[dict[str, Any]], columns: list[Column] | None = None) -> bytes:
    columns = columns or list(spec.template_columns)
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerow([c.title for c in columns])
    for row in rows:
        writer.writerow([csv_text(c.kind, row.get(c.key)) for c in columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def build_file(fmt: str, spec: TemplateSpec, rows: list[dict[str, Any]], columns: list[Column] | None = None) -> bytes:
    return build_csv(spec, rows, columns) if fmt == "csv" else build_xlsx(spec, rows, columns)


MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
}
