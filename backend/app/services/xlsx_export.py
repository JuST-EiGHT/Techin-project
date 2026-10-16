"""Выгрузка таблицы в Excel для чтения и работы — не шаблон для загрузки (шаблоны — services/templates, правило 14).

Лист с таблицей: жирный заголовок, закреплён, автофильтр, ширины, форматы чисел; под таблицей — итог, если задан.
Второй лист — пояснения (параметры отчёта, дата, как читать). Пользуются: «Кандидаты на вывоз» (5.1); дальше —
поставки для логиста (5.2).
"""
import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_FONT = {"name": "Calibri", "size": 11}
_HEADER_FONT = Font(bold=True, **_FONT)
_HEADER_FILL = PatternFill("solid", fgColor="E3E6EE")
_FORMATS = {"text": "@", "int": "#,##0", "number": "#,##0.0", "money": "#,##0.00", "date": "DD.MM.YYYY"}


@dataclass(frozen=True)
class Col:
    title: str
    width: int = 14
    kind: str = "text"  # text | int | number | money | date


def _value(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    return v


def table_xlsx(
    sheet: str,
    columns: list[Col],
    rows: list[list[Any]],
    *,
    total: list[Any] | None = None,
    notes_title: str = "Пояснения",
    notes: list[str] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    for j, col in enumerate(columns, start=1):
        head = ws.cell(row=1, column=j, value=col.title)
        head.font, head.fill = _HEADER_FONT, _HEADER_FILL
        head.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = col.width
    for i, row in enumerate(rows, start=2):
        for j, (col, v) in enumerate(zip(columns, row), start=1):
            c = ws.cell(row=i, column=j, value=_value(v))
            c.number_format = _FORMATS[col.kind] if not isinstance(v, date) else _FORMATS["date"]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    if total is not None:
        r = len(rows) + 3
        for j, (col, v) in enumerate(zip(columns, total), start=1):
            c = ws.cell(row=r, column=j, value=_value(v))
            c.font, c.number_format = _HEADER_FONT, _FORMATS[col.kind]
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if notes:
        info = wb.create_sheet(notes_title[:31])
        info.column_dimensions["A"].width = 120
        for r, text in enumerate(notes, start=1):
            c = info.cell(row=r, column=1, value=text)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if r == 1:
                c.font = Font(bold=True, name="Calibri", size=14)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
